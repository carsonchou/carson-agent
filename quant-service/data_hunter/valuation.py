# -*- coding: utf-8 -*-
"""
valuation.py — 台股數據獵手「財務估值模型」agent（Anthropic 金融 agent 範例的台股版第3支）

純計算，不靠 LLM，可單元測試。輸入股票代號 → 算出「合理價參考區間（便宜/合理/昂貴）＋
現價落點＋財務比率＋EPS 歷史」，供看板互動表看＋下載 Excel。

四法估值（各法可算出的才列，缺資料優雅降級，全路徑 try/except、絕不炸）：
  1. PE 法      歷史 PE 分位(P25/P50/P75)；不足 60 點退產業預設 PE 區間
  2. 殖利率法    歷史殖利率分位(cash_div_ttm/歷史收盤)；不足 60 點退等級預設(高息/一般)
  3. PB 法      歷史 PB 分位(BVPS 用現值近似固定，非精算)
  4. 成長模型    PEG(clamp 年增率 5-30 當合理 PE) ＋ 穩定配息股另算 Gordon 股利成長模型

定位：估值參考區間，非投資建議，不喊單一目標價（PERSONA 紅線一致）。

用法
  python valuation.py 2330      # 印估值結果(合理價區間/四法明細/財務比率/EPS歷史)
"""
from __future__ import annotations

import io
import json
import math
import statistics
import sys
import time
from datetime import date, datetime
from pathlib import Path

HERE = Path(__file__).resolve().parent

DISCLAIMER = "估值參考區間，基於公開財報推估，非投資建議，不保證。"

# ── 產業 PE 預設區間（歷史 PE 樣本不足 60 點時退用；非精算，粗略夠用） ─────────
_INDUSTRY_PE_DEFAULT = {"半導體": (15.0, 25.0), "金融": (10.0, 15.0), "傳產": (12.0, 18.0)}
_TRADITIONAL_KEYWORDS = (
    "水泥", "食品", "塑膠", "紡織", "電機機械", "鋼鐵", "橡膠", "汽車", "造紙",
    "建材", "營造", "化學", "玻璃", "陶瓷", "油電", "燃氣", "農業科技", "百貨",
)


def _industry_pe_band(industry: str | None) -> tuple[float, float, float]:
    """(cheap_pe, fair_pe, expensive_pe)；fair 取 cheap/expensive 中值。"""
    ind = industry or ""
    if "半導體" in ind:
        lo, hi = _INDUSTRY_PE_DEFAULT["半導體"]
    elif "金融" in ind or "保險" in ind or "保险" in ind:
        lo, hi = _INDUSTRY_PE_DEFAULT["金融"]
    elif any(k in ind for k in _TRADITIONAL_KEYWORDS):
        lo, hi = _INDUSTRY_PE_DEFAULT["傳產"]
    else:
        lo, hi = (12.0, 20.0)
    return lo, (lo + hi) / 2.0, hi


def _percentile(sorted_vals: list, pct: float):
    """線性內插百分位（同 numpy.percentile 預設法），sorted_vals 需已由小到大排序。"""
    if not sorted_vals:
        return None
    n = len(sorted_vals)
    if n == 1:
        return sorted_vals[0]
    k = (n - 1) * pct
    f, c = math.floor(k), math.ceil(k)
    if f == c:
        return sorted_vals[int(k)]
    return sorted_vals[f] * (c - k) + sorted_vals[c] * (k - f)


def _clamp(x, lo, hi):
    return max(lo, min(hi, x))


def _classify_position(price, cheap, fair, expensive) -> tuple[str, float]:
    """現價落點分5級 + 距合理價%；price/fair 缺任一 → ("資料不足", None)。
    cheap/expensive 缺任一時退用 fair 補位(不炸，區間退化成一點)。"""
    if price is None or fair is None:
        return "資料不足", None
    cheap = cheap if cheap is not None else fair
    expensive = expensive if expensive is not None else fair
    if price < cheap:
        pos = "便宜"
    elif price < fair * 0.95:
        pos = "偏低"
    elif price <= fair * 1.05:
        pos = "合理"
    elif price <= expensive:
        pos = "偏高"
    else:
        pos = "昂貴"
    gap = round((price / fair - 1.0) * 100, 1)
    return pos, gap


def _quarter_label(d: str) -> str:
    try:
        y, m = d.split("-")[0], d.split("-")[1]
        q = (int(m) - 1) // 3 + 1
        return f"{y}Q{q}"
    except Exception:
        return d


def _quarterly_eps(fs_rows: list, num_fn) -> list:
    """FinancialStatements rows → 由舊到新排序去重的 [(date, eps), ...]。"""
    dd: dict[str, float] = {}
    for r in fs_rows or []:
        if r.get("type") != "EPS":
            continue
        d = r.get("date")
        v = num_fn(r.get("value"))
        if d and v is not None:
            dd[d] = v
    return sorted(dd.items())


def _rolling_ttm(eps_q_series: list) -> list:
    """滾動近四季 EPS 合計 → [(date, eps_ttm), ...]（由舊到新，前3筆資料不足no輸出）。"""
    out = []
    for i in range(3, len(eps_q_series)):
        window = eps_q_series[i - 3:i + 1]
        out.append((eps_q_series[i][0], round(sum(v for _, v in window), 2)))
    return out


# ── 四法之一：PE 法 ──────────────────────────────────────────────────────────
def _method_pe(df, eps_ttm_series: list, eps_ttm_now, industry) -> dict | None:
    if df is None or not len(df) or not eps_ttm_now or eps_ttm_now <= 0:
        return None

    pe_series: list = []
    if eps_ttm_series:
        try:
            import pandas as pd
            ttm_df = pd.DataFrame(eps_ttm_series, columns=["date", "ttm"])
            ttm_df["date"] = pd.to_datetime(ttm_df["date"])
            ttm_df = ttm_df.sort_values("date").reset_index(drop=True)
            px_df = pd.DataFrame({"date": pd.to_datetime(df.index), "close": df["Close"].values})
            px_df = px_df.sort_values("date").reset_index(drop=True)
            merged = pd.merge_asof(px_df, ttm_df, on="date", direction="backward")
            merged = merged.dropna(subset=["ttm"])
            merged = merged[merged["ttm"] > 0]
            pe_series = sorted((merged["close"] / merged["ttm"]).tolist())
        except Exception:
            pe_series = []

    if len(pe_series) >= 60:
        cheap_pe = _percentile(pe_series, 0.25)
        fair_pe = _percentile(pe_series, 0.50)
        exp_pe = _percentile(pe_series, 0.75)
        source = "歷史PE分位"
        assume = f"歷史PE(P25/P50/P75)={round(cheap_pe, 1)}/{round(fair_pe, 1)}/{round(exp_pe, 1)}倍（{len(pe_series)}個交易日樣本）"
    else:
        cheap_pe, fair_pe, exp_pe = _industry_pe_band(industry)
        source = "產業預設"
        assume = f"歷史PE樣本不足({len(pe_series)}點)，退產業預設PE {cheap_pe}-{exp_pe}倍"

    return {
        "name": "PE法",
        "cheap": round(eps_ttm_now * cheap_pe, 2),
        "fair": round(eps_ttm_now * fair_pe, 2),
        "expensive": round(eps_ttm_now * exp_pe, 2),
        "assume": assume,
        "source": source,
    }


# ── 四法之二：殖利率法 ────────────────────────────────────────────────────────
def _method_dividend(df, fnd: dict) -> dict | None:
    cash_div_ttm = fnd.get("cash_div_ttm")
    if not cash_div_ttm or cash_div_ttm <= 0:
        return None
    if df is None or not len(df):
        return None
    closes = [float(c) for c in df["Close"].tolist() if c is not None and c > 0]
    if not closes:
        return None

    yields = sorted(cash_div_ttm / c for c in closes)
    if len(yields) >= 60:
        y_lo = _percentile(yields, 0.25)     # 低殖利率 → 貴
        y_mid = _percentile(yields, 0.50)
        y_hi = _percentile(yields, 0.75)     # 高殖利率 → 便宜
        source = "歷史殖利率分位"
        assume = (f"歷史殖利率(P25/P50/P75)={round(y_lo * 100, 2)}%/"
                  f"{round(y_mid * 100, 2)}%/{round(y_hi * 100, 2)}%（{len(yields)}個交易日樣本）")
    else:
        dy = fnd.get("dividend_yield")
        high_div = dy is not None and dy >= 5
        y_hi_pct, y_lo_pct = (7.0, 5.0) if high_div else (4.0, 3.0)
        y_hi, y_lo = y_hi_pct / 100.0, y_lo_pct / 100.0
        y_mid = (y_hi + y_lo) / 2.0
        source = "等級預設"
        assume = f"歷史樣本不足({len(yields)}點)，退{'高息股' if high_div else '一般股'}目標殖利率{y_lo_pct}-{y_hi_pct}%"

    if not y_hi or not y_mid or not y_lo:
        return None
    return {
        "name": "殖利率法",
        "cheap": round(cash_div_ttm / y_hi, 2),
        "fair": round(cash_div_ttm / y_mid, 2),
        "expensive": round(cash_div_ttm / y_lo, 2),
        "assume": assume,
        "source": source,
    }


# ── 四法之三：PB 法 ──────────────────────────────────────────────────────────
def _method_pb(df, price, fnd: dict) -> dict | None:
    pb = fnd.get("pb")
    if not pb or pb <= 0 or not price or price <= 0:
        return None
    if df is None or not len(df):
        return None
    bvps = price / pb
    if bvps <= 0:
        return None
    closes = [float(c) for c in df["Close"].tolist() if c is not None and c > 0]
    if not closes:
        return None

    pb_series = sorted(c / bvps for c in closes)
    cheap_pb = _percentile(pb_series, 0.25)
    fair_pb = _percentile(pb_series, 0.50)
    exp_pb = _percentile(pb_series, 0.75)
    n = len(pb_series)
    source = "歷史PB分位" if n >= 60 else f"歷史PB分位(僅{n}筆，樣本偏少)"
    assume = (f"BVPS≈{round(bvps, 2)}(現價/現PB推估，歷史BVPS變動未精算)；"
              f"歷史PB(P25/P50/P75)={round(cheap_pb, 2)}/{round(fair_pb, 2)}/{round(exp_pb, 2)}倍")
    return {
        "name": "PB法",
        "cheap": round(bvps * cheap_pb, 2),
        "fair": round(bvps * fair_pb, 2),
        "expensive": round(bvps * exp_pb, 2),
        "assume": assume,
        "source": source,
    }


# ── 四法之四：成長模型（PEG + 穩定配息股另算 Gordon） ─────────────────────────
def _estimate_dividend_growth(code: str, fundamentals_mod) -> float | None:
    """近5年年度現金股利 CAGR（首尾年皆須 >0）；資料不足回 None。"""
    five_years_ago = f"{date.today().year - 5}-01-01"
    dv = fundamentals_mod._finmind("TaiwanStockDividend", code, five_years_ago)
    if not dv:
        return None
    by_year: dict[str, float] = {}
    for r in dv:
        d = r.get("date") or ""
        y = d[:4]
        if not y:
            continue
        v = fundamentals_mod._num(r.get("CashEarningsDistribution"))
        if v and v > 0:
            by_year[y] = by_year.get(y, 0.0) + v
    years = sorted(by_year.keys())
    if len(years) < 2:
        return None
    first_y, last_y = years[0], years[-1]
    first_v, last_v = by_year[first_y], by_year[last_y]
    n = int(last_y) - int(first_y)
    if n <= 0 or first_v <= 0:
        return None
    try:
        return (last_v / first_v) ** (1.0 / n) - 1.0
    except Exception:
        return None


def _method_growth(code: str, eps_ttm_now, fnd: dict, fundamentals_mod) -> list:
    methods = []
    eps_yoy = fnd.get("eps_yoy")
    if eps_ttm_now and eps_ttm_now > 0 and eps_yoy is not None:
        peg_pe = _clamp(eps_yoy, 5.0, 30.0)
        fair = eps_ttm_now * peg_pe
        methods.append({
            "name": "成長模型(PEG)",
            "cheap": round(fair * 0.8, 2), "fair": round(fair, 2), "expensive": round(fair * 1.2, 2),
            "assume": f"合理PE≈年增率(clamp 5-30%)={round(peg_pe, 1)}倍；便宜/昂貴＝PEG 0.8/1.2倍",
            "source": "PEG推估",
        })

    div_freq = fnd.get("div_freq")
    cash_div_ttm = fnd.get("cash_div_ttm")
    if div_freq is not None and cash_div_ttm and cash_div_ttm > 0 and fundamentals_mod is not None:
        try:
            g_raw = _estimate_dividend_growth(code, fundamentals_mod)
        except Exception:
            g_raw = None
        if g_raw is not None:
            g = _clamp(g_raw, 0.0, 0.08)
            r = 0.09
            if g < r:
                d1 = cash_div_ttm * (1.0 + g)
                fair = d1 / (r - g)
                r_cheap = r + 0.01
                r_exp = r - 0.01
                cheap = d1 / (r_cheap - g)
                expensive = d1 / (r_exp - g) if r_exp > g else fair
                methods.append({
                    "name": "股利成長模型(Gordon)",
                    "cheap": round(cheap, 2), "fair": round(fair, 2), "expensive": round(expensive, 2),
                    "assume": f"D1={round(d1, 2)} 折現率r=9%±1% 股利成長率g={round(g * 100, 1)}%(近5年CAGR推估,clamp 0-8%)",
                    "source": "Gordon股利成長模型",
                })
    return methods


# ── 對外主函式 ───────────────────────────────────────────────────────────────
def build_valuation(code: str) -> dict:
    """代號 → 估值 dict（schema 見檔頭註解／README）。全路徑 try/except，缺資料優雅降級，絕不炸。"""
    ts = int(time.time())
    code_in = (code or "").strip()
    out = {
        "code": code_in, "name": code_in, "price": None, "ts": ts,
        "range": {"cheap": None, "fair": None, "expensive": None},
        "position": "資料不足", "gap_to_fair_pct": None,
        "methods": [], "ratios": {}, "eps_history": [],
        "disclaimer": DISCLAIMER,
    }
    if not code_in:
        return out

    resolved_code = code_in
    try:
        try:
            import fundamentals
        except Exception:
            fundamentals = None

        name, industry = code_in, "其他"
        try:
            import query
            resolved_code = query._resolve_code(code_in) or code_in
            name, industry = query._meta(resolved_code)
        except Exception:
            pass
        out["code"] = resolved_code
        out["name"] = name or resolved_code

        df = None
        try:
            import analyst
            df = analyst._load_ohlcv(resolved_code)
        except Exception:
            df = None

        price = None
        if df is not None and len(df):
            try:
                price = round(float(df["Close"].iloc[-1]), 2)
            except Exception:
                price = None
        out["price"] = price

        fnd: dict = {}
        rat: dict = {}
        if fundamentals is not None:
            try:
                fnd = fundamentals.load_fundamentals(resolved_code, offline=True) or {}
            except Exception:
                fnd = {}
            try:
                rat = fundamentals.load_financial_ratios(resolved_code) or {}
            except Exception:
                rat = {}
        out["ratios"] = {
            "current_ratio": rat.get("current_ratio"), "debt_ratio": rat.get("debt_ratio"),
            "fcf": rat.get("fcf"), "roe": rat.get("roe"),
            "gross_margin": fnd.get("gross_margin"), "op_margin": fnd.get("op_margin"),
            "dividend_yield": fnd.get("dividend_yield"), "grade": rat.get("grade") or {},
        }

        eps_q_series: list = []
        eps_ttm_series: list = []
        if fundamentals is not None:
            try:
                five_years_ago = f"{date.today().year - 5}-01-01"
                fs = fundamentals._finmind("TaiwanStockFinancialStatements", resolved_code, five_years_ago)
                eps_q_series = _quarterly_eps(fs, fundamentals._num)
                eps_ttm_series = _rolling_ttm(eps_q_series)
            except Exception:
                eps_q_series, eps_ttm_series = [], []
        out["eps_history"] = [{"q": _quarter_label(d), "eps": v} for d, v in eps_q_series[-8:]]

        eps_ttm_now = fnd.get("eps_ttm")
        if eps_ttm_now is None and eps_ttm_series:
            eps_ttm_now = eps_ttm_series[-1][1]

        methods: list = []
        try:
            m = _method_pe(df, eps_ttm_series, eps_ttm_now, industry)
            if m:
                methods.append(m)
        except Exception:
            pass
        try:
            m = _method_dividend(df, fnd)
            if m:
                methods.append(m)
        except Exception:
            pass
        try:
            m = _method_pb(df, price, fnd)
            if m:
                methods.append(m)
        except Exception:
            pass
        try:
            methods.extend(_method_growth(resolved_code, eps_ttm_now, fnd, fundamentals))
        except Exception:
            pass
        out["methods"] = methods

        cheap_vals = [m["cheap"] for m in methods if m.get("cheap") is not None]
        fair_vals = [m["fair"] for m in methods if m.get("fair") is not None]
        exp_vals = [m["expensive"] for m in methods if m.get("expensive") is not None]
        rng = {"cheap": None, "fair": None, "expensive": None}
        if cheap_vals:
            rng["cheap"] = round(statistics.median(cheap_vals), 2)
        if fair_vals:
            rng["fair"] = round(statistics.median(fair_vals), 2)
        if exp_vals:
            rng["expensive"] = round(statistics.median(exp_vals), 2)
        out["range"] = rng

        pos, gap = _classify_position(price, rng["cheap"], rng["fair"], rng["expensive"])
        out["position"] = pos
        out["gap_to_fair_pct"] = gap
    except Exception as e:
        out["_error"] = f"{type(e).__name__}: {e}"

    try:
        import ai_agents
        ai_agents.save_state(f"valuation_{out.get('code') or code_in}", out)
    except Exception:
        pass
    return out


# ── Excel 匯出 ───────────────────────────────────────────────────────────────
def build_valuation_xlsx(code: str) -> bytes:
    """4 sheet：①估值總表 ②財務比率 ③EPS歷史 ④假設與來源。openpyxl 未裝時拋明確例外(上層接住回錯)。"""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception as e:
        raise RuntimeError(f"openpyxl 未安裝，無法產生 Excel：{e}")

    data = build_valuation(code)
    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2F4F6F")

    def _style_header_row(ws, row_idx: int) -> None:
        for cell in ws[row_idx]:
            cell.font = hdr_font
            cell.fill = hdr_fill

    wb = Workbook()

    # ① 估值總表
    ws1 = wb.active
    ws1.title = "估值總表"
    ts = data.get("ts")
    when = datetime.fromtimestamp(ts).strftime("%Y-%m-%d %H:%M:%S") if ts else "—"
    ws1.append(["股票", f"{data.get('name')}（{data.get('code')}）"])
    ws1.append(["現價", data.get("price")])
    ws1.append(["估值時間", when])
    ws1.append([])
    rng = data.get("range") or {}
    ws1.append(["合理價區間", "便宜", "合理", "昂貴"])
    _style_header_row(ws1, ws1.max_row)
    ws1.append(["", rng.get("cheap"), rng.get("fair"), rng.get("expensive")])
    ws1.append(["現價落點", data.get("position")])
    ws1.append(["距合理價%", data.get("gap_to_fair_pct")])
    ws1.append([])
    ws1.append(["估值方法", "便宜", "合理", "昂貴", "假設", "來源"])
    _style_header_row(ws1, ws1.max_row)
    for m in data.get("methods") or []:
        ws1.append([m.get("name"), m.get("cheap"), m.get("fair"), m.get("expensive"),
                    m.get("assume"), m.get("source")])
    ws1.column_dimensions["A"].width = 16
    for col in ("B", "C", "D"):
        ws1.column_dimensions[col].width = 12
    ws1.column_dimensions["E"].width = 46
    ws1.column_dimensions["F"].width = 14

    # ② 財務比率
    ws2 = wb.create_sheet("財務比率")
    ratios = data.get("ratios") or {}
    grade = ratios.get("grade") or {}
    labels = [
        ("current_ratio", "流動比"), ("debt_ratio", "負債比%"), ("fcf", "自由現金流FCF"),
        ("roe", "ROE%(推估)"), ("gross_margin", "毛利率%"), ("op_margin", "營益率%"),
        ("dividend_yield", "殖利率%"),
    ]
    ws2.append(["指標", "數值", "評級(好/普/差)"])
    _style_header_row(ws2, 1)
    for k, lbl in labels:
        ws2.append([lbl, ratios.get(k), grade.get(k, "")])
    ws2.column_dimensions["A"].width = 18
    ws2.column_dimensions["B"].width = 14
    ws2.column_dimensions["C"].width = 14

    # ③ EPS 歷史
    ws3 = wb.create_sheet("EPS歷史")
    ws3.append(["季度", "EPS"])
    _style_header_row(ws3, 1)
    for row in data.get("eps_history") or []:
        ws3.append([row.get("q"), row.get("eps")])
    ws3.column_dimensions["A"].width = 12
    ws3.column_dimensions["B"].width = 10

    # ④ 假設與來源
    ws4 = wb.create_sheet("假設與來源")
    ws4.append(["估值方法", "假設", "來源"])
    _style_header_row(ws4, 1)
    for m in data.get("methods") or []:
        ws4.append([m.get("name"), m.get("assume"), m.get("source")])
    ws4.append([])
    ws4.append(["免責聲明", data.get("disclaimer")])
    ws4.column_dimensions["A"].width = 18
    ws4.column_dimensions["B"].width = 60
    ws4.column_dimensions["C"].width = 16

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


if __name__ == "__main__":
    _code = sys.argv[1] if len(sys.argv) > 1 else "2330"
    _r = build_valuation(_code)
    print(json.dumps(_r, ensure_ascii=False, indent=2))
