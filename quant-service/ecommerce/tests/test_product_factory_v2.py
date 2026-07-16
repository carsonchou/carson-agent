# -*- coding: utf-8 -*-
"""test_product_factory_v2.py — 一次性 SKU 工廠 v2 單元測試(不渲染 PDF、不連網)。

覆蓋:誠信守門拒絕未溯源數字、SKU 組裝完整性、listing 欄位完整、
tokenizer 切片對齊只補真來源大數字的子數(不放行憑空數字)、殖利率不再×100。
2026-07-16 增(釘死 VERIFY_REPORT_phase3b 的修復,防回歸):
  - TestBlockerRegressions:維度數不得硬編、存證不得空殼、英文守門不得失明、A4 紙張不得退回 Letter
  - TestMethodDisclosure:C1 方法揭露的每個常數都對回 strategy.py / tw_adaptive.py(防文件漂移)
  - TestXlsxCorrectness:T1 公式列數/空列留白、C2 回撤色階方向
pytest:python -m pytest quant-service/ecommerce/tests -q
unittest:python quant-service/ecommerce/tests/test_product_factory_v2.py
"""
from __future__ import annotations

import inspect
import re
import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
ECOM = HERE.parent
for p in (str(ECOM), str(ECOM.parent.parent / "youtube_channel" / "scripts"), str(ECOM.parent.parent)):
    if p not in sys.path:
        sys.path.insert(0, p)

import product_factory_v2 as pf  # noqa: E402
import render_kit as RK  # noqa: E402
from product_factory import Provenance  # noqa: E402


def _ctx():
    return pf.load_ctx()


class TestProvenanceGate(unittest.TestCase):
    def test_gate_rejects_fabricated_number(self):
        """憑空塞一個沒入池的績效數字 → 守門必須擋下(fail-closed)。"""
        prov = Provenance()
        pf._pool(prov, 25.1)  # 只有 25.1 是有來源的
        item = {"sku": "TEST", "prov": prov,
                "gate_units": ['<div class="unit">勝率 <b>87.3%</b>,年化 <b>25.1%</b></div>']}
        bad = pf.gate_sku(item)
        vals = {round(c["value"], 1) for c in bad}
        self.assertIn(87.3, vals)            # 未入池的 87.3 被擋
        self.assertNotIn(25.1, vals)         # 有來源的 25.1 放行

    def test_gate_passes_when_all_pooled(self):
        prov = Provenance()
        pf._pool(prov, 25.1, 46.5)
        item = {"sku": "T", "prov": prov,
                "gate_units": ['<div>年化 25.1%,回撤 46.5%</div>']}
        self.assertEqual(pf.gate_sku(item), [])

    def test_tokenizer_reconcile_forgives_only_real_slices(self):
        """FG 對 5 位數% 會切子數;對齊只補真來源大數字的子數,不補憑空數字。"""
        prov = Provenance()
        pf._pool(prov, 10027.7)              # 真來源大數字
        pf._reconcile_tokenizer(prov)
        self.assertTrue(any(abs(p - 27.7) < 0.2 for p in prov.pool))   # 27.7 是 10027.7 的切片 → 補上
        self.assertFalse(any(abs(p - 88.8) < 0.2 for p in prov.pool))  # 88.8 非任何來源切片 → 不補


class TestSKUAssembly(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ctx = _ctx()

    def test_all_builders_pass_gate(self):
        """六個 builder 產出的每個 SKU 變體都要通過誠信守門(真實資料)。"""
        for sid, fn in pf.BUILDERS.items():
            items = fn(self.ctx)
            self.assertTrue(items, f"{sid} 應有產出")
            for it in items:
                bad = pf.gate_sku(it)
                self.assertEqual(bad, [], f"{it['sku']}_{it['lang']} 不該被守門擋:{bad[:2]}")

    def test_item_shape_complete(self):
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                for k in ("sku", "lang", "platform", "prov", "gate_units", "listing", "artifacts"):
                    self.assertIn(k, it, f"{sid} 缺 {k}")
                self.assertTrue(it["artifacts"], f"{sid} 無成品檔")
                # 每個成品 kind 合法
                for fnpath, kind, payload in it["artifacts"]:
                    self.assertIn(kind, ("pdf", "xlsx", "csv"))

    def test_listing_fields_complete(self):
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                lst = it["listing"]
                for k in ("sku", "lang", "platform", "title", "description", "tags", "price"):
                    self.assertIn(k, lst)
                self.assertTrue(lst["title"] and lst["description"] and lst["tags"])
                self.assertLessEqual(len(lst["title"]), 140)
                self.assertLessEqual(len(lst["tags"]), 13)
                # 定價鍵對應語系
                self.assertIn("NTD" if it["lang"] == "zh" else "USD", lst["price"])

    def test_listing_no_hype_words(self):
        ban = ("穩賺", "保證", "必賺", "翻倍", "guaranteed")
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                for tag in it["listing"]["tags"]:
                    self.assertFalse(any(b in tag for b in ban), f"{sid} tag 含誇大詞:{tag}")

    def test_c1_langs_and_c2_langs(self):
        """C1/C2 出中英雙版;T1/T2/M1/M2 只出中版(對齊商業規格 §3 收斂)。"""
        langs = {sid: sorted({it["lang"] for it in fn(self.ctx)}) for sid, fn in pf.BUILDERS.items()}
        self.assertEqual(langs["C1"], ["en", "zh"])
        self.assertEqual(langs["C2"], ["en", "zh"])
        self.assertEqual(langs["T1"], ["zh"])
        self.assertEqual(langs["M1"], ["zh"])


class TestCorrectness(unittest.TestCase):
    def test_dividend_yield_not_x100(self):
        """殖利率欄位已是百分比,體檢卡不得再×100(0.91 應顯示 ~0.9 而非 91)。"""
        ctx = _ctx()
        prov = Provenance()
        ff = pf.facts_for(ctx["checkup"], "2330")
        html = pf.checkup_card(prov, "2330", "台積電", ff)
        self.assertIn("殖利率", html)
        # 91.0 不該出現在殖利率 KPI(0.91 已是 %)
        self.assertNotIn('>91.0<small>%', html)

    def test_checkup_card_pools_all_shown_numbers(self):
        """體檢卡渲染的每個 %數字都要能被守門找到來源(池已含)。"""
        ctx = _ctx()
        prov = Provenance()
        ff = pf.facts_for(ctx["checkup"], "2330")
        html = pf.checkup_card(prov, "2330", "台積電", ff)
        pf._reconcile_tokenizer(prov)
        bad = prov.gate(pf.RK.gate_text([html]), "2330")
        self.assertEqual(bad, [], f"漏綁來源:{[(c['value'], c['clause'][:30]) for c in bad][:3]}")


class TestBlockerRegressions(unittest.TestCase):
    """釘死 VERIFY_REPORT_phase3b 的 4 個 BLOCKER,任一復發即紅燈。"""

    @classmethod
    def setUpClass(cls):
        cls.ctx = _ctx()

    # ── B1:A4 紙張(不渲染,直接釘參數與 CSS 宣告)────────────────────────────
    def test_b1_pdf_paper_is_a4_not_letter(self):
        """`.page` 是 A4 幾何,若沒宣告 @page 就會退回 Letter → 每頁溢出一張空白頁。
        兩道保險(html_doc 的 @page + render_pdf 的 format)必須都在。"""
        doc = RK.html_doc(pf.CFG.BRAND, "p", "s", "t", "<section></section>", ["<div></div>"])
        self.assertRegex(doc, r"@page\s*\{\s*size:\s*A4", "html_doc 少了 @page size:A4 → PDF 會退回 Letter")
        src = inspect.getsource(RK.render_pdf)
        self.assertIn('format="A4"', src, "render_pdf 少了 format=A4")

    # ── B2:維度數不得硬編 ────────────────────────────────────────────────
    def test_b2_dimension_count_is_dynamic_and_true(self):
        """封面/listing 的維度數必須等於該檔實際 facts 數(硬編 11 對 3/9 檔不實)。"""
        ck = self.ctx["checkup"]
        for code in ck.get("by_code", {}):
            ff = pf.facts_for(ck, code)
            if ff:
                self.assertEqual(pf.n_dims(ff), len(ff))
        # T2 出貨哪一檔就宣稱哪一檔的維度數
        for it in pf.build_T2(self.ctx):
            code = it["listing"]["title"]
            nd = pf.n_dims(pf.facts_for(ck, re.search(r"\((\w+)\)", code).group(1)))
            self.assertIn(f"{nd} 項", it["listing"]["description"])

    def test_b2_no_hardcoded_11_in_covers(self):
        """封面 stats 不得再出現寫死的『11項』字面量(來源:A3(b)/(c)/(d))。"""
        src = inspect.getsource(pf)
        self.assertNotIn("'11<small>項</small>'", src)
        self.assertNotIn('"11<small>項</small>"', src)
        self.assertNotIn("(\"Facts each\", '11')", src)

    def test_b2_c2_discloses_thin_tickers(self):
        """C2 合輯各檔維度參差 → 必須揭露範圍與較薄的檔,不得挑一個數字充場面。"""
        ck = self.ctx["checkup"]
        dims = {c: pf.n_dims(pf.facts_for(ck, c)) for c in ck.get("by_code", {}) if pf.facts_for(ck, c)}
        lo, hi = min(dims.values()), max(dims.values())
        if lo == hi:
            self.skipTest("本次資料各檔維度一致,無揭露需求")
        for it in pf.build_C2(self.ctx):
            self.assertIn(str(lo), it["listing"]["description"])
            self.assertIn(str(hi), it["listing"]["description"])

    # ── B3:英文守門不得失明 ──────────────────────────────────────────────
    def test_b3_fg_is_blind_to_pure_english(self):
        """釘住『為什麼需要英文抽取器』這個前提:FG 對純英文回 0 claim。
        哪天 FG 自己補了英文(此測試轉紅),就該回頭簡化 _extract_claims_en 而不是留兩套。"""
        import fact_source_guard as FG
        self.assertEqual(FG.extract_claims("Win rate 99.9%, max drawdown 88.8%."), [])

    def test_b3_en_gate_catches_fabricated_number(self):
        """英文捏造的績效數字必須被擋(否則英文化=順手關掉守門)。"""
        prov = Provenance()
        pf._pool(prov, 25.1)
        item = {"sku": "T", "prov": prov,
                "gate_units": ['<div class="unit">CAGR 25.1% but win rate 99.9%.</div>']}
        vals = {round(c["value"], 1) for c in pf.gate_sku(item)}
        self.assertIn(99.9, vals, "英文捏造數字沒被擋 → 英文版守門形同虛設")
        self.assertNotIn(25.1, vals, "有來源的英文數字不該被誤擋")

    def test_b3_en_units_have_no_chinese_framework_strings(self):
        """英文版內容 unit 不得含中文框架字串(個股名/品牌名等專有名詞除外)。"""
        ck = self.ctx["checkup"]
        code = next(iter(ck["by_code"]))
        names = {v.get("name", "") for v in ck["by_code"].values()}
        html = pf.checkup_card(Provenance(), code, ck["by_code"][code].get("name", code),
                               pf.facts_for(ck, code), "en")
        text = RK.gate_text([html])
        for nm in sorted(names, key=len, reverse=True):
            if nm:
                text = text.replace(nm, "")
        self.assertFalse(re.findall(r"[一-鿿]", text),
                         f"英文卡殘留中文:{set(re.findall(r'[一-鿿]+', text))}")

    def test_b3_en_card_uses_no_fullwidth_space(self):
        """全形空白(U+3000)是中文排版字元,不該出現在英文成品裡。"""
        ck = self.ctx["checkup"]
        code = next(iter(ck["by_code"]))
        html = pf.checkup_card(Provenance(), code, "x", pf.facts_for(ck, code), "en")
        self.assertNotIn("　", html)

    def test_b3_every_listing_has_real_disclaimer(self):
        """中英版都必須真的有免責,且 seo_note 不得自稱含免責卻其實沒有(存證欄位不得說謊)。"""
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                lst = it["listing"]
                self.assertTrue(lst["disclaimer_present"],
                                f'{sid}_{lst["lang"]} listing 無免責')
                self.assertEqual(lst["disclaimer_present"], "已含免責" in lst["seo_note"],
                                 f'{sid}_{lst["lang"]} seo_note 與事實不符')

    # ── B4:存證不得空殼 ─────────────────────────────────────────────────
    def test_b4_provenance_records_not_empty(self):
        """8/8 SKU 的 records 必須非空(v1 baseline 有 15 筆,不可退化為 0)。"""
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                self.assertTrue(it["prov"].records,
                                f'{sid}_{it["lang"]} records=[] → 存證空殼化(規格違反)')

    def test_b4_records_have_resolvable_source_and_value(self):
        """每筆存證都要能回答『哪個檔的哪個欄位』;數值型要存實際 value 不得全 None。"""
        for sid, fn in pf.BUILDERS.items():
            for it in fn(self.ctx):
                recs = it["prov"].records
                for r in recs[:50]:
                    self.assertTrue(r.get("source"), f"{sid} 有 record 沒來源")
                    self.assertTrue(r.get("field"), f"{sid} 有 record 沒欄位名")
                    self.assertTrue(r.get("value") is not None or r.get("text") is not None)
                self.assertTrue(any(r["value"] is not None for r in recs),
                                f"{sid} 全部 record 都沒有數值 → 稽核無法程式化比對")

    def test_b4_checkup_source_points_at_real_fact_key(self):
        """存證的 source 必須指到事實庫裡真的存在的 fact_key + 欄位(不是漂亮的假路徑)。"""
        ck = self.ctx["checkup"]
        prov = Provenance()
        ff = pf.facts_for(ck, "2330")
        pf.checkup_card(prov, "2330", "台積電", ff)
        checked = 0
        for r in prov.records:
            m = re.match(r"youtube_channel/STUDIO/stock_checkup_facts\.json:results\.([^.]+)\.data\.(.+)$",
                         r["source"])
            if not m or r["value"] is None or "[" in m.group(2):
                continue
            node = ck["results"].get(m.group(1), {}).get("data")
            self.assertIsNotNone(node, f'source 指到不存在的 fact_key:{m.group(1)}')
            for part in m.group(2).split("."):
                node = node.get(part) if isinstance(node, dict) else None
            if not isinstance(node, (int, float)):
                continue
            want = node * 100 if "×100" in r["field"] else node
            self.assertLessEqual(abs(want - r["value"]), max(0.25, abs(want) * 0.005),
                                 f'{r["field"]} 存證值 {r["value"]} 對不上來源 {want}')
            checked += 1
        self.assertGreater(checked, 8, "可機械回查的存證太少 → 稽核價值不足")


class TestM1Freshness(unittest.TestCase):
    """M1 賣點是盤前防呆,清單過期**有實害** → 新鮮度措辭必須跟著資料走。"""

    def test_iso_timestamp_updated_still_yields_stale_warning(self):
        """回歸:`updated` 是完整 ISO 時戳(2026-07-13T09:05:36),不是純日期。
        直接餵 date.fromisoformat() 會 ValueError 被吞掉 → 過期警語靜默消失(修了等於沒修)。"""
        from datetime import date, timedelta
        old = (date.today() - timedelta(days=3)).isoformat() + "T09:05:36"
        ctx = dict(_ctx())
        ctx["daytrade"] = {"disposition": ["2330"], "attention": [],
                           "updated": old, "_file": "daytrade_eligibility_test.json"}
        it = pf.build_M1(ctx)[0]
        body = "".join(it["gate_units"])
        self.assertIn("距今 3 天", body, "ISO 時戳沒被正確解析 → 過期警語沒出現")
        self.assertNotIn("09:05:36", body, "機器時戳不該出現在買家看到的文案裡")
        self.assertNotIn("09:05:36", it["listing"]["description"])

    def test_unparsable_date_fails_safe_to_warning(self):
        """解析不出日期時必須示警(fail-safe),不可默認為新鮮。"""
        ctx = dict(_ctx())
        ctx["daytrade"] = {"disposition": ["2330"], "attention": [],
                           "updated": "(unknown)", "_file": "x.json"}
        self.assertIn("以證交所當日公告為準", "".join(pf.build_M1(ctx)[0]["gate_units"]))

    def test_listing_does_not_claim_today_or_daily_update(self):
        """舊 listing 寫「每日更新/附今日快照」但抓取器已停 3 天 → 不得再宣稱新鮮度。"""
        for it in pf.build_M1(_ctx()):
            lst = it["listing"]
            self.assertNotIn("每日更新", lst["title"])
            self.assertNotIn("今日快照", lst["description"])


class TestMethodDisclosure(unittest.TestCase):
    """C1 的方法揭露是對外承諾,必須永遠等於源碼真值——參數漂移時這裡要先紅。"""

    def test_cost_numbers_match_strategy_source(self):
        import strategy
        cm = strategy.COST_MODELS["tw_real"]
        zh = " ".join(b for _, b in pf.METHOD_ZH)
        en = " ".join(b for _, b in pf.METHOD_EN)
        for txt in (zh, en):
            self.assertIn(f'{cm["fee_buy"] * 100:g}%', txt)      # 0.1425%
            self.assertIn(f'{cm["fee_sell"] * 100:g}%', txt)
            self.assertIn(f'{cm["tax_sell"] * 100:g}%', txt)     # 0.3%
        # slip_ticks=0 → 必須明講「未模擬滑價」,不得含糊
        self.assertEqual(cm["slip_ticks"], 0)
        self.assertIn("未模擬滑價", zh)
        self.assertIn("Slippage is NOT modelled", en)

    def test_regime_params_match_tw_adaptive_source(self):
        import tw_adaptive
        p = tw_adaptive.AdaptiveParams
        zh = " ".join(b for _, b in pf.METHOD_ZH)
        en = " ".join(b for _, b in pf.METHOD_EN)
        for txt in (zh, en):
            self.assertIn(f"{p.adxTrend:g}", txt)        # 18
            self.assertIn(f"{p.erTrend:g}", txt)         # 0.26
            self.assertIn(f"{p.bbLen:g}", txt)           # 30
            self.assertIn(f"{p.bbK:g}", txt)             # 2.5
            self.assertIn(f"{p.rsiBuy:g}", txt)          # 30
            self.assertIn(f"{p.rsiSell:g}", txt)         # 50
            self.assertIn(f"{p.wideMult:g}", txt)        # 11
            self.assertIn(f"{p.trendExposure:g}", txt)   # 0.95
            self.assertIn(f"{p.mrExposure:g}", txt)      # 0.6
            self.assertIn(f"{int(p.barsPerYear)}", txt)  # 252
        self.assertIn(f"{tw_adaptive.MIN_YEARS:g}", zh)
        self.assertIn(f"{tw_adaptive.MIN_BARS:g}", zh)

    def test_survivorship_bias_is_disclosed(self):
        """get_universe() 取現存上市櫃 → 倖存者偏誤為真,且對我們不利,必須照實揭露。"""
        zh = " ".join(b for _, b in pf.METHOD_ZH)
        en = " ".join(b for _, b in pf.METHOD_EN)
        self.assertIn("下市", zh)
        self.assertIn("高估", zh)
        self.assertIn("delisted", en)
        self.assertIn("inflates", en)


class TestXlsxCorrectness(unittest.TestCase):
    """xlsx 出貨物的正確性(A6 平均成本 92.5 / A7 色階反向)。"""

    @classmethod
    def setUpClass(cls):
        cls.ctx = _ctx()
        cls.tmp = Path(__file__).resolve().parent / "_tmp_xlsx"
        cls.tmp.mkdir(exist_ok=True)

    @classmethod
    def tearDownClass(cls):
        import shutil
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_t1_formulas_cover_all_rows_and_blank_rows_stay_blank(self):
        """A6:舊版只有第 2/3 列有公式,且 D3 有股數但 E3 無價 → 平均成本算出 92.5。
        新版:公式鋪滿資料列,且沒填股數的列用 IF 留白(不會蹦出 0 或荒謬均價)。"""
        import openpyxl
        refs = [("0050", "元大台灣50", 10.0, 5.0, 7.0, "src")]
        p = self.tmp / "t1.xlsx"
        pf._build_dca_xlsx(p, refs)
        ws = openpyxl.load_workbook(p)["定投追蹤"]
        frows = [r for r in range(2, 60)
                 if isinstance(ws.cell(row=r, column=9).value, str)
                 and str(ws.cell(row=r, column=9).value).startswith("=")]
        self.assertGreaterEqual(len(frows), 20, "公式列太少 → 買家從第 4 列登就不會自動算")
        self.assertEqual(frows, list(range(2, max(frows) + 1)), "公式列不連續")
        for r in frows:                       # 每列都要有空值防護,否則空列會算出假均價
            self.assertIn('=""', str(ws.cell(row=r, column=9).value).replace('=IF(OR(', ''))
            self.assertTrue(str(ws.cell(row=r, column=9).value).startswith("=IF(OR("))
        # 示範列只有一列有資料;第 3 列不得同時「有股數、無價格」(那正是 92.5 的成因)
        self.assertIsNotNone(ws["E2"].value)
        self.assertIsNone(ws["D3"].value)

    def test_t1_note_is_below_data_rows(self):
        """MINOR 14:說明文字不得放在資料列裡(買家往下登會撞到)。"""
        import openpyxl
        refs = [("0050", "元大台灣50", 10.0, 5.0, 7.0, "src")]
        p = self.tmp / "t1b.xlsx"
        pf._build_dca_xlsx(p, refs)
        ws = openpyxl.load_workbook(p)["定投追蹤"]
        last_formula = max(r for r in range(2, 60)
                           if isinstance(ws.cell(row=r, column=9).value, str)
                           and str(ws.cell(row=r, column=9).value).startswith("="))
        for r in range(2, last_formula + 1):
            v = ws.cell(row=r, column=1).value
            self.assertFalse(isinstance(v, str) and len(v) > 20,
                             f"第 {r} 列(資料區)混進說明文字:{v!r}")

    def test_c2_drawdown_colorscale_matches_cagr_direction(self):
        """A7:回撤欄是**負值**,與 C1 的正值幅度符號相反。本表自訂『紅=好』,
        故最慘(min,-98.5)必須是綠、最輕(max)是紅——與年化欄同方向。"""
        import openpyxl
        ck = self.ctx["checkup"]
        codes = [c for c in ck.get("by_code", {}) if pf.facts_for(ck, c)]
        p = self.tmp / "c2.xlsx"
        pf._build_checkup_overview_xlsx(p, ck, codes)
        ws = openpyxl.load_workbook(p)["總覽"]
        got = {}
        for rng, rules in ws.conditional_formatting._cf_rules.items():
            for r in rules:
                if r.type == "colorScale":
                    col = str(rng.sqref)[0]
                    got[col] = [c.rgb[-6:] for c in r.colorScale.color]
        self.assertIn("C", got); self.assertIn("D", got)
        self.assertEqual(got["C"], got["D"],
                         "回撤欄(負值)色階方向必須與年化欄一致,否則同表紅色一下代表好一下代表糟")
        # 回撤欄的值確實是負的(若來源改成正值幅度,本測試要跟著改方向)
        vals = [ws.cell(row=r, column=4).value for r in range(2, len(codes) + 2)]
        self.assertTrue(all(v is None or v <= 0 for v in vals), "回撤欄符號變了 → 色階方向要重新檢討")

    def test_c2_overview_discloses_per_ticker_dimensions(self):
        import openpyxl
        ck = self.ctx["checkup"]
        codes = [c for c in ck.get("by_code", {}) if pf.facts_for(ck, c)]
        p = self.tmp / "c2b.xlsx"
        pf._build_checkup_overview_xlsx(p, ck, codes)
        ws = openpyxl.load_workbook(p)["總覽"]
        self.assertIn("體檢維度", [c.value for c in ws[1]])
        for i, c in enumerate(codes, start=2):
            self.assertEqual(ws.cell(row=i, column=8).value, pf.n_dims(pf.facts_for(ck, c)))

    def test_en_xlsx_headers_and_sheetnames_have_no_chinese(self):
        """英文買家收到中文表頭/分頁名 = 看不懂自己買了什麼(B3)。"""
        import openpyxl
        ck = self.ctx["checkup"]
        codes = [c for c in ck.get("by_code", {}) if pf.facts_for(ck, c)]
        p = self.tmp / "c2en.xlsx"
        pf._build_checkup_overview_xlsx(p, ck, codes, None, "en")
        wb = openpyxl.load_workbook(p)
        for ws in wb.worksheets:
            self.assertFalse(re.findall(r"[一-鿿]", ws.title), f"分頁名有中文:{ws.title}")
            for c in ws[1]:
                if c.value:
                    self.assertFalse(re.findall(r"[一-鿿]", str(c.value)),
                                     f"表頭有中文:{c.value}")


if __name__ == "__main__":
    unittest.main()
