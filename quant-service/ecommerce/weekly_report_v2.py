#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""weekly_report_v2.py —【旗艦】台股全市場週報引擎 v2。

把 v1 只吐純文字 .md 的 subscription_report 升級成品牌化深色 PDF 週報(+ xlsx 附件)。
主體吃 data_hunter/state.json 全市場掃描(1001 檔 wave_top + 34 板塊 + 強弱 + 法人籌碼 +
真實訊號追蹤),體檢降為深度加值層。規格:
  - 商業篇 docs/ecommerce/REDESIGN_SPEC_business.md §2(8 sections / 資料源 / 誠信)
  - 產品篇 docs/ecommerce/REDESIGN_SPEC_product.md(視覺 token / 分頁 / 渲染管線)

誠信鐵則(生死線,與產線同一條):
  - 寫進 PDF 的每個績效數字都經 Provenance 綁來源欄位並入池(復用 product_factory.Provenance,
    不自造弱化版);組完各 section 後用嚴容差 gate 複驗,查無來源 → 該 section 降級為
    「資料異常已隱藏」卡(fail-closed),絕不炸整份報告、絕不編數字。
  - 任一資料源缺/空(休市/circuit_breaker/signals=[]) → 該 section 降級為明確「本週無資料」卡。
  - 「介紹≠推薦」:估值用位階條不用買賣燈;免責頁必印。

交付:load_send_list() 讀 STUDIO/ecommerce_subscribers.json(不存在回空+log,不硬依賴);
      send 一律 dry_run 預設,本引擎不寄信、不打外部 API。

用法:
  python weekly_report_v2.py                 # 用當前真實資料產一份完整週報 PDF+xlsx
  python weekly_report_v2.py --tier basic    # 只出基礎版 section
  python weekly_report_v2.py --out <dir>
驗證:python -m py_compile quant-service/ecommerce/weekly_report_v2.py
"""
from __future__ import annotations

import argparse
import csv
import glob
import json
import math
import re
import statistics
import sys
from datetime import date, datetime, timedelta
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

# ── 路徑 ─────────────────────────────────────────────────────────────────────
HERE = Path(__file__).resolve().parent            # quant-service/ecommerce
QUANT = HERE.parent                                # quant-service
ROOT = QUANT.parent                                # repo root
DATA_HUNTER = QUANT / "data_hunter"
TWDATA = ROOT / "twdata"
STUDIO = ROOT / "youtube_channel" / "STUDIO"
SCRIPTS = ROOT / "youtube_channel" / "scripts"
MOCKUP = HERE / "mockup"

STATE = DATA_HUNTER / "state.json"
CHECKUP_FACTS = STUDIO / "stock_checkup_facts.json"
ADAPTIVE_CSV = TWDATA / "adaptive_per_stock.csv"
SUBSCRIBERS = STUDIO / "ecommerce_subscribers.json"
THEME_CSS = HERE / "report_theme.css"
OUT_DEFAULT = QUANT / "output" / "ecommerce_ready" / "weekly_v2"

# ── 復用產線既有溯源守門 + Provenance(不自造弱化版)──────────────────────────
sys.path.insert(0, str(SCRIPTS))
try:
    import fact_source_guard as FG  # noqa: E402
except Exception as exc:  # noqa: BLE001
    print(f"[weekly_v2] 致命:無法載入 fact_source_guard(溯源守門)——{exc}")
    raise SystemExit(2)
try:
    import config as CFG  # ecommerce/config.py（同目錄，sys.path 已含 HERE? 保險加入）
except Exception:  # noqa: BLE001
    sys.path.insert(0, str(HERE))
    import config as CFG  # noqa: E402

# Provenance 定義在 product_factory —— 直接復用同一個類別
sys.path.insert(0, str(HERE))
from product_factory import Provenance  # noqa: E402

TODAY = date.today()
DISCLAIMER = ("本內容為程式化的歷史數據彙整與教學,只做「事實介紹」,不是投資建議、不喊單、"
              "不報明牌;歷史數據非未來保證,投資有風險,據此進出盈虧自負。「介紹」不等於「推薦」。")


# ══════════════════════════════════════════════════════════════════════════════
#  小工具
# ══════════════════════════════════════════════════════════════════════════════
def _esc(s) -> str:
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


_NUM_RE = re.compile(r"-?\d+\.?\d*")


def _pool_fg(prov: Provenance, s: str) -> None:
    """把 FG 百分比抽取器『會從 s 抽到的形式』也入池。

    FG._RX_PCT_ARABIC 的整數位只吃 \\d{1,4},對 ≥10000% 的台股怪物報酬(如 21051.3%)
    會截成 1051.3 這種幻影數 —— 那是 tokenizer 侷限,不是造假。用 FG 自己的抽取器
    對來源值/來源字串預抽並入池,讓 gate 對『合法大數』不誤殺,同時仍擋得下『憑空手打、
    不在任何來源字串也非任何來源欄位』的績效數(那種數字不會被 _pool 收進池)。
    """
    for m in FG._RX_PCT_ARABIC.finditer(s or ""):
        tok = m.group(1) or m.group(2)
        try:
            v = float(tok)
            prov.pool.add(abs(v)); prov.pool.add(abs(round(v, 1)))
        except (TypeError, ValueError):
            pass


def _pool(prov: Provenance, *items) -> None:
    """把「來源提供的」數字入池(數字直接入,字串抽出其中數字入,並補 FG-tokenized 形式)。

    只該對『真的來自來源欄位/來源字串』的值呼叫 —— 這正是 gate 的鑑別點:
    憑空手打、沒經過 _pool 的績效數字不會在池裡 → gate 擋下。
    """
    for it in items:
        if it is None:
            continue
        if isinstance(it, (int, float)):
            v = float(it)
            prov.pool.add(abs(v)); prov.pool.add(abs(round(v, 1)))
            _pool_fg(prov, f"{v}%"); _pool_fg(prov, f"{v:.2f}%")
        else:
            s = str(it)
            for m in _NUM_RE.findall(s):
                try:
                    v = float(m)
                    prov.pool.add(abs(v)); prov.pool.add(abs(round(v, 1)))
                except ValueError:
                    pass
            _pool_fg(prov, s)


def _bind(prov: Provenance, src: str, **fields) -> None:
    """具名綁定:入池(給 gate 用)+ 留存證(給稽核檔用),每個數字綁「來源檔:欄位」。

    A1 修(VERIFY_REPORT_phase3a):v1 只有 S7 寫 prov.records,其餘 7 段雖然有入池
    (所以 gate 全段有效、憑空造假擋得下),但**持久化的稽核檔只覆蓋 1/8 段** ——
    外部稽核者無法只憑該檔回查 S1–S6/S8。現在所有 section 一律走 _bind,
    存證檔可逐筆回答「這個數字來自哪個檔的哪個欄位」。

    A3 修:數值型直接存 value(不再全 null),稽核可程式化比對而不只靠文字。
    """
    for name, v in fields.items():
        if v is None:
            continue
        _pool(prov, v)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            prov.records.append({"field": name, "value": float(v), "text": None, "source": src})
        else:
            prov.records.append({"field": name, "value": None, "text": str(v), "source": src})


def wilson_ci(k: int, n: int, z: float = 1.96) -> tuple[float, float] | None:
    """勝率的 Wilson score 95% 信賴區間(回 0~1 的 (lo, hi));n<=0 回 None。

    為什麼要它:REVIEW 要求「標明 n=19 樣本不足、不可外推」。但**光寫「樣本不足」是空話**——
    讀者無從判斷有多不足。Wilson 區間把它變成可查證的事實:19 筆、勝率 15.8% 的 95% 區間是
    5.5%~37.6%,**寬到什麼都不能斷言**。這是標準閉式解(比 normal approximation 在小樣本/
    極端比例下更正確),不是自創統計。
    驗證:k=50,n=100 → 0.4038~0.5962(對得上教科書值);見 tests。
    """
    if n <= 0:
        return None
    p = k / n
    d = 1 + z * z / n
    c = p + z * z / (2 * n)
    m = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return ((c - m) / d, (c + m) / d)


def _pct(v, digits=2, sign=True) -> tuple[str, str]:
    """回傳 (顯示字串, 台股色 class)。正=紅(pos) 負=綠(neg)。"""
    try:
        v = float(v)
    except (TypeError, ValueError):
        return "—", ""
    s = f"{v:+.{digits}f}%" if sign else f"{v:.{digits}f}%"
    return s, ("pos" if v > 0 else "neg" if v < 0 else "")


def _gate_text_from_units(units: list[str]) -> str:
    joined = "\n".join(units)
    return re.sub(r"<[^>]+>", " ", joined)


def _degrade_unit(sid: str, title: str, reason: str) -> str:
    return (f'<div class="unit"><div class="sec-hd"><span class="sn">{sid}</span>'
            f'<h2>{_esc(title)}</h2></div>'
            f'<div class="card degrade"><div class="dt">本週此段無可用資料</div>'
            f'<div class="db">{_esc(reason)}</div></div></div>')


# ══════════════════════════════════════════════════════════════════════════════
#  資料載入(全部 fail-safe:缺檔回空 dict/list,不炸)
# ══════════════════════════════════════════════════════════════════════════════
def load_state() -> dict:
    try:
        return json.loads(STATE.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {}


def load_valuation_latest() -> dict:
    files = sorted(glob.glob(str(TWDATA / "fundamentals" / "valuation_*.json")))
    if not files:
        return {}
    try:
        d = json.loads(Path(files[-1]).read_text(encoding="utf-8"))
        return {"date": d.get("date"), "data": d.get("data", {}), "_file": Path(files[-1]).name}
    except Exception:  # noqa: BLE001
        return {}


def load_chips_week(days: int = 5) -> list[dict]:
    files = sorted(glob.glob(str(TWDATA / "chips" / "*.json")))[-days:]
    out = []
    for f in files:
        try:
            out.append(json.loads(Path(f).read_text(encoding="utf-8")))
        except Exception:  # noqa: BLE001
            continue
    return out


def load_checkup() -> dict:
    try:
        return json.loads(CHECKUP_FACTS.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {"results": {}, "by_code": {}}


def load_adaptive() -> list[dict]:
    rows = []
    if not ADAPTIVE_CSV.exists():
        return rows
    with open(ADAPTIVE_CSV, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            try:
                r["a_net"] = float(r["a_net"]); r["a_win"] = float(r["a_win"])
            except Exception:  # noqa: BLE001
                continue
            rows.append(r)
    return rows


def _load_csv_floats(path: Path, floats: tuple) -> list[dict]:
    """通用 CSV 載入(指定欄位轉 float,轉不動的列丟掉)。給 S8 主題輪替的另外兩份基準用。"""
    rows: list[dict] = []
    if not path.exists():
        return rows
    try:
        with open(path, encoding="utf-8-sig") as f:
            for r in csv.DictReader(f):
                ok = True
                for k in floats:
                    try:
                        r[k] = float(r[k])
                    except (KeyError, ValueError, TypeError):
                        ok = False
                if ok:
                    rows.append(r)
    except Exception:  # noqa: BLE001
        return []
    return rows


def _snapshot_age(path: Path) -> tuple[str, int]:
    """回 (快照日 YYYY-MM-DD, 距今天數)。檔案不存在回 ("—", -1)。"""
    try:
        d = datetime.fromtimestamp(path.stat().st_mtime).date()
        return d.isoformat(), (TODAY - d).days
    except Exception:  # noqa: BLE001
        return "—", -1


def _twstock_meta() -> dict:
    """twstock 內建代號表 → {code: (name, industry)}。**本地套件資料集,不打網路**。

    這是全市場唯一「代號→名稱/產業」都齊全的權威來源(實測 12 個缺名碼 12/12 命中)。
    載入失敗就回空 dict —— name map 還有 state/adaptive 兩層來源,不因此炸掉。
    """
    try:
        import twstock
    except Exception:  # noqa: BLE001
        return {}
    out = {}
    try:
        for code, info in twstock.codes.items():
            if getattr(info, "type", "") == "股票" and getattr(info, "name", ""):
                out[str(code)] = (info.name, getattr(info, "group", "") or "")
    except Exception:  # noqa: BLE001
        return {}
    return out


def build_name_map(state: dict, checkup: dict, adaptive: list[dict] | None = None) -> dict:
    """{code: name}。**查不到就不放進 map**(不再拿代號當名字)。

    🔴 REVIEW_weekly_value 續訂殺手 #5:舊版每個來源都寫 `r.get("name", code)`、
    呼叫端又寫 `nmap.get(code, code)` —— 兩層「查不到就用代號」疊起來,渲染成
    `{name}<span class=code>{code}</span>` 就變成「**2891 2891**」「3045 3045」,實測 12 檔中鏢
    (中信金/台灣大/亞泥都是大型股,散戶一眼看出「連中信金的名字都不會寫」)。
    根因:那 12 檔只出現在 valuation/chips(它們**沒有 name 欄位**),不在 state.wave_top 的
    掃描宇宙裡,所以三個來源都查不到。
    修法:①補 adaptive_per_stock.csv(1770 檔帶 name)與 twstock 內建表(全市場,權威)
    ②map 不再自我填充代號 —— 查不到就是查不到,由 _stock_cell() 決定只印一次代號。
    """
    m: dict[str, str] = {}

    def _put(code, name, overwrite=False):
        code = str(code or "").strip()
        name = str(name or "").strip()
        # 名稱等於代號(來源自己就沒名字)視同沒有,不要讓它佔位
        if not code or not name or name == code:
            return
        if overwrite or code not in m:
            m[code] = name

    for r in state.get("wave_top", []) or []:
        _put(r.get("code"), r.get("name"), overwrite=True)
    for key in ("strong", "weak", "watch_long"):
        for r in state.get(key, []) or []:
            _put(r.get("code"), r.get("name"))
    for side in ("long", "short"):
        for r in (state.get("signals", {}) or {}).get(side, []) or []:
            _put(r.get("code"), r.get("name"))
    for grp in ("foreign_top", "trust_top", "consec_top", "retail_exit_top"):
        for r in (state.get("chips", {}) or {}).get(grp, []) or []:
            _put(r.get("code"), r.get("name"))
    for code, meta in (checkup.get("by_code", {}) or {}).items():
        _put(code, (meta or {}).get("name"))
    for r in adaptive or []:
        _put(r.get("code"), r.get("name"))
    for code, (name, _ind) in _twstock_meta().items():
        _put(code, name)
    return m


def build_industry_map(state: dict) -> dict:
    """{code: industry}。給 S5 做產業集中度揭露用(twstock 的 group = 產業別)。"""
    m: dict[str, str] = {}
    for code, (_n, ind) in _twstock_meta().items():
        if ind:
            m[str(code)] = ind
    for r in state.get("wave_top", []) or []:
        if r.get("code") and r.get("industry"):
            m[str(r["code"])] = r["industry"]
    return m


def _stock_cell(nmap: dict, code, cls: str = "l tkr") -> str:
    """個股欄位 HTML:有名字就「名稱+小字代號」,**查不到就只印一次代號**(不印兩次露餡)。"""
    code = str(code or "")
    name = nmap.get(code)
    if name:
        return f'<td class="{cls}">{_esc(name)}<span class="code">{_esc(code)}</span></td>'
    return f'<td class="{cls}">{_esc(code)}</td>'


def _stock_inline(nmap: dict, code) -> str:
    """同 _stock_cell 但不含 <td>(給 cmp bar 的 label 用)。"""
    code = str(code or "")
    name = nmap.get(code)
    if name:
        return f'{_esc(name)} <span class="code">{_esc(code)}</span>'
    return f'<span class="code">{_esc(code)}</span>'


# ══════════════════════════════════════════════════════════════════════════════
#  S7 點播佇列 —— 唯一護城河:完整版訂戶每月可指定 1 檔深度體檢
# ══════════════════════════════════════════════════════════════════════════════
# 🟠 REVIEW 續訂殺手 #3:S7 是全報告唯一「別處拿不到」的東西(評審 8 分),但它是**樂透**——
# 每週隨機 7 檔、訂戶不能指定,全輪 1925 檔要 5.3 年。付 149 的人想看的是**自己手上那檔**,
# 不是隨機發牌。開放點播 = 把唯一的護城河從樂透變成服務,同時製造 basic→full 的升級動機。
#
# 設計取捨:
#  - 佇列是**檔案**(STUDIO/checkup_requests.json),不是 DB —— 與本專案其餘狀態一致,
#    且 0 訂閱的現在必然是空的 → 空佇列時 S7 行為必須與改版前**完全相同**(有測試釘死)。
#  - 現算走 `ensure_fn` **可注入**:預設打真引擎(會抓行情),測試注入 stub → 測試零外部網路。
#  - 額度/資格在 submit_request() 就擋掉(完整版才有、每月 1 檔),不是產週報時才發現。
CHECKUP_REQUESTS = STUDIO / "checkup_requests.json"
REQ_QUOTA_PER_MONTH = 1      # 每位完整版訂戶每月可點播幾檔
REQ_MAX_PER_ISSUE = 3        # 單期最多處理幾筆(其餘留到下期:保護版面 + 現算成本)
REQ_MAX_ATTEMPTS = 3         # 現算連續失敗幾次才判 failed(暫時性抓取失敗不該直接槍斃)


def load_requests() -> dict:
    """讀點播佇列;不存在/壞檔一律回空結構(不炸、不硬依賴)。"""
    try:
        d = json.loads(CHECKUP_REQUESTS.read_text(encoding="utf-8"))
        if isinstance(d, dict) and isinstance(d.get("requests"), list):
            return d
    except Exception:  # noqa: BLE001
        pass
    return {"updated": TODAY.isoformat(), "requests": []}


def save_requests(doc: dict) -> bool:
    """寫回佇列。寫不進去只 log 不炸(週報本身不該因為佇列寫入失敗而產不出來)。"""
    try:
        doc["updated"] = TODAY.isoformat()
        CHECKUP_REQUESTS.parent.mkdir(parents=True, exist_ok=True)
        CHECKUP_REQUESTS.write_text(json.dumps(doc, ensure_ascii=False, indent=2), encoding="utf-8")
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"[weekly_v2] 點播佇列寫入失敗({exc});本期不更新狀態。")
        return False


_RX_CODE = re.compile(r"^\d{4,6}[A-Z]?$")


def submit_request(email: str, code: str, *, doc: dict | None = None,
                   subscribers: list[dict] | None = None, nmap: dict | None = None,
                   today: date | None = None) -> tuple[bool, str]:
    """訂戶點播一檔。回 (是否受理, 給使用者看的訊息)。**不寄信、不打網路**,只動佇列檔。

    這是給漏斗/webhook 之後接上來的 API;現在沒有 UI,但契約與額度規則先立好並測試。
    資格檢查一律 fail-closed:查不到訂戶 → 不受理(現在 0 訂閱,所以佇列必然是空的,符合預期)。
    """
    d = today or TODAY
    doc = load_requests() if doc is None else doc
    email = (email or "").strip().lower()
    code = (code or "").strip().upper()
    if not email or "@" not in email:
        return False, "請提供有效的 Email。"
    if not _RX_CODE.match(code):
        return False, f"「{code}」不是有效的台股代號格式。"
    # 代號要真的存在(用全市場名冊驗;查不到就不受理,免得排進去才發現算不出來)
    known = nmap if nmap is not None else {c: n for c, (n, _i) in _twstock_meta().items()}
    if known and code not in known:
        return False, f"查無代號 {code};請確認是上市/上櫃股票代號。"
    # 資格:完整版(含年繳)才有點播
    subs = load_send_list("full") if subscribers is None else subscribers
    emails = {(s.get("email") or "").strip().lower() for s in subs}
    if email not in emails:
        return False, "深度體檢點播是「完整版」的功能,你目前的方案沒有這項;升級後即可每月指定 1 檔。"
    # 額度:每人每月 REQ_QUOTA_PER_MONTH 檔(以 requested_at 的年月計)
    ym = d.strftime("%Y-%m")
    used = [r for r in doc["requests"]
            if (r.get("email") or "").lower() == email
            and str(r.get("requested_at", ""))[:7] == ym
            and r.get("status") != "rejected"]
    if len(used) >= REQ_QUOTA_PER_MONTH:
        return False, f"你這個月的點播額度已用完(每月 {REQ_QUOTA_PER_MONTH} 檔),下個月 1 號重置。"
    if any(r.get("code") == code and r.get("status") == "pending" for r in doc["requests"]):
        # 別人已經點過同一檔且還沒出 → 不佔他額度,直接告訴他會出現在本期
        return True, f"{code} 已在本期排程中,你會在最近一期週報看到它(未扣你的額度)。"
    doc["requests"].append({
        "email": email, "code": code, "requested_at": d.isoformat(),
        "status": "pending", "attempts": 0, "fulfilled_in": None, "last_error": None,
    })
    return True, f"已收到 {code} 的點播,會排進最近一期的完整版週報。"


def pending_requests(doc: dict | None, limit: int = REQ_MAX_PER_ISSUE) -> list[dict]:
    """待處理點播(先到先服務);doc 為 None/空 → 回空 list(空佇列 = 行為不變)。"""
    if not doc:
        return []
    out = [r for r in doc.get("requests", []) if r.get("status") == "pending"]
    out.sort(key=lambda r: str(r.get("requested_at", "")))
    return out[:limit]


def _ensure_checkup_via_engine(code: str) -> bool:
    """把不在事實庫的點播檔**現算**出來:呼叫體檢引擎 → 寫回共用事實庫。

    ⚠️ 這條路徑會抓外部行情(FinMind/Yahoo),故:
      - 一切例外都吞成 False(fail-safe):算不出來就降級說明,**絕不炸整份週報、絕不編數字**。
      - 測試一律注入 stub ensure_fn,不打網路(見 tests)。
    """
    try:
        if str(SCRIPTS) not in sys.path:
            sys.path.insert(0, str(SCRIPTS))
        import stock_checkup_facts as SCF
        facts, _ = SCF.build_checkup(code)
        if not facts:
            return False
        SCF.merge_and_write(facts)
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"[weekly_v2] 點播 {code} 現算失敗({type(exc).__name__}: {exc});本期降級說明。")
        return False


# ══════════════════════════════════════════════════════════════════════════════
#  Section 建造器 —— 每個回傳 dict(id,title,tier,units,prov,degraded)
#  units 為 HTML 片段清單(供分頁);第一個 unit 含 sec-hd。
# ══════════════════════════════════════════════════════════════════════════════
def _sec_head(sid: str, title: str, tier: str) -> str:
    tcls = "b" if tier == "basic" else ""
    tlabel = "基礎版" if tier == "basic" else "完整版"
    return (f'<div class="sec-hd"><span class="sn">{sid}</span><h2>{_esc(title)}</h2>'
            f'<span class="tier {tcls}">{tlabel}</span></div>')


def sec_S1(state: dict) -> dict:
    sid, title, tier = "S1", "市場溫度與體質", CFG.SECTION_TIERS["S1"]
    prov = Provenance()
    g = state.get("gauge", {}) or {}
    idx = state.get("index", {}) or {}
    if not g:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "state.json 無 gauge 溫度資料(可能休市或掃描未跑)。")]}
    temp = g.get("temperature")
    _bind(prov, "state.json:gauge+index",
          temperature=temp, breadth=g.get("breadth"), adr=g.get("adr"), nh=g.get("nh"),
          nl=g.get("nl"), avg_rsi=g.get("avg_rsi"), adv=g.get("adv"), dec=g.get("dec"),
          flat=g.get("flat"), index_chg=idx.get("chg"), index_price=idx.get("price"))
    label = g.get("label", "")
    idx_chg, idx_cls = _pct(idx.get("chg"))
    mkpos = max(0.0, min(100.0, float(temp) if temp is not None else 50.0))
    head = _sec_head(sid, title, tier)
    lead = ('<div class="lead">溫度計是全市場體質的綜合讀數(RSI/廣度/漲跌比/新高低/量能加權),'
            '只陳述體質、不判斷方向。</div>')
    kpis = (
        f'<div class="kpis">'
        f'<div class="kpi"><div class="k">市場溫度</div><div class="v">{_esc(f"{temp:.1f}") if temp is not None else "—"}'
        f'<small> {_esc(label)}</small></div></div>'
        f'<div class="kpi"><div class="k">站上20MA</div><div class="v">{_esc(g.get("breadth","—"))}<small>%</small></div></div>'
        f'<div class="kpi"><div class="k">漲跌比 ADR</div><div class="v">{_esc(g.get("adr","—"))}</div></div>'
        f'<div class="kpi"><div class="k">平均 RSI</div><div class="v">{_esc(g.get("avg_rsi","—"))}</div></div>'
        f'</div>')
    gauge = (f'<div class="gaugebar"><span class="mk" style="left:{mkpos:.0f}%"></span></div>'
             f'<div class="gaugescale"><span>0 冷</span><span>50 中性</span><span>100 熱</span></div>')
    detail = (
        f'<div class="block" style="margin-top:11px"><div class="metricrow">'
        f'漲 <b>{_esc(g.get("adv",0))}</b> ／ 跌 <b>{_esc(g.get("dec",0))}</b> ／ 平 <b>{_esc(g.get("flat",0))}</b>'
        f'　·　60日新高 <b>{_esc(g.get("nh","—"))}</b> ／ 新低 <b>{_esc(g.get("nl","—"))}</b><br>'
        f'大盤 {_esc(idx.get("name","0050"))} 收 <b>{_esc(idx.get("price","—"))}</b>,當日 '
        f'<span class="{idx_cls}">{idx_chg}</span>,趨勢 <b>{_esc(idx.get("trend","—"))}</b>'
        f'{"（站上年線）" if idx.get("above_yearline") else ""}</div></div>')
    html = f'<div class="unit">{head}{lead}{kpis}{gauge}{detail}</div>'
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [html]}


def sec_S2(state: dict) -> dict:
    sid, title, tier = "S2", "板塊輪動熱力", CFG.SECTION_TIERS["S2"]
    prov = Provenance()
    secs = state.get("sectors", []) or []
    if not secs:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "state.json 無 sectors 板塊資料。")]}
    ranked = sorted(secs, key=lambda s: s.get("score", 0), reverse=True)
    top, bot = ranked[:5], ranked[-5:]

    def rows(items):
        out = []
        for s in items:
            chg, cls = _pct(s.get("avg_chg"))
            _bind(prov, f"state.json:sectors[{s.get('name','?')}]",
                  avg_chg=s.get("avg_chg"), bull_pct=s.get("bull_pct"), score=s.get("score"),
                  count=s.get("count"), inst_count=s.get("inst_count"), leader=s.get("leader"))
            bull = f"{s.get('bull_pct', 0):.0f}"
            score = f"{s.get('score', 0):.1f}"
            out.append(
                f'<tr><td class="l tkr">{_esc(s.get("name","—"))}</td>'
                f'<td class="{cls}">{chg}</td>'
                f'<td>{bull}%</td>'
                f'<td>{score}</td>'
                f'<td>{_esc(s.get("count","—"))}</td>'
                f'<td>{_esc(s.get("inst_count","—"))}</td>'
                f'<td class="l">{_esc(s.get("leader","—"))}</td></tr>')
        return "".join(out)

    head = _sec_head(sid, title, tier)
    lead = (f'<div class="lead">全 {len(secs)} 板塊依強弱分數排序,列最強 5 / 最弱 5;'
            f'完整 {len(secs)} 板塊附 xlsx 可自行排序篩選。</div>')
    thead = ('<thead><tr><th class="l">板塊</th><th>均漲跌</th><th>多方%</th><th>分數</th>'
             '<th>檔數</th><th>法人買</th><th class="l">領漲</th></tr></thead>')
    u_top = (f'<div class="unit">{head}{lead}<table class="grid">{thead}'
             f'<tbody>{rows(top)}</tbody></table></div>')
    u_bot = (f'<div class="unit"><div class="lead" style="color:var(--tx3)">最弱 5 板塊</div>'
             f'<table class="grid">{thead}<tbody>{rows(bot)}</tbody></table></div>')
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [u_top, u_bot]}


def _strength_table(prov, rows_data, nmap, n, ascending=False, label=""):
    ranked = sorted(rows_data, key=lambda r: r.get("score", 0), reverse=not ascending)[:n]
    body = []
    for r in ranked:
        chg, cls = _pct(r.get("chg"))
        _bind(prov, f"state.json:wave_top[{r.get('code','?')}]",
              chg=r.get("chg"), rsi=r.get("rsi"), score=r.get("score"), price=r.get("price"))
        rsi = f"{r.get('rsi', 0):.1f}"
        score = f"{r.get('score', 0):.1f}"
        body.append(
            f'<tr>{_stock_cell(nmap, r.get("code"))}'
            f'<td class="l">{_esc(r.get("industry","—"))}</td>'
            f'<td>{_esc(r.get("price","—"))}</td>'
            f'<td class="{cls}">{chg}</td>'
            f'<td>{rsi}</td>'
            f'<td>{score}</td>'
            f'<td>{_esc(r.get("st","—"))}</td></tr>')
    thead = ('<thead><tr><th class="l">個股</th><th class="l">產業</th><th>價</th><th>漲跌</th>'
             '<th>RSI</th><th>強弱分</th><th>趨勢</th></tr></thead>')
    lead = f'<div class="lead" style="color:var(--tx3)">{label}</div>' if label else ""
    return f'<table class="grid">{thead}<tbody>{"".join(body)}</tbody></table>', lead


def sec_S3(state: dict, nmap: dict) -> dict:
    sid, title, tier = "S3", "全市場強弱榜", CFG.SECTION_TIERS["S3"]
    prov = Provenance()
    wave = state.get("wave_top", []) or []
    if not wave:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "state.json 無 wave_top 強弱榜資料。")]}
    head = _sec_head(sid, title, tier)
    lead = (f'<div class="lead">全市場 {len(wave)} 檔依強弱分數排序,報告本體列最強 15 / 最弱 15;'
            f'完整 {len(wave)} 檔附 xlsx 供 Excel/Sheets 排序篩選。</div>')
    t_up, _ = _strength_table(prov, wave, nmap, 15, ascending=False)
    t_dn, l_dn = _strength_table(prov, wave, nmap, 15, ascending=True, label="最弱 15(強弱分數末段)")
    u1 = f'<div class="unit">{head}{lead}{t_up}</div>'
    u2 = f'<div class="unit">{l_dn}{t_dn}</div>'
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [u1, u2]}


def sec_S4(state: dict, chips_week: list[dict], nmap: dict) -> dict:
    sid, title, tier = "S4", "法人與籌碼週流向", CFG.SECTION_TIERS["S4"]
    prov = Provenance()
    if not chips_week:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "本週 twdata/chips 無日檔可加總。")]}
    # 加總本週各日 foreign_net → 排序
    agg: dict[str, float] = {}
    days = []
    for doc in chips_week:
        days.append(doc.get("date"))
        for code, v in (doc.get("data", {}) or {}).items():
            fn = v.get("foreign_net")
            if isinstance(fn, (int, float)):
                agg[code] = agg.get(code, 0) + fn
    if not agg:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "本週 chips 日檔皆無 foreign_net 欄位。")]}
    top = sorted(agg.items(), key=lambda kv: kv[1], reverse=True)[:12]
    bot = sorted(agg.items(), key=lambda kv: kv[1])[:8]
    maxabs = max(abs(top[0][1]), abs(bot[0][1])) or 1

    def cmp_rows(items):
        out = []
        for code, val in items:
            _bind(prov, f"twdata/chips[{code}]:foreign_net(本週日檔加總)", foreign_net_5d=val)
            w = min(100, abs(val) / maxabs * 100)
            cls = "" if val >= 0 else "red"
            vcls = "pos" if val > 0 else "neg"
            out.append(
                f'<div class="row"><span class="lb">{_stock_inline(nmap, code)}</span>'
                f'<div class="track"><div class="fill {cls}" style="width:{w:.0f}%"></div></div>'
                f'<span class="vn {vcls}">{val:+,.0f}</span></div>')
        return "".join(out)

    # 連買榜(consec_top)
    consec = (state.get("chips", {}) or {}).get("consec_top", []) or []
    consec_rows = []
    for r in consec[:6]:
        _bind(prov, f"state.json:chips.consec_top[{r.get('code','?')}]",
              consec=r.get("consec"), net=r.get("net"))
        consec_rows.append(f'<div class="metricrow">{_stock_inline(nmap, r.get("code"))}:連買 '
                           f'<b>{_esc(r.get("consec","—"))}</b> 日({_esc(r.get("side",""))})</div>')
    head = _sec_head(sid, title, tier)
    lead = (f'<div class="lead">加總本週 {len([d for d in days if d])} 個交易日外資買賣超(單位:張),'
            f'排出本週外資淨買/淨賣。日期窗:{_esc(days[0])} ~ {_esc(days[-1])}。</div>')
    u1 = (f'<div class="unit">{head}{lead}'
          f'<div class="two"><div class="block"><div class="bt">本週外資淨買 Top</div>'
          f'<div class="cmp">{cmp_rows(top)}</div></div>'
          f'<div class="block"><div class="bt">本週外資淨賣 Top</div>'
          f'<div class="cmp">{cmp_rows(bot)}</div></div></div>')
    if consec_rows:
        u1 += (f'<div class="block" style="margin-top:11px"><div class="bt">法人連買天數榜</div>'
               f'{"".join(consec_rows)}</div>')
    u1 += "</div>"
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [u1]}


def sec_S5(valdoc: dict, nmap: dict, imap: dict | None = None) -> dict:
    sid, title, tier = "S5", "估值位階雷達", CFG.SECTION_TIERS["S5"]
    prov = Provenance()
    data = (valdoc or {}).get("data", {}) or {}
    if not data:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "無 valuation 估值檔可用。")]}
    yields = [(c, v.get("dividend_yield"), v.get("pe"), v.get("pb")) for c, v in data.items()]
    yld = [(c, y, pe, pb) for c, y, pe, pb in yields if isinstance(y, (int, float)) and y > 0]
    pes = [v.get("pe") for v in data.values() if isinstance(v.get("pe"), (int, float)) and v["pe"] > 0]
    if not yld and not pes:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "valuation 檔內 pe/dividend_yield 全為 null。")]}
    yld.sort(key=lambda t: t[1], reverse=True)
    top_y = yld[:15]
    p25 = statistics.quantiles(pes, n=4)[0] if len(pes) >= 4 else (pes[0] if pes else 0)
    pmed = statistics.median(pes) if pes else 0
    p75 = statistics.quantiles(pes, n=4)[2] if len(pes) >= 4 else (pes[-1] if pes else 0)
    _bind(prov, "valuation:data[*].pe 分布(濾 null 且 pe>0)",
          pe_p25=p25, pe_median=pmed, pe_p75=p75, pe_sample_n=len(pes))

    imap = imap or {}
    body = []
    inds: dict[str, int] = {}
    for c, y, pe, pb in top_y:
        _bind(prov, f"valuation:data[{c}]", dividend_yield=y, pe=pe, pb=pb)
        ind = imap.get(str(c), "")
        if ind:
            inds[ind] = inds.get(ind, 0) + 1
        body.append(
            f'<tr>{_stock_cell(nmap, c)}'
            f'<td class="l">{_esc(ind or "—")}</td>'
            f'<td>{y:.2f}%</td>'
            f'<td>{_esc(f"{pe:.2f}") if isinstance(pe,(int,float)) else "—"}</td>'
            f'<td>{_esc(f"{pb:.2f}") if isinstance(pb,(int,float)) else "—"}</td></tr>')
    thead = ('<thead><tr><th class="l">個股</th><th class="l">產業</th><th>殖利率</th>'
             '<th>本益比</th><th>股價淨值比</th></tr></thead>')
    head = _sec_head(sid, title, tier)
    lead = (f'<div class="lead">濾除 null 後,列全市場現金殖利率 Top 15,並給全市場本益比分布'
            f'(共 {len(pes)} 檔有 PE)。只陳述位置,不判斷貴賤、不構成買賣建議。</div>')
    dist = (f'<div class="block" style="margin-bottom:11px"><div class="bt">全市場本益比分布</div>'
            f'<div class="metricrow">P25 <b>{p25:.1f}</b> 倍　·　中位數 <b>{pmed:.1f}</b> 倍　·　'
            f'P75 <b>{p75:.1f}</b> 倍</div></div>')
    # 🔴 REVIEW #6:榜單前段幾乎全是營建股(高殖利率多為建案認列的一次性配發),而本頻道受眾
    # 正是「怕被割的小白」——只列榜不警示,這段對他們不是沒用,是**危險**。
    # 揭露必須是「事實陳述+機制說明」,不能變成買賣建議:講清楚殖利率怎麼算出來的、
    # 為什麼會不可持續、產業集中在哪,讓讀者自己判斷。集中度數字由榜單實算,綁 provenance。
    top_ind = sorted(inds.items(), key=lambda kv: -kv[1])[:2]
    conc = ""
    if top_ind and top_ind[0][1] >= 2:
        _bind(prov, "derived:Top15 榜單的 twstock 產業別計數",
              **{f"top15_industry_count[{top_ind[0][0]}]": top_ind[0][1]})
        parts = "、".join(f"{k} {v} 檔" for k, v in top_ind)
        conc = (f'本期 Top 15 的產業分布集中在 <b>{_esc(parts)}</b>(共 {len(top_y)} 檔)。')
    warn = (f'<div class="block" style="margin-top:11px"><div class="bt">讀這張表之前(風險揭露)</div>'
            f'<div class="metricrow">{conc}'
            f'殖利率 = <b>過去已配發的現金股利 ÷ 現價</b>,是<b>回頭看</b>的數字,不是未來會配多少。<br>'
            f'· <b>景氣循環股/建案認列型公司</b>(如營建、航運)常因一次性獲利而配出高股利,'
            f'該筆獲利認列完就可能大幅下降——<b>高殖利率不等於好標的,也不等於配得久</b>。<br>'
            f'· 殖利率變高也可能是<b>股價跌下來</b>造成的(分母變小),不必然是好事。<br>'
            f'· 本表<b>只陳述位置、不推薦任何個股</b>;要判斷可持續性,請自行看該公司的獲利'
            f'來源是否為常態性業務。<span style="color:var(--tx3)">介紹 ≠ 推薦。</span></div></div>')
    u1 = (f'<div class="unit">{head}{lead}{dist}'
          f'<table class="grid">{thead}<tbody>{"".join(body)}</tbody></table>{warn}</div>')
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [u1]}


def sec_S6(state: dict, nmap: dict | None = None) -> dict:
    # 🔴 REVIEW 續訂殺手 #2:舊標題「誠實成績單」把這段framing成「我們的訊號表現」,
    # 於是 15.8% 勝率/-7.77% 就變成「我在付錢訂閱一個會賠錢的訊號」——而且它在 basic,
    # 最低價層的門面自曝其短。內容一個字都不用改(那是全報告最可敬的一段),
    # 改的是**它在回答什麼問題**:不是「我們的訊號多準」,而是「為什麼這份報告不賣訊號」。
    # 同一批數字,從勸退變成差異化信任錨。
    sid, title, tier = "S6", "為什麼我們不賣訊號", CFG.SECTION_TIERS["S6"]
    prov = Provenance()
    tr = state.get("track", {}) or {}
    if not tr or tr.get("n_closed") is None:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "state.json 無 track 訊號追蹤資料。")]}
    n_closed = tr.get("n_closed", 0)
    wr = (tr.get("win_rate") or 0) * 100
    lwr = (tr.get("long_win_rate") or 0) * 100
    swr = (tr.get("short_win_rate") or 0) * 100
    avg_r = tr.get("avg_r")
    avg_ret = tr.get("avg_ret_pct")
    # 招牌數字:全部綁 state.json:track 的原始欄位(勝率/R/報酬皆由 track 聚合直出)
    _bind(prov, "state.json:track",
          n_closed=n_closed, win_rate_pct=wr, long_win_rate_pct=lwr, short_win_rate_pct=swr,
          avg_r=avg_r, avg_ret_pct=avg_ret, n_open=tr.get("n_open"))
    ret_s, ret_cls = _pct(avg_ret)
    head = _sec_head(sid, title, tier)
    lead = ('<div class="lead">市面上的老師只曬贏單。我們把自家程式訊號的<b>真實平倉結果全部攤開'
            '(含輸單、不挑不藏)</b>,而它現在是<b>虧的</b>——這正是這份報告<b>只賣數據、'
            '不喊單、不報明牌</b>的原因。<br>如果連我們自己跑出來的訊號都證明不了穩定的 edge,'
            '那任何人叫你「跟單」都該被你懷疑。你付的錢買的是<b>能自己查證的原始數據</b>,'
            '不是明牌。</div>')
    kpis = (
        f'<div class="kpis">'
        f'<div class="kpi"><div class="k">已平倉</div><div class="v">{_esc(n_closed)}<small> 筆</small></div></div>'
        f'<div class="kpi"><div class="k">總勝率</div><div class="v">{wr:.1f}<small>%</small></div></div>'
        f'<div class="kpi"><div class="k">平均 R 值</div><div class="v {"pos" if (avg_r or 0)>0 else "neg"}">'
        f'{_esc(f"{avg_r:+.2f}") if avg_r is not None else "—"}</div></div>'
        f'<div class="kpi"><div class="k">平均報酬</div><div class="v {ret_cls}">{ret_s}</div></div>'
        f'</div>')
    # 樣本不足揭露:用 Wilson 區間把「n=19 不可外推」從口號變成可查證的事實。
    # k 由 win_rate×n_closed 還原(track 只給比例不給勝場數),故標「約」。
    k_win = int(round((tr.get("win_rate") or 0) * n_closed))
    ci = wilson_ci(k_win, n_closed)
    ci_html = ""
    if ci and n_closed > 0:
        lo, hi = ci[0] * 100, ci[1] * 100
        _bind(prov, "derived:Wilson score 95% CI(k=round(win_rate×n_closed), n=n_closed)",
              win_ci_lo_pct=lo, win_ci_hi_pct=hi, k_wins=k_win)
        # ⚠️ 句子為什麼要這樣斷:守門(FG)是用「同一子句內有沒有績效語境詞」判斷一個 % 是不是
        # 績效宣稱,而子句只以「。！?\n」分界。若把「95% 信賴區間」寫在含「勝率」的同一句裡,
        # **信心水準的 95% 會被誤判成一個查無來源的績效數字**,整段 S6 就被 fail-closed 隱藏
        # (實測如此)。正解**不是**把 95 灌進池——那會讓池多一個 95,日後真有人捏造「勝率 95%」
        # 就撞得到鄰居,等於為了讓自己的文案過關而弱化守門。正解是把「信心水準」獨立成一句:
        # 它本來就不是績效宣稱,不該被當成績效宣稱。區間端點(5.5/37.6)仍照常綁來源入池。
        ci_html = (f'<br><b>樣本不足,不可外推</b>:僅 <b>{n_closed}</b> 筆已平倉,'
                   f'這個勝率的信賴區間約 <b>{lo:.1f}% ~ {hi:.1f}%</b> —— '
                   f'區間寬到<b>無法下任何結論</b>(不論好壞)。'
                   f'(此區間以 Wilson score 法計算,信心水準 95%。)'
                   f'把 {wr:.1f}% 當成「這套訊號的真實勝率」是錯的;'
                   f'把它當成「訊號無效的鐵證」也一樣不嚴謹。'
                   f'我們照實貼出來,是要你知道<b>樣本這麼小的時候,誰都不該給你結論</b>——'
                   f'包括我們自己。')
    detail = (f'<div class="block" style="margin-top:11px"><div class="metricrow">'
              f'多方勝率 <b>{lwr:.1f}%</b>　·　空方勝率 <b>{swr:.1f}%</b>　·　'
              f'目前未平倉 <b>{_esc(tr.get("n_open","—"))}</b> 筆<br>'
              f'<span style="color:var(--tx3)">R 值 = 報酬 ÷ 進場風險;負值代表這批訊號目前是虧的——'
              f'我們照實呈現,不美化。</span>{ci_html}</div></div>')
    # 近期逐筆(只列已平倉的真實結果)
    recent = [r for r in (tr.get("recent", []) or []) if r.get("result") not in (None, "open")][:6]
    rec_rows = ""
    if recent:
        body = []
        for r in recent:
            rs, rcls = _pct(r.get("ret_pct"))
            _bind(prov, f"state.json:track.recent[{r.get('code','?')}]",
                  ret_pct=r.get("ret_pct"), r=r.get("r"), entry=r.get("entry"), exit=r.get("exit"))
            body.append(
                f'<tr>{_stock_cell(nmap or {}, r.get("code"))}'
                f'<td>{_esc(r.get("side","—"))}</td><td>{_esc(r.get("entry","—"))}</td>'
                f'<td>{_esc(r.get("exit","—"))}</td><td class="{rcls}">{rs}</td>'
                f'<td>{_esc(r.get("exit_reason","—"))}</td></tr>')
        thead = ('<thead><tr><th class="l">個股</th><th>方向</th><th>進</th><th>出</th>'
                 '<th>報酬</th><th>出場原因</th></tr></thead>')
        rec_rows = (f'<div class="block" style="margin-top:11px"><div class="bt">近期已平倉逐筆</div>'
                    f'<table class="grid">{thead}<tbody>{"".join(body)}</tbody></table></div>')
    html = f'<div class="unit">{head}{lead}{kpis}{detail}{rec_rows}</div>'
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": [html]}


def sec_S7(checkup: dict, window_days: int = 7, requests_doc: dict | None = None,
           ensure_fn=None, nmap: dict | None = None) -> dict:
    """本期深度體檢:**點播優先**,其餘用本週輪替補。

    requests_doc:點播佇列(會**就地更新 status/attempts**,由 generate_weekly 負責存檔);
                 None 或無 pending → 完全等同改版前的行為(有測試釘死)。
    ensure_fn   :把不在庫的檔現算出來的函式(可注入;預設打真引擎會抓行情,測試注入 stub)。
    """
    sid, title, tier = "S7", "本期深度體檢個股", CFG.SECTION_TIERS["S7"]
    prov = Provenance()
    results = checkup.get("results", {}) or {}
    by_code = checkup.get("by_code", {}) or {}

    def _fact_ok(f):
        return (isinstance(f, dict) and f.get("key") and (f.get("claim") or f.get("summary"))
                and f.get("source") and f.get("data") not in (None, "", [], {}))

    def _parse(s):
        try:
            return datetime.strptime((s or "")[:10], "%Y-%m-%d").date()
        except Exception:  # noqa: BLE001
            return None

    def _facts_of(code: str, res: dict) -> list:
        """該 code 全部通過守門的事實(**不看時間窗**)——點播檔要的是「我要的那檔」,
        不是「這週剛好算到的那檔」。"""
        out = []
        for key, f in res.items():
            if not _fact_ok(f):
                continue
            parts = key.split("__")
            if len(parts) < 2 or parts[1] != code:
                continue
            out.append(f)
        return out

    since = TODAY - timedelta(days=window_days)
    # ── ① 先處理點播(佇列空 → pend=[] → 完全走原本的輪替路徑,行為不變)──────────
    pend = pending_requests(requests_doc, REQ_MAX_PER_ISSUE)
    req_codes: list[str] = []
    req_notes: list[str] = []
    ensure = ensure_fn if ensure_fn is not None else _ensure_checkup_via_engine
    for r in pend:
        code = str(r.get("code", "")).strip()
        if not code:
            continue
        if _facts_of(code, results):
            r["status"] = "fulfilled"; r["fulfilled_in"] = TODAY.isoformat()
            req_codes.append(code)
            continue
        # 不在事實庫 → 現算(fail-safe:算不出來只降級說明,不炸、不編)
        r["attempts"] = int(r.get("attempts") or 0) + 1
        ok = False
        try:
            ok = bool(ensure(code))
        except Exception as exc:  # noqa: BLE001
            r["last_error"] = f"{type(exc).__name__}: {exc}"[:160]
        if ok:
            checkup = load_checkup()
            results = checkup.get("results", {}) or {}
            by_code = checkup.get("by_code", {}) or {}
        if ok and _facts_of(code, results):
            r["status"] = "fulfilled"; r["fulfilled_in"] = TODAY.isoformat(); r["last_error"] = None
            req_codes.append(code)
        else:
            r["last_error"] = r.get("last_error") or "體檢引擎本期未能算出此檔(可能資料源暫時不可用)"
            if r["attempts"] >= REQ_MAX_ATTEMPTS:
                r["status"] = "failed"
                req_notes.append(f'點播的 <b>{_esc(code)}</b> 連續 {r["attempts"]} 期都沒能算出來,'
                                 f'已從佇列移除並會個別通知;<b>不佔用你的點播額度</b>。')
            else:
                req_notes.append(f'點播的 <b>{_esc(code)}</b> 這期沒能算出來(資料源暫時不可用),'
                                 f'<b>已保留在佇列、下期優先處理</b>,不佔用你的額度。')

    # ── ② 本週輪替:窗內新完成的體檢事實(排除已被點播納入的檔)────────────────
    by_stock: dict[str, list] = {}
    for code in req_codes:                       # 點播檔排最前面
        fs = _facts_of(code, results)
        if fs:
            by_stock[code] = fs
    for key, f in results.items():
        if not _fact_ok(f):
            continue
        d = _parse(f.get("computed_at"))
        if d is not None and d < since:
            continue
        # key 形如 checkup_<type>__<code>[__<window>] —— code 一律在 index 1
        # (crash 為 checkup_crash__<code>__<crisis2008|covid2020|bear2022>,不能用 [-1])
        parts = key.split("__")
        if len(parts) < 2:
            continue
        code = parts[1]
        if code in req_codes:                    # 已由點播納入,別重複
            continue
        by_stock.setdefault(code, []).append(f)
    if not by_stock:
        reason = f"本週窗({since}~{TODAY})內無新完成、且通過溯源守門的體檢事實。"
        if req_notes:
            reason += " " + re.sub(r"<[^>]+>", "", " ".join(req_notes))
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, reason)]}
    units = []
    head = _sec_head(sid, title, tier)
    n_rot = len(by_stock) - len(req_codes)
    if req_codes:
        lead = (f'<div class="lead">本期含 <b>{len(req_codes)} 檔訂戶點播</b>'
                f'(完整版每月可指定 1 檔,先到先服務)+ {n_rot} 檔本週輪替;'
                f'以下每則都逐字引用體檢引擎算出的既有事實(含息還原/套牢/腰斬/崩盤三段/'
                f'毛利/股利/估值位階),各附來源。</div>')
    else:
        lead = (f'<div class="lead">本週體檢管線新覆蓋 {len(by_stock)} 檔;以下每則都逐字引用體檢引擎'
                f'算出的既有事實(含息還原/套牢/腰斬/崩盤三段/毛利/股利/估值位階),各附來源。'
                f'<br><span style="color:var(--tx3)">完整版訂戶每月可<b>點播 1 檔</b>指定個股'
                f'——想看自己手上那檔,不用等輪到。</span></div>')
    if req_notes:
        lead += (f'<div class="block" style="margin-top:9px"><div class="metricrow" '
                 f'style="color:var(--tx3)">{"<br>".join(req_notes)}</div></div>')
    first = True
    order = ["checkup_long_horizon", "checkup_annual_extremes", "checkup_underwater",
             "checkup_valuation_position", "checkup_three_way", "checkup_crash",
             "checkup_gross_margin", "checkup_dividend_history"]
    for code, facts in by_stock.items():
        name = by_code.get(code, {}).get("name") or (nmap or {}).get(code) or code
        badge = ('<span class="tier" style="margin-left:6px">訂戶點播</span>'
                 if code in req_codes else "")
        facts_sorted = sorted(facts, key=lambda f: order.index(f["key"].split("__")[0])
                              if f["key"].split("__")[0] in order else 99)
        rows = []
        for f in facts_sorted[:8]:
            txt = (f.get("claim") or f.get("summary") or "").strip()
            if not txt:
                continue
            # 體檢數字為既有事實 —— 把 fact.data 內數字入池(來源=fact_key),原樣引用 claim
            for val in FG._walk_numbers(f.get("data", {})):
                prov.pool.add(abs(val)); prov.pool.add(abs(round(val, 1)))
            _pool(prov, txt)
            # A2 修:不截斷。v1 存 txt[:60],51/71 筆被切斷、部分斷在數字中間
            #(如「卡瑪比率 0.」),PDF 渲染的是完整文字,但稽核檔的尾巴壞掉 →
            # 稽核者看到的像是壞數字。存證檔要能取信於人就不能自己先失真。
            prov.records.append({"field": f.get("key"), "value": None, "text": txt,
                                 "source": f.get("source", ""),
                                 "note": "體檢引擎既有事實(逐字引用)"})
            rows.append(f'<div class="metricrow">• {_esc(txt)}</div>')
        card = (f'<div class="card"><div class="kpis" style="margin-bottom:8px">'
                f'<div class="kpi" style="flex:none;text-align:left;min-width:0"><div class="k">個股</div>'
                f'<div class="v" style="font-size:17px">{_esc(name)} <span class="code" '
                f'style="font-size:11px;color:var(--tx3)">{_esc(code)}</span>{badge}</div></div></div>'
                f'{"".join(rows)}</div>')
        if first:
            units.append(f'<div class="unit">{head}{lead}{card}</div>')
            first = False
        else:
            units.append(f'<div class="unit">{card}</div>')
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False, "units": units}


def s8_theme_index(today: date | None = None) -> int:
    """本月主題編號(0/1/2)。以年月決定 → 同一個月固定、跨月必換,可預測可測試。"""
    d = today or TODAY
    return (d.year * 12 + d.month) % 3


def sec_S8(adaptive: list[dict], nmap: dict | None = None, today: date | None = None) -> dict:
    """結構基準:每月換一個主題,三個主題輪替。

    🔴 REVIEW 續訂殺手 #1(最確定的退訂觸發器):舊版標題寫「月度輪替」,實際是把同一份
    2026-06-12 的靜態快照**每週渲染出完全相同的數字**——資料 35 天沒動,訂戶連看 4 週就會
    發現八分之一是複製貼上,而且「輪替」這個詞當時是**假的**(什麼都沒在輪)。

    修法與取捨(三選一之外的第四條,理由寫在這裡供覆核):
      ① 「每月真重跑回測」是**正解但不在我的權限內**——重跑 1770 檔要抓外部行情(踩「不打
         外部 API」紅線)且需要排程,那是 ops 變更。→ 已在回報裡標為建議另開 task。
      ② 「直接拿掉」會把 S8 唯一的價值(「趨勢策略無腦套全市場只有 53% 正報酬」這種紮實的
         觀念矯正、且反著自己利益講)一起丟掉,可惜。
      ③ 所以採「**真的輪替**」:三個主題各用一份**真實存在**的基準資料(趨勢/多空/風險調整),
         依年月決定,跨月必換 → 「輪替」從假話變成真話,且每月內容真的不同。
    ⚠️ 但要誠實面對一件事:**月度輪替並不能讓「每週不一樣」**——同一個月內的四週本來就會
    一樣。所以真正的修不是假裝它每週都新,而是**把規則講白**:本段明寫「每月換一次主題、
    月內固定」+ 明標快照日與距今天數。訂戶第 2 週看到一樣的東西時,那是**他早就知道的設計**,
    不是被騙。undisclosed 的重複才是退訂觸發器,disclosed 的月更節奏不是。
    """
    sid, tier = "S8", CFG.SECTION_TIERS["S8"]
    prov = Provenance()
    ti = s8_theme_index(today)
    snap_map = {0: ADAPTIVE_CSV, 1: TWDATA / "longshort_per_stock.csv", 2: TWDATA / "per_stock_results.csv"}
    snap_day, snap_age = _snapshot_age(snap_map[ti])
    theme_names = {0: "趨勢策略無腦套全市場", 1: "加了放空會更好嗎", 2: "風險調整後還剩多少"}
    title = f"結構基準 · 本月主題:{theme_names[ti]}"
    head = _sec_head(sid, title, tier)

    def _shell(inner: str, extra_note: str = "") -> str:
        stale = ""
        if snap_age >= 0:
            _bind(prov, f"derived:{snap_map[ti].name} 檔案時間 vs 今天", snapshot_age_days=snap_age)
            stale = (f'資料為 <b>{_esc(snap_day)}</b> 的靜態快照(距今 <b>{snap_age}</b> 天),'
                     f'非即時、非可交易訊號。')
        rule = ('<div class="block" style="margin-top:11px"><div class="metricrow" '
                'style="color:var(--tx3)">'
                f'<b>這段的更新規則(先講清楚,免得你以為我們在灌水)</b>:本段是<b>教育性基準</b>,'
                f'<b>每月換一次主題</b>(共 3 個主題輪替),<b>同一個月內的每週內容相同</b>——'
                f'它要傳達的是不隨盤勢變動的結構性事實,不是每週追新聞。{stale}'
                f'{extra_note}</div></div>')
        return f'<div class="unit">{head}{inner}{rule}</div>'

    # ── 主題 0:趨勢策略無腦套全市場(adaptive)────────────────────────────────
    if ti == 0:
        if not adaptive:
            return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                    "units": [_degrade_unit(sid, title, "無 adaptive_per_stock.csv 全市場回測快照。")]}
        n = len(adaptive)
        nets = [r["a_net"] for r in adaptive]
        med = statistics.median(nets)
        n_pos = len([x for x in nets if x > 0])
        pct_pos = n_pos / n * 100
        top = sorted(adaptive, key=lambda r: r["a_net"], reverse=True)[:5]
        _bind(prov, "twdata/adaptive_per_stock.csv(全樣本未濾)",
              sample_n=n, a_net_median=med, n_positive=n_pos, pct_positive=pct_pos)
        lead = ('<div class="lead">把趨勢策略<b>無腦套全市場</b>,真正該看的是<b>中位數</b>,'
                '別被最好那幾檔騙走——這是本月的觀念矯正。</div>')
        kpis = (
            f'<div class="kpis">'
            f'<div class="kpi"><div class="k">樣本檔數</div><div class="v">{n}</div></div>'
            f'<div class="kpi"><div class="k">淨報酬中位數</div><div class="v {"pos" if med>0 else "neg"}">{med:+.1f}<small>%</small></div></div>'
            f'<div class="kpi"><div class="k">正報酬佔比</div><div class="v">{pct_pos:.1f}<small>%</small></div></div>'
            f'</div>')
        body = []
        for r in top:
            _bind(prov, f"twdata/adaptive_per_stock.csv[{r.get('code','?')}]",
                  a_net=r["a_net"], a_win=r["a_win"])
            s, cls = _pct(r["a_net"])
            body.append(f'<tr>{_stock_cell(nmap or {}, r.get("code"))}'
                        f'<td class="{cls}">{s}</td><td>{r["a_win"]:.1f}%</td></tr>')
        tbl = (f'<div class="block" style="margin-top:11px"><div class="bt">自適應淨報酬前 5(教育示例)</div>'
               f'<table class="grid"><thead><tr><th class="l">個股</th><th>淨報酬</th><th>勝率</th></tr></thead>'
               f'<tbody>{"".join(body)}</tbody></table></div>')
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False,
                "units": [_shell(f"{lead}{kpis}{tbl}")]}

    # ── 主題 1:純多 vs 多空(longshort)—— 「加放空會更好嗎」──────────────────
    if ti == 1:
        ls = _load_csv_floats(snap_map[1], ("l_net", "ls_net"))
        if not ls:
            return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                    "units": [_degrade_unit(sid, title, "無 longshort_per_stock.csv 多空基準快照。")]}
        n = len(ls)
        l_med = statistics.median([r["l_net"] for r in ls])
        ls_med = statistics.median([r["ls_net"] for r in ls])
        better = len([r for r in ls if r["ls_net"] > r["l_net"]])
        pct_better = better / n * 100
        _bind(prov, "twdata/longshort_per_stock.csv(全樣本未濾)",
              sample_n=n, long_only_net_median=l_med, longshort_net_median=ls_med,
              n_better_with_short=better, pct_better_with_short=pct_better)
        lead = ('<div class="lead">很多人以為「會放空才是高手」。本月拿同一套策略跑<b>純做多</b> vs '
                '<b>多空都做</b>,看加了放空到底有沒有比較好——答案不一定是你想的那樣。</div>')
        kpis = (
            f'<div class="kpis">'
            f'<div class="kpi"><div class="k">樣本檔數</div><div class="v">{n}</div></div>'
            f'<div class="kpi"><div class="k">純多中位數</div><div class="v {"pos" if l_med>0 else "neg"}">{l_med:+.1f}<small>%</small></div></div>'
            f'<div class="kpi"><div class="k">多空中位數</div><div class="v {"pos" if ls_med>0 else "neg"}">{ls_med:+.1f}<small>%</small></div></div>'
            f'<div class="kpi"><div class="k">加空後變好</div><div class="v">{pct_better:.1f}<small>%</small></div></div>'
            f'</div>')
        note = (f'<div class="block" style="margin-top:11px"><div class="metricrow">'
                f'同一批 <b>{n}</b> 檔、同一套策略,只差「有沒有做空」:'
                f'加了放空之後淨報酬變好的只有 <b>{better} / {n}</b>(<b>{pct_better:.1f}%</b>)。'
                f'<br><span style="color:var(--tx3)">這是歷史統計,不是叫你別放空、也不是叫你放空——'
                f'台股放空另有借券成本與平盤下限制,本回測未必涵蓋你的實際條件。介紹 ≠ 推薦。</span>'
                f'</div></div>')
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False,
                "units": [_shell(f"{lead}{kpis}{note}")]}

    # ── 主題 2:風險調整後(per_stock:sharpe / 報酬回撤比)────────────────────
    ps = _load_csv_floats(snap_map[2], ("sharpe", "max_dd_pct", "net_profit_pct"))
    if not ps:
        return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": True,
                "units": [_degrade_unit(sid, title, "無 per_stock_results.csv 風險基準快照。")]}
    n = len(ps)
    sh = [r["sharpe"] for r in ps]
    sh_med = statistics.median(sh)
    n_sh1 = len([x for x in sh if x >= 1.0])
    pct_sh1 = n_sh1 / n * 100
    dd_med = statistics.median([abs(r["max_dd_pct"]) for r in ps])
    _bind(prov, "twdata/per_stock_results.csv(全樣本未濾)",
          sample_n=n, sharpe_median=sh_med, n_sharpe_ge_1=n_sh1, pct_sharpe_ge_1=pct_sh1,
          max_dd_pct_median=dd_med)
    # ⚠️ 這句原本寫「同樣賺 50%,一路平順 vs 中間先跌一半」——那個 50% 是**修辭舉例**、
    # 不是資料,但守門看到績效語境裡的裸 % 就會擋(實測擋下)。正解不是把 50 灌進池
    # (那會讓池多一個整十數,日後捏造的「報酬 50%」就撞得到鄰居 → 弱化守門),
    # 而是**把不必要的數字拿掉**:這句話不靠那個 50% 也成立。
    lead = ('<div class="lead">只看報酬會騙人:同樣的總報酬,一路平順 vs 中間先讓你腰斬一次,'
            '是兩種完全不同的東西。本月看<b>風險調整後</b>還剩多少——Sharpe,'
            '以及你實際要吞下去的最大回撤。</div>')
    kpis = (
        f'<div class="kpis">'
        f'<div class="kpi"><div class="k">樣本檔數</div><div class="v">{n}</div></div>'
        f'<div class="kpi"><div class="k">Sharpe 中位數</div><div class="v {"pos" if sh_med>0 else "neg"}">{sh_med:.2f}</div></div>'
        f'<div class="kpi"><div class="k">Sharpe ≥ 1 佔比</div><div class="v">{pct_sh1:.1f}<small>%</small></div></div>'
        f'<div class="kpi"><div class="k">最大回撤中位數</div><div class="v neg">{dd_med:.1f}<small>%</small></div></div>'
        f'</div>')
    note = (f'<div class="block" style="margin-top:11px"><div class="metricrow">'
            f'全市場 <b>{n}</b> 檔裡,Sharpe 站上 1.0 的只有 <b>{n_sh1}</b> 檔(<b>{pct_sh1:.1f}%</b>);'
            f'而中位數的最大回撤是 <b>{dd_med:.1f}%</b> —— 意思是<b>一半以上的標的,'
            f'歷史上都曾經讓你帳面腰斬級別的難受</b>。<br>'
            f'<span style="color:var(--tx3)">Sharpe = 每承擔一單位波動換到的超額報酬,'
            f'越高代表報酬相對於波動越划算。這是歷史統計,不預測未來。介紹 ≠ 推薦。</span>'
            f'</div></div>')
    return {"id": sid, "title": title, "tier": tier, "prov": prov, "degraded": False,
            "units": [_shell(f"{lead}{kpis}{note}")]}


# ══════════════════════════════════════════════════════════════════════════════
#  誠信 gate:組完各 section 後嚴容差複驗;查無來源 → 該 section 降級(fail-closed)
# ══════════════════════════════════════════════════════════════════════════════
def gate_or_degrade(section: dict) -> dict:
    """對非降級 section 跑 Provenance.gate;有查無來源的績效數字 → 整段換降級卡。"""
    if section.get("degraded"):
        return section
    text = _gate_text_from_units(section["units"])
    bad = section["prov"].gate(text, section["id"])
    if bad:
        vals = ", ".join(str(c["value"]) for c in bad[:4])
        section["units"] = [_degrade_unit(
            section["id"], section["title"],
            f"溯源守門擋下查無來源的數字({vals}）,為誠信 fail-closed 已隱藏本段(不出未溯源數字)。")]
        section["degraded"] = True
        section["gated_out"] = True
    return section


# ══════════════════════════════════════════════════════════════════════════════
#  HTML 組裝 + 分頁 + 渲染
# ══════════════════════════════════════════════════════════════════════════════
def _cover_html(state: dict, sections: list[dict], week_label: str, tier: str) -> str:
    g = state.get("gauge", {}) or {}
    n_sec = sum(1 for s in sections if not s.get("degraded"))
    wave_n = len(state.get("wave_top", []) or [])
    sect_n = len(state.get("sectors", []) or [])
    tr = state.get("track", {}) or {}
    b = CFG.BRAND
    sub = CFG.SUBSCRIPTION
    tinfo = sub["full"] if tier == "full" else sub["basic"]
    price = (f'訂閱:{sub["basic"]["name_zh"]} NT${sub["basic"]["ntd_month"]}/月　·　'
             f'{sub["full"]["name_zh"]} NT${sub["full"]["ntd_month"]}/月　·　'
             f'{sub["annual"]["name_zh"]} NT${sub["annual"]["ntd_year"]}/年')
    return f'''<section class="page cover">
  <div class="frame"></div>
  <div class="inner">
    <div class="runhead">
      <div class="brand"><span class="logo"></span><b>{_esc(b["name_zh"])}</b>&nbsp;{_esc(b["name_en"])}</div>
      <div class="vol"><div class="eyebrow">訂閱週報 · WEEKLY</div>
        <div style="margin-top:6px">{_esc(week_label)}</div>
        <div>{_esc(tinfo["name_zh"])}</div></div>
    </div>
    <div class="masthead">
      <div class="kick">全市場掃描 · 數據週報</div>
      <h1>台股全市場<br><span class="thin">週報</span></h1>
      <div class="sub">{_esc(b["tagline_zh"])}——主體吃全市場強弱掃描,體檢為深度加值層。
        每個數字都綁真實來源欄位,查無來源該段自動隱藏。</div>
      <div class="cover-stats">
        <div class="cstat"><div class="k">市場溫度</div><div class="v">{_esc(f"{g.get('temperature','—')}")}<small>{_esc(g.get('label',''))}</small></div></div>
        <div class="cstat"><div class="k">強弱榜覆蓋</div><div class="v">{wave_n}<small>檔</small></div></div>
        <div class="cstat"><div class="k">板塊</div><div class="v">{sect_n}<small>類</small></div></div>
        <div class="cstat"><div class="k">訊號追蹤</div><div class="v">{_esc(tr.get('n_closed','—'))}<small>已平倉</small></div></div>
      </div>
      <div class="cover-strip">
        <span class="badge">介紹 ≠ 推薦</span>
        <span class="txt">本報告為<b>歷史/當期數據彙整</b>,中性陳述、不喊買賣、不報明牌。含<b>真實訊號成績單(含輸單)</b>。</span>
      </div>
      <div class="price-line">{_esc(price)}　｜　平台:{_esc(sub["platform_tw"])}</div>
    </div>
    <div class="runfoot"><span>{_esc(b["name_zh"])} {_esc(b["name_en"])} · {_esc(b["product_zh"])}</span>
      <span class="disc">{_esc(week_label)}</span><span>封面 · 共 <span class="pagetotal"></span> 頁</span></div>
  </div>
</section>'''


def _page_template(brand: dict, product: str, week_label: str) -> str:
    return f'''<template id="pagetpl"><section class="page">
  <div class="frame"></div>
  <div class="inner">
    <div class="runhead">
      <div class="brand"><span class="logo"></span><b>{_esc(brand["name_zh"])}</b>&nbsp;{_esc(product)}</div>
      <div class="eyebrow">{_esc(week_label)}</div>
    </div>
    <div class="flow"></div>
    <div class="runfoot"><span>{_esc(brand["name_zh"])} {_esc(brand["name_en"])} · {_esc(product)}</span>
      <span class="disc">介紹 ≠ 推薦</span>
      <span><span class="pageno"></span> / <span class="pagetotal"></span></span></div>
  </div>
</section></template>'''


def _disclaimer_unit(state: dict, sections: list[dict]) -> str:
    srcs = ["data_hunter/state.json(全市場強弱掃描)",
            "twdata/chips/*.json(法人買賣超)",
            "twdata/fundamentals/valuation_*.json(估值:pe/殖利率/pb)",
            "youtube_channel/STUDIO/stock_checkup_facts.json(深度體檢,FinMind + Yahoo 含息還原)",
            "twdata/adaptive_per_stock.csv(全市場回測靜態快照,2026-06-12)"]
    degraded = [s["id"] for s in sections if s.get("degraded")]
    dnote = ("　本期自動降級/隱藏的段落:" + "、".join(degraded)) if degraded else "　本期 8 段全數呈現。"
    return (f'<div class="unit"><div class="disc-title">免責與資料來源</div>'
            f'<div class="disc-body">{_esc(DISCLAIMER)}<br><br>'
            f'本週報主體為<b>當期全市場掃描</b>與<b>歷史數據彙整</b>,不是即時可交易訊號;'
            f'S8 結構基準為 <b>2026-06-12 靜態回測快照</b>,僅供教育,不代表現在或未來。'
            f'S6 訊號追蹤為程式訊號的真實平倉紀錄(含輸單),為誠實揭露、非邀約跟單。'
            f'估值段(S5/S7)只陳述數據位置,<b>不判斷貴賤、不構成買賣建議</b>。</div>'
            f'<div class="disc-src"><b>資料來源</b>:' + "；".join(_esc(s) for s in srcs) + "。" + _esc(dnote) +
            f'　每個績效數字經溯源守門(fail-closed)驗證,查無來源的段落已自動隱藏。</div></div>')


_PAGINATE_JS = '''<script>
(function(){
  var src = document.getElementById('src');
  var host = document.getElementById('pages');
  var tpl = document.getElementById('pagetpl');
  var units = Array.prototype.slice.call(src.children);
  var pageNo = 1; // 封面=第1頁
  var cur=null, flow=null;
  function newPage(){
    var p = tpl.content.firstElementChild.cloneNode(true);
    host.appendChild(p);
    cur=p; flow=p.querySelector('.flow'); pageNo++;
    p.querySelector('.pageno').textContent = pageNo;
    return p;
  }
  newPage();
  for(var i=0;i<units.length;i++){
    var u = units[i];
    flow.appendChild(u);
    if(flow.scrollHeight > flow.clientHeight + 1){
      flow.removeChild(u);
      newPage();
      flow.appendChild(u);
      // 若單一 unit 仍超高(理論上罕見),就留在該頁不再無限開頁
    }
  }
  var total = host.querySelectorAll('.page').length;
  Array.prototype.forEach.call(document.querySelectorAll('.pagetotal'), function(e){ e.textContent = total; });
  src.parentNode.removeChild(src);
  window.__paginated__ = true;
})();
</script>'''


def build_html(state: dict, sections: list[dict], week_label: str, tier: str) -> str:
    css = THEME_CSS.read_text(encoding="utf-8")
    b = CFG.BRAND
    product = b["product_zh"]
    # 依 tier 過濾 section(basic 不含 full-only)
    visible = [s for s in sections
               if tier == "full" or CFG.SECTION_TIERS.get(s["id"]) == "basic"]
    unit_html = []
    for s in visible:
        unit_html.extend(s["units"])
    unit_html.append(_disclaimer_unit(state, visible))
    cover = _cover_html(state, visible, week_label, tier)
    pagetpl = _page_template(b, product, week_label)
    return f'''<!doctype html><html lang="zh-Hant"><head><meta charset="utf-8">
<title>{_esc(b["name_zh"])} {_esc(product)} {_esc(week_label)}</title>
<style>{css}</style></head><body>
<div id="pages">{cover}</div>
{pagetpl}
<div id="src" style="position:absolute;left:-99999px;top:0;width:186mm">{"".join(unit_html)}</div>
{_PAGINATE_JS}
</body></html>'''


def render_pdf(html: str, out_pdf: Path) -> Path:
    """用 mockup/render.py 同款 Playwright 參數印 PDF(print_background/A4/繁中),
       並等分頁 JS 跑完(window.__paginated__)。"""
    from playwright.sync_api import sync_playwright
    out_pdf.parent.mkdir(parents=True, exist_ok=True)
    html_path = out_pdf.with_suffix(".html")
    html_path.write_text(html, encoding="utf-8")
    with sync_playwright() as p:
        br = p.chromium.launch()
        pg = br.new_page()
        pg.goto(html_path.resolve().as_uri(), wait_until="networkidle")
        pg.wait_for_function("window.__paginated__ === true", timeout=15000)
        pg.emulate_media(media="print")
        # format="A4" 與 report_theme.css 的 @page{size:A4} 是雙保險:prefer_css_page_size
        # 只在 CSS 真的宣告頁面尺寸時才有東西可 prefer,否則**靜默退回 Letter** → 每頁溢
        # 50pt 擠出整頁空白(實測 full 24 頁裡 11 頁空白)。兩者都指 A4,拿掉任一都會復發。
        pg.pdf(path=str(out_pdf), format="A4", prefer_css_page_size=True, print_background=True,
               margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        br.close()
    return out_pdf


# ══════════════════════════════════════════════════════════════════════════════
#  xlsx 附件(§6:全 1001 檔強弱 + 34 板塊,深表頭/凍結/篩選/紅綠條件格式)
# ══════════════════════════════════════════════════════════════════════════════
def build_xlsx(state: dict, out_xlsx: Path) -> Path:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    BG, GOLD, UP, DN, MID = "0B0E14", "E3B93E", "E5484D", "2FB877", "F2F2F2"

    def head(ws, ncol):
        ws.row_dimensions[1].height = 26
        for c in range(1, ncol + 1):
            cell = ws.cell(row=1, column=c)
            cell.fill = PatternFill(start_color=BG, end_color=BG, fill_type="solid")
            cell.font = Font(color=GOLD, bold=True, size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")

    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "全市場強弱"
    ws.append(["代號", "名稱", "產業", "價", "漲跌%", "RSI", "強弱分", "趨勢"])
    head(ws, 8)
    wave = sorted(state.get("wave_top", []) or [], key=lambda r: r.get("score", 0), reverse=True)
    for r in wave:
        ws.append([str(r.get("code", "")), r.get("name", ""), r.get("industry", ""),
                   r.get("price"), r.get("chg"), r.get("rsi"), r.get("score"), r.get("st", "")])
    last = len(wave) + 1
    if last >= 2:
        ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:H{last}"
        for col, fmt in (("E", '0.00"%"'), ("F", "0.0"), ("G", "0.0")):
            for row in range(2, last + 1):
                ws[f"{col}{row}"].number_format = fmt
        # 漲跌%:台股 低綠→高紅
        ws.conditional_formatting.add(f"E2:E{last}", ColorScaleRule(
            start_type="min", start_color=DN, mid_type="num", mid_value=0, mid_color=MID,
            end_type="max", end_color=UP))
        ws.conditional_formatting.add(f"G2:G{last}", DataBarRule(start_type="min", end_type="max", color=GOLD))
    for i, w in enumerate([9, 16, 16, 8, 9, 8, 9, 8], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    ws2 = wb.create_sheet("板塊輪動")
    ws2.append(["板塊", "均漲跌%", "多方%", "分數", "檔數", "法人買", "領漲"])
    head(ws2, 7)
    secs = sorted(state.get("sectors", []) or [], key=lambda s: s.get("score", 0), reverse=True)
    for s in secs:
        ws2.append([s.get("name", ""), s.get("avg_chg"), s.get("bull_pct"), s.get("score"),
                    s.get("count"), s.get("inst_count"), s.get("leader", "")])
    l2 = len(secs) + 1
    if l2 >= 2:
        ws2.freeze_panes = "B2"; ws2.auto_filter.ref = f"A1:G{l2}"
        ws2.conditional_formatting.add(f"B2:B{l2}", ColorScaleRule(
            start_type="min", start_color=DN, mid_type="num", mid_value=0, mid_color=MID,
            end_type="max", end_color=UP))
        ws2.conditional_formatting.add(f"D2:D{l2}", DataBarRule(start_type="min", end_type="max", color=GOLD))
    for i, w in enumerate([18, 10, 8, 8, 7, 8, 20], start=1):
        ws2.column_dimensions[chr(64 + i)].width = w

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_xlsx))
    return out_xlsx


# ══════════════════════════════════════════════════════════════════════════════
#  訂閱名冊介面(不硬依賴 Task #6 產物;不存在回空 + log)
# ══════════════════════════════════════════════════════════════════════════════
_TIER_RANK = {"basic": 1, "full": 2, "full_annual": 3, "unknown": 0}


def load_send_list(tier: str | None = None) -> list[dict]:
    """讀 STUDIO/ecommerce_subscribers.json 的寄送名單,回 [{email,name,tier,platform}]。

    ── 介面契約(名冊 schema 為金流層真相,本端遷就)──────────────────────────
    名冊由 quant-service/webhook/subscribers.py 維護與導出。schema =
    **扁平 dict**  {email_lower: {email,name,tier,platform,status,...}};active 的定義是
    entry["status"] == "active"(不是布林 active 欄位)。tier 分層權重見 webhook 的
    TIER_RANK(basic<full<full_annual);tier="full" 只回完整版及以上。

    實作:優先直接呼叫 webhook.subscribers.export_active(path, tier)(單一事實來源,
    格式永遠對得上);import 不到(webhook 套件缺依賴等)才走等價的本地讀取——刻意**不硬依賴**
    webhook 套件,名冊檔不存在/壞一律回空 + log,絕不炸。交付端據此 dry_run,不在此寄信。
    """
    if not SUBSCRIBERS.exists():
        print(f"[weekly_v2] 訂閱名冊尚未產生({SUBSCRIBERS.name}),回空清單(交付端會 no-op)。")
        return []
    # 首選:直接用金流層的 export_active(單一事實來源)
    try:
        if str(QUANT) not in sys.path:
            sys.path.insert(0, str(QUANT))
        from webhook.subscribers import export_active  # noqa: E402
        return export_active(SUBSCRIBERS, tier=tier)
    except Exception as exc:  # noqa: BLE001
        print(f"[weekly_v2] 直呼 export_active 失敗({type(exc).__name__}),改用本地等價讀取。")
    # 後備:本地實作同一契約(扁平 dict + status==active + tier 權重)
    try:
        book = json.loads(SUBSCRIBERS.read_text(encoding="utf-8"))
    except Exception as exc:  # noqa: BLE001
        print(f"[weekly_v2] 訂閱名冊解析失敗:{exc};回空清單。")
        return []
    if not isinstance(book, dict):
        return []
    want = _TIER_RANK.get(tier, 0) if tier else 0
    out = []
    for e in book.values():
        if not isinstance(e, dict) or e.get("status") != "active":
            continue
        if tier and _TIER_RANK.get(e.get("tier", "unknown"), 0) < want:
            continue
        out.append({"email": e.get("email", ""), "name": e.get("name", ""),
                    "tier": e.get("tier", "unknown"), "platform": e.get("platform", "")})
    out.sort(key=lambda r: (-_TIER_RANK.get(r["tier"], 0), r["email"].lower()))
    return out


# ══════════════════════════════════════════════════════════════════════════════
#  主流程
# ══════════════════════════════════════════════════════════════════════════════
def generate_weekly(tier: str = "full", out_dir: Path | None = None,
                    requests_doc: dict | None = None, ensure_fn=None,
                    save_queue: bool = True) -> dict:
    out_dir = out_dir or OUT_DEFAULT
    state = load_state()
    valdoc = load_valuation_latest()
    chips_week = load_chips_week(5)
    checkup = load_checkup()
    adaptive = load_adaptive()
    nmap = build_name_map(state, checkup, adaptive)
    imap = build_industry_map(state)
    # 點播佇列:sec_S7 會就地更新 status/attempts,產完後存回(佇列空 → 全程 no-op)
    reqs = load_requests() if requests_doc is None else requests_doc

    sections = [
        sec_S1(state), sec_S2(state), sec_S3(state, nmap),
        sec_S4(state, chips_week, nmap), sec_S5(valdoc, nmap, imap),
        sec_S6(state, nmap), sec_S7(checkup, requests_doc=reqs, ensure_fn=ensure_fn, nmap=nmap),
        sec_S8(adaptive, nmap),
    ]
    sections = [gate_or_degrade(s) for s in sections]
    if save_queue and reqs.get("requests"):
        save_requests(reqs)

    d = state.get("date") or TODAY.isoformat()
    week_label = f"{d}(本週)"
    html = build_html(state, sections, week_label, tier)

    stamp = TODAY.isoformat()
    base = out_dir / f"weekly_{stamp}_{tier}"
    pdf = render_pdf(html, base.with_suffix(".pdf"))
    xlsx = build_xlsx(state, base.parent / f"weekly_{stamp}_全市場數據.xlsx")

    # provenance 存證(每個綁定數字 → 來源)
    # A3 修:只存「本 tier 真的有交付」的段落 —— 對齊 build_html 的可見性過濾。
    # v1 的 basic 存證檔含 71 筆 S7,但 basic 買家根本收不到 S7 → 稽核檔描述了未交付的內容。
    visible = [s for s in sections
               if tier == "full" or CFG.SECTION_TIERS.get(s["id"]) == "basic"]
    prov_records = []
    for s in visible:
        for rec in s["prov"].records:
            rec = dict(rec); rec["section"] = s["id"]; prov_records.append(rec)
    cover = sorted({r["section"] for r in prov_records})
    (base.parent / f"weekly_{stamp}_{tier}_provenance.json").write_text(
        json.dumps({"generated_at": stamp, "tier": tier,
                    "gate": "product_factory.Provenance.gate (reused, strict, fail-closed)",
                    "sections_covered": cover, "n_records": len(prov_records),
                    "records": prov_records}, ensure_ascii=False, indent=2), encoding="utf-8")

    status = [{"id": s["id"], "title": s["title"], "tier": CFG.SECTION_TIERS[s["id"]],
               "state": ("GATED" if s.get("gated_out") else "DEGRADED" if s.get("degraded") else "OK"),
               "units": len(s["units"])} for s in sections]
    return {"pdf": pdf, "xlsx": xlsx, "html": base.with_suffix(".html"),
            "tier": tier, "sections": status,
            "n_ok": sum(1 for x in status if x["state"] == "OK"),
            "send_list_n": len(load_send_list())}


def main() -> int:
    ap = argparse.ArgumentParser(description="台股全市場週報引擎 v2(旗艦)")
    ap.add_argument("--tier", choices=["basic", "full"], default="full")
    ap.add_argument("--out", default=str(OUT_DEFAULT))
    args = ap.parse_args()
    res = generate_weekly(tier=args.tier, out_dir=Path(args.out))
    print(f"\n[weekly_v2] tier={res['tier']}  →  {res['pdf']}")
    print(f"[weekly_v2] xlsx 附件 → {res['xlsx']}")
    print(f"[weekly_v2] 8 sections 狀態:")
    for s in res["sections"]:
        print(f"    {s['id']} [{s['tier']:5}] {s['state']:8} · {s['units']} unit · {s['title']}")
    print(f"[weekly_v2] OK {res['n_ok']}/8　訂閱名冊 {res['send_list_n']} 人(dry-run,不寄)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
