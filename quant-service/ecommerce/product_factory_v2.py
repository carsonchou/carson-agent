#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""product_factory_v2.py —【電商商品工廠 v2】把台股數據資產包成「值得掏錢」等級的一次性 SKU。

v1(product_factory.py,保留不動供對照)的成品是 reportlab 淺底 office 排版 + 裸傾印 xlsx,
被 Carson 判「不夠好」。v2 換渲染外殼、不動誠信骨:
  - 成品視覺 = 暗色數據卡 PDF(HTML+CSS → Chromium,復用 render_kit + report_theme.css,
    與旗艦週報同一套 token)+ 專業級 xlsx(深表頭/凍結/篩選/台股紅綠條件格式/data bar)。
  - 誠信原封搬入:Provenance 空池起步、小池嚴容差(0.25 絕對/0.5% 相對)、查無來源整個 SKU
    中止不出檔;所有文案「介紹≠推薦」、靜態回測快照標明非即時;數字全由 disp() 綁到來源。

SKU 線(對齊 REDESIGN_SPEC_business.md §3 / config.py):
  L0 磁鐵  M1 當沖適格清單 / M2 單檔旗艦體檢(台積電)         —— 免費、PDF
  L1 tripwire T1 定投追蹤模板 / T2 個股體檢單檔報告            —— zh、PDF(+T1 xlsx)
  L2 core   C1 全市場回測數據包 / C2 權值股體檢合輯            —— zh+en、xlsx+PDF

用法:
  python product_factory_v2.py --list
  python product_factory_v2.py                # 產全部(真實資料 → output/ecommerce_ready/v2/)
  python product_factory_v2.py --sku C1
驗證:python -m pytest quant-service/ecommerce/tests -q
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import statistics
import sys
from datetime import date
from pathlib import Path

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
except Exception:  # noqa: BLE001
    pass

HERE = Path(__file__).resolve().parent            # quant-service/ecommerce
QUANT = HERE.parent                               # quant-service
ROOT = QUANT.parent                               # repo root
TWDATA = ROOT / "twdata"
STUDIO = ROOT / "youtube_channel" / "STUDIO"
SCRIPTS = ROOT / "youtube_channel" / "scripts"
OUT_ROOT = QUANT / "output" / "ecommerce_ready" / "v2"

ADAPTIVE_CSV = TWDATA / "adaptive_per_stock.csv"
LONGSHORT_CSV = TWDATA / "longshort_per_stock.csv"
PERSTOCK_CSV = TWDATA / "per_stock_results.csv"
CHECKUP_FACTS = STUDIO / "stock_checkup_facts.json"

sys.path.insert(0, str(HERE))
sys.path.insert(0, str(SCRIPTS))
import config as CFG            # noqa: E402  ecommerce/config.py(共用定價/品牌)
import render_kit as RK         # noqa: E402  共用渲染工具箱
from product_factory import Provenance  # noqa: E402  復用 v1 誠信 Provenance(空池/嚴容差)

try:
    import fact_source_guard as FG  # noqa: E402
except Exception as exc:  # noqa: BLE001
    print(f"[pf2] 致命:無法載入 fact_source_guard(溯源守門)——{exc}")
    raise SystemExit(2)

TODAY = date.today().isoformat()
DISCLAIMER = ("本商品為歷史/當期數據體檢與教學工具,所有數字均為歷史統計,不預測未來、"
              "不構成投資建議、不喊單、不報明牌。投資有風險。「介紹」不等於「推薦」。")
DISCLAIMER_EN = ("This product is a historical-data health-check and educational tool. All figures "
                 "are historical statistics; nothing here predicts the future or is investment advice.")
SOURCES_CHECKUP = ["youtube_channel/STUDIO/stock_checkup_facts.json(體檢引擎:FinMind 財報/月營收/股利/估值 + Yahoo 含息還原價)"]
SOURCES_BACKTEST = ["twdata/adaptive_per_stock.csv、longshort_per_stock.csv、per_stock_results.csv(全市場回測靜態快照,2026-06-12 產)"]
SOURCES_DAYTRADE = ["twdata/daytrade_eligibility_*.json(證交所 TWSE OpenAPI 當日處置/注意股)"]
SOURCES_CHECKUP_EN = ["youtube_channel/STUDIO/stock_checkup_facts.json (health-check engine: FinMind "
                      "statements/monthly revenue/dividends/valuation + Yahoo dividend-adjusted prices)"]
SOURCES_BACKTEST_EN = ["twdata/adaptive_per_stock.csv, longshort_per_stock.csv, per_stock_results.csv "
                       "(full-market backtest static snapshot, generated 2026-06-12)"]

# ── C1 方法與限制揭露(MAJOR 9)────────────────────────────────────────────────
# ⚠️ 每一項都必須能回源碼查證,不得憑印象寫。本區逐項出處(2026-07-16 逐行讀碼確認):
#   tw_adaptive.py:AdaptiveParams(regime 門檻/子策略參數)、:119 backtest_adaptive(cost_model="tw_real")
#   strategy.py:COST_MODELS["tw_real"] = fee_buy .001425 / fee_sell .001425 / tax_sell .003 / slip_ticks 0
#   tw_data.py:get_universe() 讀 twstock.codes(=**現存**上市櫃)→ 倖存者偏誤為真
#   tw_adaptive.py:45-46 MIN_BARS=1000 / MIN_YEARS=5
# 「滑價未模擬」與「倖存者偏誤」是對我們不利的事實,但買家花 NT$990 有權知道 → 照實揭露。
METHOD_ZH = [
    ("「自適應」到底是什麼", "每根 K 線先判 regime 再切子策略(單一部位、純多、不放空):"
     "<b>ADX(14) > 18 且效率比 ER(20) > 0.26 → 判趨勢段</b>,走三重 SuperTrend 順勢跟蹤"
     "(基礎長度 10、倍數 2.5/6.0、係數 1.5/2.5/3.5,倉位 0.95,寬停損 11×ATR);"
     "<b>否則判盤整段</b>,走均值回歸(布林 30 期 2.5σ 或 RSI(14)<30 買進,回中軌或 RSI>50 賣出,"
     "倉位 0.60,破下軌再 1.5×ATR 停損,單筆最長持有 30 根)。"),
    ("成本假設(含什麼、不含什麼)", "已扣:買進手續費 <b>0.1425%</b>、賣出手續費 <b>0.1425%</b>、"
     "賣出證交稅 <b>0.3%</b>(即台股實際費率 tw_real)。"
     "<b>未模擬滑價</b>(slip_ticks=0)——這是對本數據包不利的事實,但你有權知道:"
     "實際下單成交價會比回測差,流動性差的個股差更多。"),
    ("倖存者偏誤(有,且方向對我們有利)", "選股池取自 twstock 的<b>現存</b>上市櫃清單,"
     "<b>已下市/已合併的公司不在裡面</b>。也就是說這 1770 檔全是「活到今天」的股票,"
     "整體績效因此被高估。看到正報酬佔比時請把這點折進去。"),
    ("其他前提", "資料已修正分割(個股反分割會污染回測);年化以 252 根/年計;"
     "每檔至少需 5 年、1000 根日 K 才納入(新股/資料太短者已排除)。"),
]
METHOD_EN = [
    ("What “adaptive” actually means", "Each bar picks a regime, then a sub-strategy (single position, "
     "long-only, no shorting): <b>ADX(14) &gt; 18 and efficiency ratio ER(20) &gt; 0.26 → trend regime</b>, "
     "trading a triple-SuperTrend trailing system (base length 10, multipliers 2.5/6.0, factors 1.5/2.5/3.5, "
     "exposure 0.95, wide stop 11×ATR); <b>otherwise range regime</b>, trading mean-reversion (buy on a "
     "30-period 2.5σ Bollinger break or RSI(14) &lt; 30; sell back at the mid-band or RSI &gt; 50; exposure 0.60, "
     "stop 1.5×ATR below the lower band, max hold 30 bars)."),
    ("Cost assumptions (what is and isn’t included)", "Deducted: buy commission <b>0.1425%</b>, sell "
     "commission <b>0.1425%</b>, sell transaction tax <b>0.3%</b> (the real Taiwan retail fee schedule). "
     "<b>Slippage is NOT modelled</b> (slip_ticks = 0). That fact cuts against this data pack, but you have "
     "a right to know it: real fills will be worse than the backtest, and materially worse in illiquid names."),
    ("Survivorship bias (present, and it flatters us)", "The universe is the list of <b>currently listed</b> "
     "Taiwan stocks, so <b>delisted and merged companies are absent</b>. All 1770 tickers here are survivors, "
     "which inflates aggregate performance. Discount the “% positive” figure accordingly."),
    ("Other assumptions", "Prices are split-adjusted (un-adjusted reverse splits corrupt backtests); "
     "annualisation uses 252 bars/year; a ticker needs at least 5 years and 1000 daily bars to be included "
     "(recent listings and short histories are excluded)."),
]
_NUM_RE = re.compile(r"-?\d+\.?\d*")


# ── 溯源:每個顯示的數字都經 disp() 綁定入池(gate 為回歸保險網)────────────────
def _pool(prov: Provenance, *items) -> None:
    for it in items:
        if it is None:
            continue
        if isinstance(it, (int, float)):
            v = float(it)
            prov.pool.add(abs(v)); prov.pool.add(abs(round(v, 1)))
        else:
            for m in _NUM_RE.findall(str(it)):
                try:
                    v = float(m)
                    prov.pool.add(abs(v)); prov.pool.add(abs(round(v, 1)))
                except ValueError:
                    pass


def disp(prov: Provenance, value, fmt: str = "{:.1f}", *, src: str = "", field: str = "") -> str:
    """格式化並把『顯示出來的數字』入池 + **留存證**(保證文字與池一致;gate 只會抓到漏綁的)。

    額外把 f"{s}%" 經 FG 抽取器入池——FG 的 %宣稱正則上限 4 位整數(\\d{1,4}),
    對 5 位數百分比(如 21051.3%)只會抽出末 4 位(1051.3);唯有把 FG 眼中的形式也入池,
    來源本就是真的大數字才不會被守門誤殺。此為對齊守門 tokenizer,非放水(值仍源自 value)。

    ⚠️ src/field(2026-07-16 修,phase3b B4):v2 初版全改走 disp()/_pool(),這兩者**只寫
    prov.pool、從不寫 prov.records** → 8/8 `_provenance.json` 的 records 恆為 []。數字本身沒
    造假(獨立重算 8,940 格零誤差)、gate 也真的會擋(靠 in-memory pool),但**持久化稽核檔是
    空的** → 外部稽核者拿到存證檔得到零資訊,違反 REDESIGN_SPEC_product.md:263 且是 line 20
    明列「不可退化」的 v1 baseline(v1 有 15 次 prov.num())。
    現在每個寫進成品的數字都必須帶 src(來源檔:欄位)+ field,存證檔才能逐筆回答「這個數字
    來自哪個檔的哪個欄位」。沿用 weekly_report_v2._bind 的 record 形狀(field/value/text/source)。
    """
    s = fmt.format(value)
    _pool(prov, s)
    for c in FG.extract_claims(s + "%"):
        prov.pool.add(abs(c["value"])); prov.pool.add(abs(round(c["value"], 1)))
    if src and field:
        try:
            v = float(value)
        except (TypeError, ValueError):
            v = None
        prov.records.append({"field": field, "value": v, "text": s,
                             "source": _complete_src(src, field)})
    return s


def _complete_src(src: str, field: str) -> str:
    """把 field 的葉欄位補進 source,讓 source 成為**完整、可機械解析**的指標。

    為什麼需要:第一版 source 只指到 `...data`,葉欄位藏在 field(如 `2330.long_horizon.cagr×100`)
    → 外部稽核者得先猜出「field 去掉 code/base 前綴才是欄位路徑」這個內規才回查得到。實測自寫
    稽核腳本 1021 筆裡有 941 筆解不開 —— 存證檔「有內容」但不可機械回查,等於 B4 只修一半。
    field 慣例固定是 `{code}.{base}.{leaf...}`,故可穩定補齊:
      source `…:results.checkup_long_horizon__2330.data` + field `2330.long_horizon.cagr×100`
      → `…:results.checkup_long_horizon__2330.data.cagr`(×100 這個 transform 仍記在 field)。
    巢狀葉(stock_allin.total_return)一併保留;含括號的說明型 field(逐字引用)不動。
    """
    if not src.endswith(".data"):
        return src
    parts = field.split(".", 2)
    if len(parts) < 3:
        return src
    rest = parts[2].replace("×100", "").strip()
    if not rest or "(" in rest or "[" in rest:
        return src
    return f"{src}.{rest}"


def _bind(prov: Provenance, src: str, **fields) -> None:
    """入池 + 留存證,但不回傳顯示字串(給『有入池但不直接印出』的數字用,如 sparkline 序列端點)。
    與 weekly_report_v2._bind 同一形狀:數值型直接存 value,稽核可程式化比對而不只靠文字。"""
    for name, v in fields.items():
        if v is None:
            continue
        _pool(prov, v)
        s = _complete_src(src, name)
        if isinstance(v, (int, float)) and not isinstance(v, bool):
            prov.records.append({"field": name, "value": float(v), "text": None, "source": s})
        else:
            prov.records.append({"field": name, "value": None, "text": str(v), "source": s})


def _ck_src(ff: dict, base: str, field: str = "") -> str:
    """體檢事實的可回查來源字串:`檔案:results.<fact_key>.data.<欄位>`。
    fact_key 由 facts_for() 附掛(`_key`),稽核者可直接拿去 grep 原始事實庫。"""
    f = ff.get(base)
    if isinstance(f, list):
        f = f[0] if f else None
    key = (f or {}).get("_key") or f"checkup_{base}"
    tail = f".{field}" if field else ""
    return f"youtube_channel/STUDIO/stock_checkup_facts.json:results.{key}.data{tail}"


def _pool_src(prov: Provenance, s: str, src: str = "", field: str = "") -> None:
    """逐字引用的來源字串(體檢 claim/summary):既入原始數字,也入 FG 抽取器眼中的 %宣稱,
    對齊守門對 5 位數百分比的 4 位截斷,避免真來源數字被誤判查無來源。
    給了 src/field 就一併留存證(逐字引用的整句 claim 存 text,稽核可回查原句)。"""
    _pool(prov, s)
    for c in FG.extract_claims(s):
        prov.pool.add(abs(c["value"])); prov.pool.add(abs(round(c["value"], 1)))
    if src and field and s:
        prov.records.append({"field": field, "value": None, "text": str(s), "source": src})


def cls_updn(v) -> str:
    try:
        return "pos" if float(v) > 0 else "neg" if float(v) < 0 else ""
    except (TypeError, ValueError):
        return ""


# ── 英文績效宣稱抽取器(補 FG 的中文盲區)──────────────────────────────────────
# 🔴 2026-07-16 修 phase3b B3 時實測發現的**結構性破口**(報告未提,但不補就會踩):
#   FG.extract_claims 要求命中點所在子句含 PERF_CTX 才算宣稱,而 PERF_CTX 全是中文詞
#   (報酬/勝率/回撤…);且 _clause() 只以「。！？\n」切句(英文句點不切,否則 25.1 會被切壞)。
#   ⇒ 對**純英文**文字實測回 0 claim:`extract_claims("Win rate 99.9%") == []`。
#   今天 C1_en/C2_en 之所以有被守門到,只是因為它們還殘留中文(C2_en 37.5% 是中文)——
#   中文詞把整段撐成一個含 PERF_CTX 的子句,英文數字才連帶被抽出來。
#   ⇒ **把英文版真的英文化(B3),等於順手把這兩支的 gate 關掉**,「修 B3」變成「放寬 gate」。
#   故在地化必須與本抽取器同批上線。比對器仍沿用 FG._sourced_strict(不自造容差),
#   只補 FG 沒有的「英文抽取」這一段;對所有 SKU 都跑(中文版夾雜英文也照樣被抽),
#   守門覆蓋率只增不減。
_EN_PERF_CTX = (
    "return", "drawdown", "cagr", "annualis", "annualiz", "win rate", "sharpe", "calmar",
    "yield", "margin", "profit", "loss", "gain", "volatil", "median", "percentile",
    "performance", "underwater", "halv", "fell", "rose", "growth", "positive", "dividend",
)
# 誠實揭露語境(對應 FG.HEDGE 的英文面):同句有這些詞 = 不是本商品的事實斷言
_EN_HEDGE = (
    "for illustration", "illustrative", "hypothetical", "for example", "e.g.", "sample row",
    "not a guarantee", "no guarantee", "does not predict", "not investment advice",
    "past performance", "educational", "placeholder",
)
_RX_PCT_EN = re.compile(r"(\d{1,4}(?:\.\d+)?)\s*%")          # 與 FG._RX_PCT_ARABIC 同款(不吃負號)
_RX_EN_SENT = re.compile(r"(?<=[.;:!?])\s+|\n+|。|！|？")     # 英文句界:標點**後接空白**才切(不切 25.1)


def _extract_claims_en(text: str) -> list[dict]:
    """抽出英文『績效類百分比宣稱』:{value, raw, clause}。與 FG.extract_claims 同精神——
    必須落在績效語境的句子裡才算宣稱(否則年份/頁碼/價格全被誤抓)。"""
    claims: list[dict] = []
    for seg in _RX_EN_SENT.split(text or ""):
        if not seg:
            continue
        low = seg.lower()
        if not any(w in low for w in _EN_PERF_CTX):
            continue
        for m in _RX_PCT_EN.finditer(seg):
            try:
                claims.append({"value": float(m.group(1)), "raw": m.group(0),
                               "clause": seg.strip()[:90]})
            except ValueError:
                pass
    return claims


def _gate_en(prov: Provenance, text: str) -> list[dict]:
    """英文路徑守門:抽取器自造(FG 無英文),比對器沿用 FG._sourced_strict(嚴容差,不放水)。"""
    bad = []
    for c in _extract_claims_en(text):
        low = c["clause"].lower()
        if any(h in low for h in _EN_HEDGE):
            continue
        if not FG._sourced_strict(c["value"], prov.pool):
            bad.append(c)
    return bad


# ── 資料載入(全部 fail-safe)──────────────────────────────────────────────────
def load_checkup() -> dict:
    try:
        return json.loads(CHECKUP_FACTS.read_text(encoding="utf-8"))
    except Exception:  # noqa: BLE001
        return {"results": {}, "by_code": {}}


def elig_trusted(d: dict | None) -> bool:
    """這份當沖適格清單可不可信(= 能不能拿來說「處置股 N 檔」)。

    🔴 這是 M1 的安全線(2026-07-17 修):舊版 daytrade_eligibility.refresh() 是 **fail-OPEN**
    —— TWSE 掛掉時 `_fetch_json` 吞成 `[]`,refresh() 照樣寫出 `{"disposition": []}`,
    與「今天真的 0 檔處置股」**在檔案層面完全無法區分**。M1 的賣點是**報單前防呆**,
    印出「處置股 0 檔」等於告訴當沖客「這檔可以沖」——安全清單上的 fail-open,有實害。
    (上一版加的「快照日 + 距今 N 天」警語擋不住這個:今天日期、內容全空的檔看起來既新鮮又安全。)

    契約與 data_hunter/daytrade_eligibility.is_trusted() 一致:
      · 新格式有 `ok` 欄位 → 直接看它;
      · 舊格式(2026-07-17 前)沒有 ok,而舊 refresh 抓不到也會寫空檔 → 全空一律當**不可信**,
        有資料才當可信(那顯然是抓成功的)。
    刻意不 import data_hunter 那支(跨專案硬依賴),但兩邊都有測試釘住同一套判定。
    """
    if not isinstance(d, dict) or not d:
        return False
    if "ok" in d:
        return bool(d["ok"])
    return bool(d.get("disposition") or d.get("attention"))


def load_daytrade_elig() -> dict:
    files = sorted(TWDATA.glob("daytrade_eligibility_*.json"))
    if not files:
        return {}
    try:
        d = json.loads(files[-1].read_text(encoding="utf-8"))
        d["_file"] = files[-1].name
        return d
    except Exception:  # noqa: BLE001
        return {}


def _load_csv(path: Path, floats: tuple = (), ints: tuple = ()) -> list[dict]:
    rows = []
    if not path.exists():
        return rows
    with open(path, encoding="utf-8-sig") as f:
        for r in csv.DictReader(f):
            ok = True
            for k in floats:
                try:
                    r[k] = float(r[k])
                except (KeyError, ValueError, TypeError):
                    ok = False
            for k in ints:
                try:
                    r[k] = int(float(r[k]))
                except (KeyError, ValueError, TypeError):
                    ok = False
            if ok:
                rows.append(r)
    return rows


def facts_for(checkup: dict, code: str) -> dict:
    """回 code 的 {suffix: fact_dict}(只收通過守門的 fact;crash 收成 list)。

    每個 fact 複製一份並附掛 `_key`(原始 fact_key),供 _ck_src() 產生可回查的存證來源字串;
    複製而不原地改,避免污染呼叫端共用的 checkup dict。

    ⚠️ `len()` 即該檔的**實際體檢維度數**(crash 多筆收斂成 1 個 base)——封面/listing 的維度
    宣稱一律取自這裡,不得硬編(phase3b B2:硬寫「11 項」對國泰金10/00878=7/國巨6 三檔不實)。
    """
    out: dict = {}
    for key, f in (checkup.get("results", {}) or {}).items():
        if not key.endswith(f"__{code}") and f"__{code}__" not in key:
            continue
        if not (isinstance(f, dict) and f.get("source") and (f.get("claim") or f.get("summary"))
                and f.get("data") not in (None, "", [], {})):
            continue
        base = key[len("checkup_"):].split("__")[0] if key.startswith("checkup_") else key
        f = dict(f)
        f["_key"] = key
        if base == "crash":
            out.setdefault("crash", []).append(f)
        else:
            out[base] = f
    return out


def n_dims(ff: dict) -> int:
    """該檔實際體檢維度數(= facts_for 的 base 數)。封面/listing 的「每檔維度 N 項」唯一來源。"""
    return len(ff)


# ══════════════════════════════════════════════════════════════════════════════
#  個股體檢卡(M2/T2/C2 共用)—— 暗色數據卡:KPI + 估值位階條 + sparkline + 三買法對照
# ══════════════════════════════════════════════════════════════════════════════
# 英文版怎麼來(phase3b B3):**不翻譯、不過 LLM**。資料層的 claim/summary 是中文逐字句,
# 直接機翻既有成本又有走樣風險(數字被改寫=誠信事故)。作法是**從同一份 data 欄位確定性重建
# 英文句**——數字仍來自欄位、由 disp() 綁來源入池,英文只是另一套模板。中文版維持逐字引用
# 資料層原句(既有行為不動)。兩邊數字同源,不可能出現「中英版數字不一致」。
_CRASH_WIN_EN = {
    "2008金融海嘯": "2008 global financial crisis",
    "2020新冠崩盤": "2020 COVID-19 crash",
    "2022台股熊市": "2022 Taiwan bear market",
}


def _crash_win_en(d: dict) -> str:
    """崩盤視窗名的英文。未知視窗**不猜**,退回以 window_start 年份陳述事實(不編造事件名)。"""
    w = str(d.get("window") or "")
    if w in _CRASH_WIN_EN:
        return _CRASH_WIN_EN[w]
    yr = str(d.get("window_start") or "")[:4]
    return f"the {yr} drawdown window" if yr.isdigit() else "this drawdown window"


_CARD_L = {
    "zh": {
        "head": "{name}（{code}）體檢", "badge": "含息還原",
        "cagr": "年化報酬", "mdd": "最大回撤", "uw": "最長套牢", "uw_unit": " 年", "dy": "殖利率",
        "val_t": "估值位階(近10年,只標位置)",
        "seg_g": "偏低區", "seg_a": "中段", "seg_r": "偏高區",
        "tw_t": "近{y}年 三種買法對照(含息還原)",
        "allin": "單筆 All-in", "dca": "每月定投", "bench": "同期 0050",
        "sp_rev": "年營收(億)", "sp_eps": "年 EPS(元)", "sp_gm": "單季毛利率(%)",
        "t_extremes": "年度極值", "t_halv": "腰斬史", "t_div": "股利", "t_crash": "崩盤韌性",
    },
    "en": {
        "head": "{code} · {name} — Health Check", "badge": "Dividend-adjusted",
        "cagr": "CAGR", "mdd": "Max drawdown", "uw": "Longest underwater", "uw_unit": " yrs",
        "dy": "Dividend yield",
        "val_t": "Valuation position (10-year, position only)",
        "seg_g": "lower zone", "seg_a": "mid zone", "seg_r": "upper zone",
        "tw_t": "Past {y} years — three ways to buy (dividend-adjusted)",
        "allin": "Lump sum (all-in)", "dca": "Monthly DCA", "bench": "0050, same period",
        "sp_rev": "Annual revenue (NT$100M)", "sp_eps": "Annual EPS (NT$)",
        "sp_gm": "Quarterly gross margin (%)",
        "t_extremes": "Yearly extremes", "t_halv": "Halvings", "t_div": "Dividends",
        "t_crash": "Crash resilience",
    },
}


def _en_lines(prov: Provenance, ff: dict, code: str) -> list[tuple[str, str]]:
    """英文版的極值/腰斬/股利/崩盤敘述:全部由 data 欄位重建(數字經 disp 綁來源)。
    回 [(tag, sentence)];缺欄位就跳過該句(不硬掰)。"""
    L = _CARD_L["en"]
    out: list[tuple[str, str]] = []

    ae = (ff.get("annual_extremes") or {}).get("data") or {}
    wy, by = ae.get("worst_year") or {}, ae.get("best_year") or {}
    if wy.get("year") and by.get("year"):
        src = _ck_src(ff, "annual_extremes")
        n = disp(prov, ae.get("n_full_years", 0), "{:.0f}", src=src, field=f"{code}.annual_extremes.n_full_years")
        wr = disp(prov, wy.get("return", 0) * 100, "{:+.1f}", src=src, field=f"{code}.annual_extremes.worst_year.return×100")
        br = disp(prov, by.get("return", 0) * 100, "{:+.1f}", src=src, field=f"{code}.annual_extremes.best_year.return×100")
        _bind(prov, src, **{f"{code}.annual_extremes.worst_year.year": wy.get("year"),
                            f"{code}.annual_extremes.best_year.year": by.get("year")})
        out.append((L["t_extremes"], f'Across {n} full calendar years: the worst year was {wy["year"]} '
                                     f'at a {wr}% return; the best was {by["year"]} at {br}%.'))

    hv = (ff.get("halvings") or {}).get("data") or {}
    if hv.get("n_halvings") is not None:
        src = _ck_src(ff, "halvings")
        nh = disp(prov, hv.get("n_halvings", 0), "{:.0f}", src=src, field=f"{code}.halvings.n_halvings")
        yr = disp(prov, hv.get("years", 0), "{:.1f}", src=src, field=f"{code}.halvings.years")
        ev = (hv.get("events") or [{}])[0]
        tail = ""
        if ev.get("from_peak_date") and ev.get("halved_date"):
            _bind(prov, src, **{f"{code}.halvings.events[0].from_peak_date": ev["from_peak_date"],
                                f"{code}.halvings.events[0].halved_date": ev["halved_date"]})
            tail = f' First one: from the {ev["from_peak_date"]} peak down to a halving by {ev["halved_date"]}.'
        out.append((L["t_halv"], f'Over {yr} years of data, the price halved (down 50% or more from a peak) '
                                 f'{nh} time(s).{tail}'))

    dh = (ff.get("dividend_history") or {}).get("data") or {}
    if dh.get("n_years_data") is not None:
        src = _ck_src(ff, "dividend_history")
        ny = disp(prov, dh.get("n_years_data", 0), "{:.0f}", src=src, field=f"{code}.dividend_history.n_years_data")
        cy = disp(prov, dh.get("consecutive_years", 0), "{:.0f}", src=src, field=f"{code}.dividend_history.consecutive_years")
        # ⚠️ avg_yield_5y 是**比率**(2330=0.0197 → 2.0%),與 valuation_position.latest_dividend_yield
        # (**已是百分比**,2330=0.91 → 0.9%)單位相反。同一份事實庫兩種慣例,弄反就是 ×100 誠信事故。
        # 佐證:2603 avg_yield_5y=0.276 而資料層中文句寫「近5年平均現金殖利率約 27.6%」。
        ay = dh.get("avg_yield_5y")
        tail = ""
        if ay is not None:
            s = disp(prov, ay * 100, "{:.1f}", src=src, field=f"{code}.dividend_history.avg_yield_5y×100")
            tail = f' 5-year average cash dividend yield about {s}%.'
        out.append((L["t_div"], f'{ny} years of dividend data; {cy} consecutive paying years '
                                f'(before the most recent break).{tail}'))

    for f in (ff.get("crash") or [])[:3]:
        d = f.get("data") or {}
        if d.get("trough_drawdown") is None:
            continue
        src = _ck_src({"crash": [f]}, "crash")
        dd = disp(prov, d["trough_drawdown"] * 100, "{:+.1f}", src=src, field=f"{code}.crash.trough_drawdown×100")
        htr = disp(prov, (d.get("hold_through_return") or 0) * 100, "{:+.1f}", src=src,
                   field=f"{code}.crash.hold_through_return×100")
        _bind(prov, src, **{f"{code}.crash.peak_date": d.get("peak_date"),
                            f"{code}.crash.trough_date": d.get("trough_date"),
                            f"{code}.crash.latest_date": d.get("latest_date")})
        out.append((L["t_crash"], f'{_crash_win_en(d)}: from the {d.get("peak_date")} peak to the '
                                  f'{d.get("trough_date")} trough it fell {dd}%; holding through to '
                                  f'{d.get("latest_date")} returned {htr}%.'))
    return out


def checkup_card(prov: Provenance, code: str, name: str, ff: dict, lang: str = "zh") -> str:
    """回一個 .unit HTML;數字全經 disp() 入池+留存證。ff = facts_for() 的結果。"""
    L = _CARD_L.get(lang, _CARD_L["zh"])
    parts = [RK.sec_head(L["head"].format(name=name, code=code), L["badge"])]

    # KPI:年化 / 最大回撤 / 最長套牢 / 殖利率
    kpis = []
    lh = ff.get("long_horizon", {}).get("data") if ff.get("long_horizon") else None
    if lh:
        src = _ck_src(ff, "long_horizon")
        cagr = lh.get("cagr", 0) * 100
        mdd = lh.get("max_drawdown", 0) * 100
        kpis.append((L["cagr"], f'<span class="{cls_updn(cagr)}">'
                                f'{disp(prov, cagr, src=src, field=f"{code}.long_horizon.cagr×100")}%</span>'))
        kpis.append((L["mdd"], f'<span class="neg">'
                               f'{disp(prov, mdd, src=src, field=f"{code}.long_horizon.max_drawdown×100")}%</span>'))
        _bind(prov, src, **{f"{code}.long_horizon.years": lh.get("years"),
                            f"{code}.long_horizon.total_return": lh.get("total_return"),
                            f"{code}.long_horizon.calmar": lh.get("calmar"),
                            f"{code}.long_horizon.start": lh.get("start"),
                            f"{code}.long_horizon.end": lh.get("end")})
    uw = ff.get("underwater", {}).get("data") if ff.get("underwater") else None
    if uw:
        src = _ck_src(ff, "underwater")
        kpis.append((L["uw"], f'{disp(prov, uw.get("max_underwater_years", 0), src=src, field=f"{code}.underwater.max_underwater_years")}'
                              f'<small>{L["uw_unit"]}</small>'))
        _bind(prov, src, **{f"{code}.underwater.max_underwater_days": uw.get("max_underwater_days")})
    vp = ff.get("valuation_position", {}).get("data") if ff.get("valuation_position") else None
    if vp:
        src = _ck_src(ff, "valuation_position")
        # 已是百分比,勿再×100(跨股量級反證:2603 陽明 8.23 若當比率=823% 荒謬)
        kpis.append((L["dy"], f'{disp(prov, vp.get("latest_dividend_yield", 0), src=src, field=f"{code}.valuation_position.latest_dividend_yield")}'
                              f'<small>%</small>'))
    if kpis:
        parts.append(RK.kpi_row(kpis))

    # 長期報酬敘述:中文逐字引用來源句;英文由 data 欄位重建
    if ff.get("long_horizon") and lh:
        src = _ck_src(ff, "long_horizon")
        if lang == "zh":
            s = ff["long_horizon"].get("summary") or ff["long_horizon"].get("claim") or ""
            _pool_src(prov, s, src=src, field=f"{code}.long_horizon.summary(逐字引用)")
        else:
            tr = disp(prov, (lh.get("total_return") or 0) * 100, "{:,.1f}", src=src,
                      field=f"{code}.long_horizon.total_return×100")
            cg = disp(prov, (lh.get("cagr") or 0) * 100, "{:.1f}", src=src, field=f"{code}.long_horizon.cagr×100")
            md = disp(prov, (lh.get("max_drawdown") or 0) * 100, "{:.1f}", src=src,
                      field=f"{code}.long_horizon.max_drawdown×100")
            cal = disp(prov, lh.get("calmar") or 0, "{:.2f}", src=src, field=f"{code}.long_horizon.calmar")
            yy = disp(prov, lh.get("years") or 0, "{:.1f}", src=src, field=f"{code}.long_horizon.years")
            s = (f'Dividend-adjusted total return {tr}% over {yy} years '
                 f'({lh.get("start")} → {lh.get("end")}) — CAGR {cg}%, max drawdown {md}%, Calmar {cal}.')
        parts.append(f'<div class="lead">{RK.esc(s)}</div>')

    two = []
    # 估值位階條(位置陳述,非買賣)
    if vp:
        src = _ck_src(ff, "valuation_position")
        pctl = vp.get("percentile_rank", 50)
        per = vp.get("latest_per"); p25 = vp.get("p25"); med = vp.get("median"); p75 = vp.get("p75")
        lc = RK.light_class(pctl)
        seg = L["seg_g"] if lc == "g" else (L["seg_a"] if lc == "a" else L["seg_r"])
        d_per = disp(prov, per, src=src, field=f"{code}.valuation_position.latest_per")
        d_pctl = disp(prov, pctl, "{:.0f}", src=src, field=f"{code}.valuation_position.percentile_rank")
        d_p25 = disp(prov, p25, src=src, field=f"{code}.valuation_position.p25")
        d_med = disp(prov, med, src=src, field=f"{code}.valuation_position.median")
        d_p75 = disp(prov, p75, src=src, field=f"{code}.valuation_position.p75")
        if lang == "zh":
            body = (f'<span class="dot {lc}"></span>本益比 <b>{d_per}</b> 倍,'
                    f'位於自身近10年第 <b>{d_pctl}</b> 百分位（{seg}）{RK.pos_bar(pctl)}<br>'
                    f'區間 P25 <b>{d_p25}</b> ／ 中位 <b>{d_med}</b> ／ '
                    f'P75 <b>{d_p75}</b> 倍　<span style="color:var(--tx3)">·不判斷貴賤</span>')
        else:
            body = (f'<span class="dot {lc}"></span>P/E <b>{d_per}</b>×, sitting at the '
                    f'<b>{d_pctl}</b>th percentile of its own 10-year range ({seg}) {RK.pos_bar(pctl)}<br>'
                    f'Range: P25 <b>{d_p25}</b> / median <b>{d_med}</b> / P75 <b>{d_p75}</b>× '
                    f'<span style="color:var(--tx3)">· position only, no view on cheap or expensive</span>')
        two.append(f'<div class="block"><div class="bt">{RK.esc(L["val_t"])}</div>'
                   f'<div class="metricrow">{body}</div></div>')

    # 三種買法對照(All-in vs 定投 vs 0050)
    tw = ff.get("three_way", {}).get("data") if ff.get("three_way") else None
    if tw and tw.get("stock_allin") and tw.get("bench"):
        src = _ck_src(ff, "three_way")
        a = tw["stock_allin"]["total_return"] * 100
        dca = tw["stock_dca"]["total_return"] * 100
        b = tw["bench"]["total_return"] * 100
        mx = max(a, dca, b) or 1
        yrs = tw.get("years", 10)
        _bind(prov, src, **{f"{code}.three_way.start": tw.get("start"),
                            f"{code}.three_way.bench.ticker": (tw.get("bench") or {}).get("ticker")})
        rows = [
            (L["allin"], a, mx, False, disp(prov, a, "{:+.1f}", src=src,
                                            field=f"{code}.three_way.stock_allin.total_return×100") + "%"),
            (L["dca"], dca, mx, False, disp(prov, dca, "{:+.1f}", src=src,
                                            field=f"{code}.three_way.stock_dca.total_return×100") + "%"),
            (L["bench"], b, mx, False, disp(prov, b, "{:+.1f}", src=src,
                                            field=f"{code}.three_way.bench.total_return×100") + "%"),
        ]
        yy = disp(prov, yrs, "{:.0f}", src=src, field=f"{code}.three_way.years")
        two.append(f'<div class="block"><div class="bt">{RK.esc(L["tw_t"].format(y=yy))}</div>'
                   f'{RK.cmp_bars(rows)}</div>')
    if two:
        parts.append(f'<div class="two">{"".join(two)}</div>')

    # sparkline:營收 / EPS / 毛利率
    sparks = []
    rev = ff.get("revenue_trend", {}).get("data") if ff.get("revenue_trend") else None
    if rev and rev.get("series"):
        ser = [x["revenue"] / 1e8 for x in rev["series"] if isinstance(x.get("revenue"), (int, float))]
        if len(ser) >= 2:
            src = _ck_src(ff, "revenue_trend", "series[*].revenue÷1e8")
            _bind(prov, src, **{f"{code}.revenue_trend.series[0]÷1e8": ser[0]})
            latest = disp(prov, ser[-1], "{:,.0f}", src=src, field=f"{code}.revenue_trend.series[-1]÷1e8")
            sparks.append((L["sp_rev"], RK.sparkline(ser),
                           f'{latest} 億' if lang == "zh" else latest))
    eps = ff.get("eps_trend", {}).get("data") if ff.get("eps_trend") else None
    if eps and eps.get("series"):
        ser = [x["eps"] for x in eps["series"] if isinstance(x.get("eps"), (int, float))]
        if len(ser) >= 2:
            src = _ck_src(ff, "eps_trend", "series[*].eps")
            _bind(prov, src, **{f"{code}.eps_trend.series[0]": ser[0]})
            latest = disp(prov, ser[-1], src=src, field=f"{code}.eps_trend.series[-1]")
            sparks.append((L["sp_eps"], RK.sparkline(ser),
                           f'{latest} 元' if lang == "zh" else latest))
    gm = ff.get("gross_margin", {}).get("data") if ff.get("gross_margin") else None
    if gm and gm.get("series"):
        ser = [x["gross_margin"] for x in gm["series"] if isinstance(x.get("gross_margin"), (int, float))]
        if len(ser) >= 2:
            src = _ck_src(ff, "gross_margin", "series[*].gross_margin")
            _bind(prov, src, **{f"{code}.gross_margin.series[0]": ser[0]})
            latest = disp(prov, ser[-1], src=src, field=f"{code}.gross_margin.series[-1]")
            sparks.append((L["sp_gm"], RK.sparkline(ser), f'{latest}%'))
    if sparks:
        cells = "".join(
            f'<div class="block" style="flex:1"><div class="bt">{RK.esc(lbl)}</div>'
            f'<div style="display:flex;align-items:center;justify-content:space-between;gap:8px">'
            f'{svg}<b style="font-size:15px">{RK.esc(latest)}</b></div></div>'
            for lbl, svg, latest in sparks)
        parts.append(f'<div style="display:flex;gap:9px;margin-top:9px">{cells}</div>')

    # 韌性/極值/股利/腰斬:中文逐字引用既有 claim;英文由 data 欄位重建
    lines = []
    if lang == "zh":
        for suffix, tag in (("annual_extremes", L["t_extremes"]), ("halvings", L["t_halv"]),
                            ("dividend_history", L["t_div"])):
            f = ff.get(suffix)
            if f:
                s = f.get("summary") or f.get("claim") or ""
                _pool_src(prov, s, src=_ck_src(ff, suffix), field=f"{code}.{suffix}.summary(逐字引用)")
                lines.append(f'<b>{tag}</b>　{RK.esc(s)}')
        for f in ff.get("crash", [])[:3]:
            s = f.get("summary") or f.get("claim") or ""
            _pool_src(prov, s, src=_ck_src({"crash": [f]}, "crash"), field=f"{code}.crash.summary(逐字引用)")
            lines.append(f'<b>{L["t_crash"]}</b>　{RK.esc(s)}')
    else:
        # 英文版用一般空白,不用全形空白「　」(U+3000)——那是中文排版字元,夾在英文句裡
        # 對英文買家就是一個看不懂的寬洞(也會被中文殘留掃描算成 CJK)。
        for tag, s in _en_lines(prov, ff, code):
            lines.append(f'<b>{tag}</b> &nbsp;{RK.esc(s)}')
    if lines:
        parts.append(f'<div class="block" style="margin-top:9px"><div class="metricrow">'
                     + "<br>".join(lines) + "</div></div>")

    return f'<div class="unit">{"".join(parts)}</div>'


# ══════════════════════════════════════════════════════════════════════════════
#  SKU 建造器
# ══════════════════════════════════════════════════════════════════════════════
# listing 免責偵測詞(中/英各一組)。英文版舊版 C2_en description **完全無免責**、
# C1_en 只有一句 "Historical stats only.",但兩者 seo_note 都自稱「含免責」——
# **存證欄位自己在說謊**(phase3b MAJOR 10)。修法不是把字加回去就好,而是讓 seo_note
# **由事實算出來**:偵測不到免責就照實寫「⚠️ 無免責」,並由測試把付費 listing 釘成必須有。
_DISC_MARK = {
    "zh": ("非投資建議", "不構成投資建議", "介紹不等於推薦", "介紹≠推薦", "教學"),
    "en": ("not investment advice", "does not predict", "not a recommendation",
           "historical statistics only", "educational"),
}


def _has_disclaimer(desc: str, lang: str) -> bool:
    low = (desc or "").lower()
    return any(m.lower() in low for m in _DISC_MARK.get(lang, _DISC_MARK["zh"]))


def _listing(sku_id: str, lang: str, title: str, desc: str, tags: list[str],
             price_ntd, price_usd, platform: str) -> dict:
    ban = ("穩賺", "必賺", "保證", "最強", "翻倍", "暴賺", "神準", "包贏", "best", "guaranteed", "profit")
    clean, seen = [], set()
    for t in tags:
        t = t.strip()
        if not t or t.lower() in seen or any(b in t.lower() for b in ban):
            continue
        seen.add(t.lower()); clean.append(t[:20])
        if len(clean) >= 13:
            break
    price = {"NTD": price_ntd} if lang == "zh" else {"USD": price_usd}
    has_disc = _has_disclaimer(desc, lang)
    return {"sku": sku_id, "lang": lang, "platform": platform, "title": title[:140],
            "description": desc.strip(), "tags": clean, "price": price,
            "disclaimer_present": has_disc,
            "seo_note": ("標題相關詞前置;tag 長尾去主觀詞;無捏造轉換率;"
                         + ("已含免責。" if has_disc else "⚠️ 偵測不到免責字樣——上架前必須補。"))}


def build_M1(ctx) -> list[dict]:
    """L0 磁鐵:當沖適格清單(zh,PDF)。資料=最新 daytrade_eligibility。"""
    dt = ctx["daytrade"]
    if not dt:
        return []
    # 🔴 fail-CLOSED:清單不可信就**整支 SKU 不出**,絕不印「處置股 0 檔」。
    # M1 的賣點就是報單前防呆;給錯的安全訊息比不給更糟(「0 檔」= 告訴人家這檔可以沖)。
    # 這裡選擇不出檔而非出一份警語 PDF,因為這是**免費磁鐵**——沒有人在等它,
    # 漏發一天零損失;發一份「查無資料」的空殼反而砸品牌。產線 log 會說明原因。
    if not elig_trusted(dt):
        print("  🔴 M1 跳過:當沖適格清單不可信(TWSE 抓取失敗或舊格式空檔)——"
              "不出檔優於印出「處置股 0 檔」誤導當沖客。"
              f"（檔案 {dt.get('_file', '?')}，ok={dt.get('ok')}，reason={dt.get('reason')}）")
        return []
    prov = Provenance()
    b = CFG.BRAND
    src_f = f'twdata/{dt.get("_file", "daytrade_eligibility_*.json")}'
    disp_list = dt.get("disposition", []) or []
    attn = dt.get("attention", []) or []
    # ⚠️ `updated` 是**完整 ISO 時戳**(如 2026-07-13T09:05:36),不是純日期。
    # 第一版直接丟給 date.fromisoformat() → 在 3.9 會 ValueError → 被 except 吞掉變 stale_days=0
    # → 過期警語靜默不顯示(「修好了」其實沒生效)。取前 10 碼才是日期;顯示也只給日期,
    # 不把 09:05:36 這種機器時戳丟到買家臉上。
    updated_raw = str(dt.get("updated", TODAY))
    updated = updated_raw[:10]
    n_disp = disp(prov, len(disp_list), "{:.0f}", src=f"{src_f}:disposition", field="disposition.count")
    n_attn = disp(prov, len(attn), "{:.0f}", src=f"{src_f}:attention", field="attention.count")
    _bind(prov, f"{src_f}:updated", snapshot_date=updated_raw)
    # 🔴 誠實新鮮度(phase3b MAJOR 8):listing 舊寫「每日更新/附今日快照」,但抓取器自 07-13 起
    # 沒再跑 → 賣點是「盤前防呆」而清單過期**有實害**(處置股每日變動)。不假裝新鮮:
    # 快照日非今天就在成品裡明講落後幾天,並要買家自行重抓。措辭跟著資料走,不跟著行銷走。
    try:
        stale_days = (date.today() - date.fromisoformat(updated)).days
    except (TypeError, ValueError):
        # 解析不出日期 = 不知道多舊 → 當作「不確定新鮮度」照樣示警,**不可**當成新鮮(fail-safe)。
        stale_days = -1
    if stale_days > 0:
        _bind(prov, "derived:date.today() - 快照日", snapshot_lag_days=stale_days)
        stale_note = (f'<br><span class="neg">⚠️ 本檔快照日為 {RK.esc(updated)},'
                      f'距今 {stale_days} 天。處置股名單每個交易日都會變——'
                      f'<b>請務必以證交所當日公告為準,不要拿這份過期清單直接下單。</b></span>')
    elif stale_days < 0:
        stale_note = ('<br><span class="neg">⚠️ 無法判定本快照日期,請一律以證交所當日公告為準。</span>')
    else:
        stale_note = ""
    kpi = RK.kpi_row([("處置股(不可現沖)", f'{n_disp}<small> 檔</small>'),
                      ("注意股(風險高)", f'{n_attn}<small> 檔</small>')])
    # 處置股代號密表(每列一碼)
    rows = []
    for i, c in enumerate(disp_list):
        _bind(prov, f"{src_f}:disposition[{i}]", **{f"disposition[{i}]": c})
        rows.append(f'<td class="l tkr">{RK.esc(c)}</td><td class="l">處置股 · 分盤 · 不可當沖</td>')
    tbl = RK.dense_table([("代號", "l"), ("狀態", "l")], [[r] for r in rows]) if rows else \
        '<div class="metricrow">本日無處置股。</div>'
    checklist = ("<div class=\"block\" style=\"margin-top:10px\"><div class=\"bt\">盤前防呆清單</div>"
                 "<div class=\"metricrow\">• 這碼在今日處置清單裡嗎?是 → 不能當沖。<br>"
                 "• 是注意股嗎?是 → 減碼、放寬風控。<br>• 盤中流動性夠不夠你出場?<br>"
                 "• 進場前設好硬停損了嗎?<br>• 這份清單每個交易日都要重抓——它每天變。</div></div>")
    unit = (f'<div class="unit">{RK.sec_head(f"當沖適格快照（{updated}）", "快照")}'
            f'<div class="lead">報當沖單前先確認這碼不是<b>處置股</b>(分盤→不可現沖)或<b>注意股</b>。'
            f'快照日 {RK.esc(updated)},來源 TWSE OpenAPI。{stale_note}</div>{kpi}'
            f'<div class="block" style="margin-top:10px"><div class="bt">快照日處置股清單</div>{tbl}</div>{checklist}</div>')
    disc = RK.disclaimer_unit(DISCLAIMER, SOURCES_DAYTRADE,
                              f"本清單為 <b>{RK.esc(updated)}</b> 的快照,<b>每個交易日都會變</b>,"
                              "下單前請以證交所當日公告為準。")
    cover = RK.cover_page(
        b, ["台股當沖", "適格清單"], "免費磁鐵 · 盤前防呆",
        f"{CFG.BRAND['tagline_zh']}——報單前 30 秒防呆,先確認不是處置/注意股。",
        [("處置股", f'{n_disp}<small>檔</small>'), ("注意股", f'{n_attn}<small>檔</small>'),
         ("快照日", f'<small>{RK.esc(updated)}</small>')],
        "免費 · 換 Email 即得", "介紹 ≠ 推薦")
    html = RK.html_doc(b, "當沖適格清單", "免費磁鐵", "台股當沖適格清單", cover, [unit, disc])
    m = CFG.MAGNETS["M1"]
    listing = _listing("M1", "zh", "台股當沖適格清單｜處置股·注意股盤前防呆(附快照日)",
                       f"報當沖單前先確認這碼不是處置股(不能現沖)、也不是注意股。"
                       f"本份為 {updated} 的快照(檔內明標快照日);處置名單每個交易日都會變,"
                       f"下單前請以證交所當日公告為準。資料取自免費證交所 TWSE OpenAPI。"
                       f"教學風控工具,非投資建議。免費索取。",
                       ["當沖", "台股當沖", "處置股", "注意股", "盤前檢查", "風控清單", "當沖資格",
                        "TWSE", "當沖防呆", "交易紀律"], 0, 0, "magnet")
    return [{"sku": "M1", "lang": "zh", "platform": m["platform_zh"], "prov": prov,
             "gate_units": [unit], "listing": listing,
             "artifacts": [("台股當沖適格清單.pdf", "pdf", html)]}]


def build_M2(ctx) -> list[dict]:
    """L0 磁鐵:單檔旗艦體檢 台積電(zh,PDF)。"""
    ck = ctx["checkup"]
    ff = facts_for(ck, "2330")
    if not ff:
        return []
    prov = Provenance()
    b = CFG.BRAND
    name = (ck.get("by_code", {}).get("2330", {}) or {}).get("name", "台積電")
    card = checkup_card(prov, "2330", name, ff)
    disc = RK.disclaimer_unit(
        DISCLAIMER, SOURCES_CHECKUP,
        "本報告為<b>歷史數據體檢</b>(含息還原),只陳述數據位置,不判斷貴賤、不構成買賣建議。")
    lh = ff.get("long_horizon", {}).get("data", {})
    src_lh = _ck_src(ff, "long_horizon")
    cagr = disp(prov, lh.get("cagr", 0) * 100, src=src_lh, field="2330.long_horizon.cagr×100")
    # 維度數取實際 facts 數,不硬編(phase3b B2)
    nd = n_dims(ff)
    d_nd = disp(prov, nd, "{:.0f}", src="derived:len(facts_for(stock_checkup_facts.json, 2330))",
                field="2330.facts_count")
    cover = RK.cover_page(
        b, [f"{name}", "個股體檢報告"], "免費樣本 · 旗艦體檢",
        "含息還原20年:總報酬/年化、最慘與最猛一年、史上最長套牢、腰斬幾次、崩盤怎麼過。"
        "一般頻道只講賺多少,這裡連你要熬幾年套牢都算給你看。",
        [("年化報酬", f'<span class="{cls_updn(lh.get("cagr",0))}">{cagr}%</span>'),
         ("資料涵蓋", f'{disp(prov, lh.get("years",0), "{:.0f}", src=src_lh, field="2330.long_horizon.years")}<small>年</small>'),
         ("體檢維度", f'{d_nd}<small>項</small>')],
        "免費 · 訂閱週報每週都有深度體檢", "介紹 ≠ 推薦")
    html = RK.html_doc(b, f"{name}體檢", "免費樣本", f"{name}個股體檢報告", cover, [card, disc])
    listing = _listing("M2", "zh", f"{name}個股體檢報告(免費樣本)｜含息還原20年·套牢·腰斬",
                       f"{name}的{nd}項歷史體檢:含息還原20年總報酬/年化、最長套牢期、腰斬次數、"
                       "2008/2020/2022崩盤各跌多少。純歷史數據,介紹不等於推薦。免費索取。",
                       ["台積電", "個股體檢", "含息還原", "存股", "長期投資", "套牢期", "最大回撤",
                        "台股", "免費報告", "定存股"], 0, 0, "magnet")
    m = CFG.MAGNETS["M2"]
    return [{"sku": "M2", "lang": "zh", "platform": m["platform_zh"], "prov": prov,
             "gate_units": [card], "listing": listing,
             "artifacts": [(f"{name}體檢報告.pdf", "pdf", html)]}]


def build_T1(ctx) -> list[dict]:
    """L1 tripwire:定投追蹤模板(zh,PDF 導引 + xlsx 模板)。附三買法真實對照。"""
    ck = ctx["checkup"]
    prov = Provenance()
    b = CFG.BRAND
    # 收集所有覆蓋股的 three_way 當參考列(帶 fact_key 供存證回查)
    refs = []
    for code in (ck.get("by_code", {}) or {}):
        ff = facts_for(ck, code)
        tw = ff.get("three_way", {}).get("data") if ff.get("three_way") else None
        if tw and tw.get("stock_allin") and tw.get("bench"):
            nm = ck["by_code"][code].get("name", code)
            refs.append((code, nm,
                         tw["stock_allin"]["total_return"] * 100,
                         tw["stock_dca"]["total_return"] * 100,
                         tw["bench"]["total_return"] * 100,
                         _ck_src(ff, "three_way")))
        if len(refs) >= 8:
            break
    if not refs:
        return []
    # PDF 導引 unit:cmp bars(取前 5)+ 對照密表
    bars_units = []
    for code, nm, a, d, bch, src in refs[:5]:
        mx = max(a, d, bch) or 1
        rows = [("單筆 All-in", a, mx, False,
                 disp(prov, a, "{:+.0f}", src=src, field=f"{code}.three_way.stock_allin.total_return×100") + "%"),
                ("每月定投", d, mx, False,
                 disp(prov, d, "{:+.0f}", src=src, field=f"{code}.three_way.stock_dca.total_return×100") + "%"),
                ("同期 0050", bch, mx, False,
                 disp(prov, bch, "{:+.0f}", src=src, field=f"{code}.three_way.bench.total_return×100") + "%")]
        bars_units.append(f'<div class="block" style="margin-bottom:9px">'
                          f'<div class="bt">{RK.esc(nm)}（{RK.esc(code)}）</div>{RK.cmp_bars(rows)}</div>')
    trows = []
    for code, nm, a, d, bch, src in refs:
        trows.append([
            f'<td class="l tkr">{RK.esc(nm)}<span class="code">{RK.esc(code)}</span></td>',
            f'<td class="pos">{disp(prov, a, "{:+.0f}", src=src, field=f"{code}.three_way.stock_allin.total_return×100")}%</td>',
            f'<td class="pos">{disp(prov, d, "{:+.0f}", src=src, field=f"{code}.three_way.stock_dca.total_return×100")}%</td>',
            f'<td>{disp(prov, bch, "{:+.0f}", src=src, field=f"{code}.three_way.bench.total_return×100")}%</td>'])
    tbl = RK.dense_table([("個股", "l"), ("單筆All-in", "r"), ("每月定投", "r"), ("同期0050", "r")], trows)
    unit = (f'<div class="unit">{RK.sec_head("定投 vs 單筆 vs 大盤：真實對照", "含息還原")}'
            f'<div class="lead">附一段真實對照:近10年 單筆 All-in vs 每月定投 vs 同期 0050(含息還原,'
            f'取自體檢引擎實際價格計算)。搭配下載的 xlsx 模板,每月登一筆就自動算你的平均成本'
            f'(第 2–26 列已內建公式,空列自動留白)。</div>'
            f'{"".join(bars_units)}'
            f'<div class="block" style="margin-top:6px"><div class="bt">全部覆蓋股對照</div>{tbl}</div></div>')
    disc = RK.disclaimer_unit(DISCLAIMER, SOURCES_CHECKUP,
                              "對照數字為<b>歷史含息還原</b>,非未來保證;定投不保證獲利。教學工具,非投資建議。")
    cover = RK.cover_page(
        b, ["台股定投", "追蹤模板"], "L1 · 定投工具",
        "台股/ETF 定期定額追蹤模板:每月買進登一筆,自動看到平均成本。附真實對照。",
        [("對照個股", f'{disp(prov, len(refs), "{:.0f}", src="derived:體檢引擎有 three_way 的檔數", field="refs.count")}<small>檔</small>'),
         ("對照基準", '<small>0050</small>'), ("格式", '<small>xlsx</small>')],
        f"NT${CFG.ONE_OFF['T1']['ntd']} · 蝦皮/Gumroad", "介紹 ≠ 推薦")
    _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['T1']",
          price_ntd=CFG.ONE_OFF["T1"]["ntd"], price_usd=CFG.ONE_OFF["T1"]["usd"])
    html = RK.html_doc(b, "定投追蹤模板", "L1 tripwire", "台股定投追蹤模板", cover, [unit, disc])

    def build_xlsx(path: Path, _refs=refs):
        _build_dca_xlsx(path, _refs)

    listing = _listing("T1", "zh", "台股定投追蹤模板 xlsx｜附10年真實對照(All-in vs 定投 vs 0050)",
                       "台股/ETF 定期定額追蹤模板:每月買進登一列,自動算出投入金額、累計股數、"
                       "累計投入與平均成本(第 2–26 列已內建公式,沒填的列自動留白)。附真實對照:"
                       "近10年單筆All-in vs 每月定投 vs 0050(含息還原)。Excel/Google 試算表可開。教學工具,非投資建議。",
                       ["定投模板", "台股ETF", "定期定額", "0050", "平均成本", "Excel模板", "試算表",
                        "含息還原", "存股表格", "理財工具"], CFG.ONE_OFF["T1"]["ntd"], CFG.ONE_OFF["T1"]["usd"],
                       "shopee")
    return [{"sku": "T1", "lang": "zh", "platform": CFG.ONE_OFF["T1"]["platform_zh"], "prov": prov,
             "gate_units": [unit], "listing": listing,
             "artifacts": [("台股定投追蹤模板_導引.pdf", "pdf", html),
                           ("台股定投追蹤模板.xlsx", "xlsx", build_xlsx)]}]


def build_T2(ctx) -> list[dict]:
    """L1 tripwire:個股體檢單檔報告(zh,PDF)。示範用第一檔非 2330 的覆蓋股。"""
    ck = ctx["checkup"]
    codes = [c for c in (ck.get("by_code", {}) or {}) if c != "2330" and facts_for(ck, c)]
    if not codes:
        return []
    code = codes[0]
    prov = Provenance()
    b = CFG.BRAND
    name = ck["by_code"][code].get("name", code)
    ff = facts_for(ck, code)
    card = checkup_card(prov, code, name, ff)
    disc = RK.disclaimer_unit(DISCLAIMER, SOURCES_CHECKUP,
                              "歷史數據體檢,只陳述數據位置,不判斷貴賤、不構成買賣建議。")
    # 🔴 phase3b A3(d):舊版封面/listing 硬寫「11 項」,但 code 是 by_code 順序動態挑的
    # (今天 codes[0]=2317 碰巧 11 項),一旦輪到 2882(10)/00878(7)/2327(6) 就會出貨不實宣稱
    # 且 gate 結構上抓不到(非百分比宣稱)。維度數必須跟著實際出貨的那一檔走。
    nd = n_dims(ff)
    d_nd = disp(prov, nd, "{:.0f}", src=f"derived:len(facts_for(stock_checkup_facts.json, {code}))",
                field=f"{code}.facts_count")
    _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['T2']",
          price_ntd=CFG.ONE_OFF["T2"]["ntd"], price_usd=CFG.ONE_OFF["T2"]["usd"])
    cover = RK.cover_page(
        b, [f"{name}", "個股體檢報告"], "L1 · 單檔體檢",
        # 「任一」是假的(phase3b A4/MAJOR 12):買家不能選,本檔固定出貨 codes[0]。照實寫是哪一檔。
        f"本報告為 {name}({code}) 的 {nd} 項完整體檢:含息還原總報酬、最長套牢、腰斬、崩盤三段、"
        "毛利/營收/EPS 趨勢、股利、估值位階。本商品出貨的就是這一檔(不可指定其他個股);"
        "想每檔都有?升級訂閱週報。",
        [("體檢維度", f'{d_nd}<small>項</small>'), ("含息還原", '<small>是</small>'),
         ("代號", f'<small>{RK.esc(code)}</small>')],
        f"NT${CFG.ONE_OFF['T2']['ntd']} · 蝦皮/Gumroad", "介紹 ≠ 推薦")
    html = RK.html_doc(b, f"{name}體檢", "L1 tripwire", f"{name}個股體檢報告", cover, [card, disc])
    listing = _listing("T2", "zh", f"{name}({code})個股體檢報告｜含息還原20年·套牢·腰斬·估值位階",
                       f"本商品出貨的是 {name}({code}) 的 {nd} 項完整歷史體檢(固定此檔,不可指定其他個股):"
                       "含息還原總報酬/年化、最長套牢、腰斬次數、崩盤三段、毛利/營收/EPS、股利、估值位階。"
                       "純歷史數據,介紹不等於推薦。",
                       ["個股體檢", "含息還原", "存股", "長期投資", "套牢期", "最大回撤", "估值位階",
                        "台股", "權值股", "定存股"], CFG.ONE_OFF["T2"]["ntd"], CFG.ONE_OFF["T2"]["usd"], "shopee")
    return [{"sku": "T2", "lang": "zh", "platform": CFG.ONE_OFF["T2"]["platform_zh"], "prov": prov,
             "gate_units": [card], "listing": listing,
             "artifacts": [(f"{name}體檢報告.pdf", "pdf", html)]}]


def build_C1(ctx) -> list[dict]:
    """L2 core:全市場回測數據包(zh+en,xlsx 多表 + 摘要 PDF)。"""
    adaptive = ctx["adaptive"]; longshort = ctx["longshort"]; perstock = ctx["perstock"]
    if not adaptive:
        return []
    ls_map = {r["code"]: r for r in longshort}
    ps_map = {r["code"]: r for r in perstock}
    n = len(adaptive)
    prof = [r for r in adaptive if r["a_net"] > 0]
    nets = [r["a_net"] for r in adaptive]
    med = statistics.median(nets)
    listed = sum(1 for r in adaptive if r["market"] == "上市")
    top = sorted(adaptive, key=lambda r: r["a_net"], reverse=True)[:20]

    out = []
    for lang in CFG.ONE_OFF["C1"]["langs"]:
        prov = Provenance()
        b = CFG.BRAND
        SRC_AD = "twdata/adaptive_per_stock.csv"
        src_agg = f"derived:{SRC_AD}(全 1770 列聚合)"
        pn = disp(prov, len(prof) / n * 100, src=src_agg, field="a_net>0 佔比 = count/total×100")
        nn = disp(prov, n, "{:.0f}", src=f"{SRC_AD}:列數", field="tickers.total")
        npf = disp(prov, len(prof), "{:.0f}", src=src_agg, field="a_net>0.count")
        ln = disp(prov, listed, "{:.0f}", src=src_agg, field="market=='上市'.count")
        otc = disp(prov, n - listed, "{:.0f}", src=src_agg, field="market!='上市'.count")
        mn = disp(prov, med, src=src_agg, field="a_net.median")
        # 個股名只有中文(來源資料無官方英文名)。英文版**不猜英譯**——1770 檔沒有可信對照表,
        # 硬掰英文名 = 憑空造資料,比留中文更糟。作法:英文版把 Ticker 提為主鍵獨立一欄、
        # 名稱欄明標 (zh-TW),並在免責註明「ticker 才是可靠識別鍵」(phase3b B3 的誠實解)。
        trows = []
        for r in top:
            src_r = f'{SRC_AD}[code={r["code"]}]'
            net = disp(prov, r["a_net"], "{:+.1f}", src=src_r, field=f'{r["code"]}.a_net')
            win = disp(prov, r["a_win"], src=src_r, field=f'{r["code"]}.a_win')
            cells = [f'<td class="{cls_updn(r["a_net"])}">{net}%</td>', f'<td>{win}%</td>']
            if lang == "zh":
                trows.append([f'<td class="l tkr">{RK.esc(r["name"])}'
                              f'<span class="code">{RK.esc(r["code"])}</span></td>'] + cells)
            else:
                trows.append([f'<td class="l tkr">{RK.esc(r["code"])}</td>',
                              f'<td class="l">{RK.esc(r["name"])}</td>'] + cells)
        # 方法與限制揭露(MAJOR 9):NT$990 不能賣黑箱。內容全部可回源碼查證(見 METHOD_ZH 上方註解)。
        meth_rows = METHOD_ZH if lang == "zh" else METHOD_EN
        method_unit = (
            f'<div class="unit">'
            f'{RK.sec_head("方法與限制(買前必讀)" if lang == "zh" else "Method & limitations (read before you buy)", "揭露" if lang == "zh" else "Disclosure")}'
            + "".join(f'<div class="block" style="margin-bottom:9px"><div class="bt">{RK.esc(t)}</div>'
                      f'<div class="metricrow">{body}</div></div>' for t, body in meth_rows)
            + '</div>')
        if lang == "zh":
            tbl = RK.dense_table([("個股", "l"), ("自適應淨報酬", "r"), ("勝率", "r")], trows)
            unit = (f'<div class="unit">{RK.sec_head("全市場回測重點(每個數字都在 xlsx 查得到)", f"{nn}檔")}'
                    f'<div class="lead">自適應策略跑遍 <b>{nn}</b> 檔上市櫃(上市 {ln}/上櫃其他 {otc})。'
                    f'附完整 xlsx(可排序篩選):淨報酬/獲利因子/最大回撤/交易數/勝率 + 多空 + Sharpe。'
                    f'策略定義、成本假設與已知偏誤見下一節「方法與限制」。</div>'
                    f'<div class="two"><div class="block"><div class="bt">誠實看法:中位數優先</div>'
                    f'<div class="metricrow">自適應淨報酬 > 0:<b>{npf} / {nn}</b>(<b>{pn}%</b>)<br>'
                    f'全市場淨報酬<b>中位數 {mn}%</b><br><span style="color:var(--tx3)">'
                    f'少數幾檔亮眼不代表策略通用——看中位數。</span></div></div>'
                    f'<div class="block"><div class="bt">自適應淨報酬 前20</div>{tbl}</div></div></div>')
            cover = RK.cover_page(
                b, ["台股全市場", "回測數據包"], "L2 · 全市場數據",
                f"自適應策略跑遍 {n} 檔上市櫃,每檔含淨報酬/獲利因子/最大回撤/交易數/勝率+多空+Sharpe。"
                "附完整方法與限制揭露(含成本假設與倖存者偏誤)。用來自己驗證『策略無腦套全市場』到底行不行。",
                [("覆蓋", f'{nn}<small>檔</small>'), ("正報酬佔比", f'{pn}<small>%</small>'),
                 ("中位淨報酬", f'<span class="{cls_updn(med)}">{mn}%</span>')],
                f"NT${CFG.ONE_OFF['C1']['ntd']} · Portaly/Gumroad", "介紹 ≠ 推薦", lang)
            _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['C1']", price_ntd=CFG.ONE_OFF["C1"]["ntd"])
            disc = RK.disclaimer_unit(DISCLAIMER, SOURCES_BACKTEST,
                                      "此回測為 <b>2026-06-12 靜態快照</b>,僅供教育與自我驗證,"
                                      "非即時可交易訊號、不代表現在或未來。"
                                      "<b>成本假設、倖存者偏誤等已知限制見「方法與限制」節。</b>", lang)
            title = f"台股全市場回測數據包 {n}檔｜自適應+多空+Sharpe(xlsx+摘要)"
            desc = (f"台股全市場回測數據包:自適應策略跑遍 {n} 檔上市櫃,每檔含淨報酬、獲利因子、"
                    "最大回撤、交易數、勝率、多空、Sharpe。專業級 xlsx(凍結/篩選/紅綠條件格式,附欄位說明表)"
                    "+摘要 PDF。附完整方法揭露:策略規則、已扣的手續費/證交稅、未模擬滑價、倖存者偏誤都寫明。"
                    "純歷史統計、非投資建議、靜態快照非即時。")
            tags = ["台股回測", "全市場數據", "量化數據包", "選股數據", "回測xlsx", "台股量化",
                    "趨勢策略", "勝率數據", "最大回撤", "獲利因子"]
            fn = "台股全市場回測數據包_摘要.pdf"; xn = f"台股全市場回測_{n}檔.xlsx"
            platform = CFG.ONE_OFF["C1"]["platform_zh"]
        else:
            tbl = RK.dense_table([("Ticker", "l"), ("Name (zh-TW)", "l"),
                                  ("Adaptive Net", "r"), ("Win %", "r")], trows)
            unit = (f'<div class="unit">{RK.sec_head("Full-market backtest — every figure is in the xlsx", f"{nn} tickers")}'
                    f'<div class="lead">An adaptive run across <b>{nn}</b> listed/OTC Taiwan '
                    f'tickers ({ln} listed / {otc} OTC). Ships with a full sortable xlsx: net return, profit '
                    f'factor, max drawdown, trades, win rate + long/short + Sharpe. The strategy definition, cost '
                    f'assumptions and known biases are in the “Method &amp; limitations” section below.</div>'
                    f'<div class="two"><div class="block"><div class="bt">Read it the honest way: median first</div>'
                    f'<div class="metricrow">Adaptive net return &gt; 0: <b>{npf} / {nn}</b> (<b>{pn}%</b>)<br>'
                    f'Market-wide <b>median {mn}%</b><br><span style="color:var(--tx3)">'
                    f'A few winners do not make a strategy universal — judge it by the median.</span></div></div>'
                    f'<div class="block"><div class="bt">Top 20 by adaptive net return</div>{tbl}</div></div></div>')
            cover = RK.cover_page(
                b, ["Taiwan Full-Market", "Backtest Pack"], "L2 · Market Data",
                f"An adaptive backtest across {n} listed & OTC Taiwan tickers — net return, "
                "profit factor, max drawdown, trades, win rate + long/short + Sharpe. Ships with a full method "
                "and limitations disclosure (cost assumptions and survivorship bias included). "
                "Rare English-language Taiwan quant data.",
                [("Tickers", f'{nn}'), ("% positive", f'{pn}<small>%</small>'),
                 ("Median net", f'<span class="{cls_updn(med)}">{mn}%</span>')],
                f"US${CFG.ONE_OFF['C1']['usd']} · Gumroad", "Description ≠ recommendation", lang)
            _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['C1']", price_usd=CFG.ONE_OFF["C1"]["usd"])
            disc = RK.disclaimer_unit(DISCLAIMER_EN, SOURCES_BACKTEST_EN,
                                      "This backtest is a <b>2026-06-12 static snapshot</b> for education and "
                                      "self-verification only; not a live tradable signal. "
                                      "<b>Cost assumptions, survivorship bias and other known limitations are set "
                                      "out in the “Method &amp; limitations” section.</b><br><br>"
                                      "<b>A note on names</b>: company names are shown in their local zh-TW form, "
                                      "because the source data carries no official English name. The ticker is the "
                                      "reliable key.", lang)
            title = f"Taiwan Full-Market Backtest Data Pack ({n} tickers, xlsx + summary)"
            desc = (f"Cleaned full-market backtest data for the Taiwan market — {n} listed & OTC tickers, each "
                    "with net return, profit factor, max drawdown, trades, win rate, long/short and Sharpe. "
                    "Professional xlsx (freeze/filter/conditional formatting, with a column-definition sheet) + "
                    "summary PDF. Full method disclosure included: strategy rules, the commissions and transaction "
                    "tax that are deducted, the fact that slippage is not modelled, and the survivorship bias in the "
                    "universe. Historical statistics only — this is not investment advice and nothing here predicts "
                    "the future. Company names appear in local zh-TW form; the ticker is the reliable key.")
            tags = ["taiwan stock data", "backtest xlsx", "quant dataset", "stock screener data",
                    "trend following", "tw market data", "win rate data", "drawdown", "sharpe", "finance dataset"]
            fn = "taiwan_fullmarket_backtest_summary.pdf"; xn = f"taiwan_fullmarket_backtest_{n}.xlsx"
            platform = CFG.ONE_OFF["C1"]["platform_en"]
        html = RK.html_doc(b, "全市場回測數據包" if lang == "zh" else "Full-Market Backtest Pack",
                           "L2 core", title, cover, [unit, method_unit, disc], lang)
        listing = _listing("C1", lang, title, desc, tags, CFG.ONE_OFF["C1"]["ntd"],
                           CFG.ONE_OFF["C1"]["usd"], platform)
        _lg = lang

        def _xlsx(path: Path, _l=_lg):
            _build_backtest_xlsx(path, adaptive, ls_map, ps_map, _l)

        out.append({"sku": "C1", "lang": lang, "platform": platform, "prov": prov,
                    # method_unit **刻意不入 gate_units**:它的數字是**源碼常數**(費率/ADX門檻/期數),
                    # 不是資料管線算出的績效數。把它們灌進 prov.pool 只會讓池變大 → 憑空數字更容易
                    # 撞到鄰居(guard 自己的「規模詛咒」),等於為了讓揭露過關而弱化守門。
                    # 這段的正確防線是 tests/test_product_factory_v2.py::TestMethodDisclosure
                    # ——直接對 strategy.COST_MODELS / AdaptiveParams 斷言,參數漂移即紅燈。
                    "gate_units": [unit], "listing": listing,
                    "artifacts": [(fn, "pdf", html), (xn, "xlsx", _xlsx)]})
    return out


def build_C2(ctx) -> list[dict]:
    """L2 core:權值股體檢合輯(zh+en,PDF 全覆蓋卡 + xlsx 總覽)。"""
    ck = ctx["checkup"]
    codes = [c for c in (ck.get("by_code", {}) or {}) if facts_for(ck, c)]
    if not codes:
        return []

    # 🔴 phase3b B2:最高價 SKU(NT$1280/US$39)封面硬寫「每檔維度 11 項」,實查 9 檔裡
    # 國泰金=10、00878=7、國巨=6 → 對 3/9 檔(33%)不實,且 gate 抓不到(非百分比宣稱)。
    # 合輯本來就參差,誠實作法不是挑一個數字充場面,而是**照實講範圍 + 逐檔揭露**
    # (順帶解掉 MINOR「00878/2327 資料較薄卻同價未揭露」:xlsx 總覽新增「體檢維度」欄)。
    dims = {c: n_dims(facts_for(ck, c)) for c in codes}
    d_lo, d_hi = min(dims.values()), max(dims.values())
    thin = sorted([c for c in codes if dims[c] < d_hi], key=lambda c: dims[c])

    def build_xlsx_lang(path: Path, _lang="zh"):
        _build_checkup_overview_xlsx(path, ck, codes, dims, _lang)

    out = []
    for lang in CFG.ONE_OFF["C2"]["langs"]:
        prov = Provenance()
        b = CFG.BRAND
        cards = [checkup_card(prov, c, ck["by_code"][c].get("name", c), facts_for(ck, c), lang)
                 for c in codes]
        src_dims = "derived:len(facts_for(stock_checkup_facts.json, <code>)) 逐檔"
        nn = disp(prov, len(codes), "{:.0f}", src=src_dims, field="codes.count")
        d_lo_s = disp(prov, d_lo, "{:.0f}", src=src_dims, field="facts_count.min")
        d_hi_s = disp(prov, d_hi, "{:.0f}", src=src_dims, field="facts_count.max")
        dim_stat = f'{d_lo_s}–{d_hi_s}<small>項</small>' if d_lo != d_hi else f'{d_hi_s}<small>項</small>'
        thin_zh = "、".join(f'{ck["by_code"][c].get("name", c)}({c}) {dims[c]} 項' for c in thin)
        thin_en = ", ".join(f'{c} ({dims[c]})' for c in thin)
        for c in codes:
            _bind(prov, src_dims, **{f"{c}.facts_count": dims[c]})
        if lang == "zh":
            disc = RK.disclaimer_unit(
                DISCLAIMER, SOURCES_CHECKUP,
                "歷史數據體檢(含息還原),只陳述數據位置,不判斷貴賤、不構成買賣建議。"
                + (f"<br><br><b>資料完整度揭露</b>:各檔可得維度不同(本合輯 {d_lo}–{d_hi} 項)。"
                   f"維度較少者:{thin_zh}——來源資料本身即無該欄位(如 ETF 無毛利率/本益比),"
                   f"我們<b>不會為了湊數而填假值</b>,缺就是留空。逐檔維度數見附帶 xlsx「體檢維度」欄。"
                   if thin else ""))
            cover = RK.cover_page(
                b, ["台股權值股", "體檢合輯"], "L2 · 深度體檢",
                "覆蓋權值股全體檢合輯:每檔含息還原20年總報酬、最長套牢、腰斬、崩盤三段、"
                "毛利/營收/EPS、股利、估值位階。一次擁有;隨體檢覆蓋成長。",
                [("覆蓋個股", f'{nn}<small>檔</small>'), ("每檔維度", dim_stat),
                 ("含息還原", '<small>是</small>')],
                f"NT${CFG.ONE_OFF['C2']['ntd']} · Portaly/Gumroad", "介紹 ≠ 推薦", lang)
            _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['C2']",
                  price_ntd=CFG.ONE_OFF["C2"]["ntd"])
            title = "台股權值股體檢合輯｜含息還原20年·套牢·腰斬·估值位階(PDF+xlsx)"
            desc = (f"台股權值股體檢報告合輯({len(codes)} 檔):含息還原20年,每檔不只講報酬,還算最長套牢、"
                    "腰斬過幾次、2008/2020/2022崩盤各跌多少、毛利/營收/EPS趨勢、股利、估值位階。"
                    f"附深色 PDF + xlsx 總覽。各檔可得維度 {d_lo}–{d_hi} 項不等(來源無該欄位者留空不補假值,"
                    "逐檔維度數見 xlsx)。純歷史數據體檢,介紹不等於推薦。")
            tags = ["台股體檢", "含息還原", "存股", "長期投資", "套牢期", "最大回撤", "估值位階",
                    "權值股", "崩盤數據", "定存股"]
            fn = "台股權值股體檢合輯.pdf"; xn = "台股權值股體檢_總覽.xlsx"
            platform = CFG.ONE_OFF["C2"]["platform_zh"]
        else:
            disc = RK.disclaimer_unit(
                DISCLAIMER_EN, SOURCES_CHECKUP_EN,
                "Dividend-adjusted historical health-checks; position statements only, "
                "not buy/sell advice."
                + (f"<br><br><b>Data-completeness disclosure</b>: the number of available check dimensions "
                   f"differs per ticker (this bundle: {d_lo}–{d_hi}). Thinner tickers: {thin_en}. "
                   f"The source data simply lacks those fields (an ETF has no gross margin or P/E, for example); "
                   f"we <b>do not invent values to pad the count</b> — missing stays blank. Per-ticker counts are "
                   f"in the “Check dimensions” column of the bundled xlsx." if thin else "")
                + "<br><br><b>A note on names</b>: company names are shown in their local zh-TW form, "
                  "because the source data carries no official English name. The ticker is the reliable key.",
                lang)
            cover = RK.cover_page(
                b, ["Taiwan Blue-Chip", "Health-Check Bundle"], "L2 · Deep Checks",
                "Dividend-adjusted 20-year health checks for major Taiwan tickers: total/annualised return, "
                "longest underwater stretch, halvings, 2008/2020/2022 crashes, margin/revenue/EPS, dividends, valuation position.",
                [("Tickers", f'{nn}'), ("Facts each", dim_stat.replace("<small>項</small>", "")),
                 ("Div-adjusted", '<small>yes</small>')],
                f"US${CFG.ONE_OFF['C2']['usd']} · Gumroad", "Description ≠ recommendation", lang)
            _bind(prov, "quant-service/ecommerce/config.py:ONE_OFF['C2']",
                  price_usd=CFG.ONE_OFF["C2"]["usd"])
            title = "Taiwan Blue-Chip Stock Health-Check Bundle (dividend-adjusted, PDF+xlsx)"
            desc = (f"A bundle of dividend-adjusted 20-year health-check reports for {len(codes)} major Taiwan "
                    "blue chips. Each shows the holding experience: longest underwater period, halvings, drawdowns "
                    "through 2008/2020/2022, margin/revenue/EPS trends, dividends, valuation position. Dark PDF + "
                    f"xlsx overview. Available check dimensions vary by ticker ({d_lo}–{d_hi}); where the source "
                    "lacks a field we leave it blank rather than invent a value. Historical statistics only — "
                    "this is not investment advice, and a description is not a recommendation. "
                    "Company names appear in local zh-TW form; the ticker is the reliable key.")
            tags = ["taiwan stocks", "stock report", "dividend adjusted", "drawdown", "tsmc data",
                    "long term investing", "buy and hold", "holding period", "valuation", "risk report"]
            fn = "taiwan_bluechip_healthcheck_bundle.pdf"; xn = "taiwan_bluechip_healthcheck_overview.xlsx"
            platform = CFG.ONE_OFF["C2"]["platform_en"]
        html = RK.html_doc(b, "權值股體檢合輯" if lang == "zh" else "Blue-Chip Health-Check",
                           "L2 core", title, cover, cards + [disc], lang)
        listing = _listing("C2", lang, title, desc, tags, CFG.ONE_OFF["C2"]["ntd"],
                           CFG.ONE_OFF["C2"]["usd"], platform)
        _lg = lang

        def _xlsx(path: Path, _l=_lg):
            build_xlsx_lang(path, _l)

        out.append({"sku": "C2", "lang": lang, "platform": platform, "prov": prov,
                    "gate_units": cards, "listing": listing,
                    "artifacts": [(fn, "pdf", html), (xn, "xlsx", _xlsx)]})
    return out


# ── xlsx builders(§6:深表頭/凍結/篩選/台股紅綠條件格式/data bar/代號留字串)──────
def _xlsx_head(ws, ncol):
    from openpyxl.styles import Font, PatternFill, Alignment
    ws.row_dimensions[1].height = 26
    for c in range(1, ncol + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = PatternFill(start_color="0B0E14", end_color="0B0E14", fill_type="solid")
        cell.font = Font(color="E3B93E", bold=True, size=11)
        cell.alignment = Alignment(horizontal="center", vertical="center")


_XL_C1 = {
    "zh": {"s1": "全市場回測", "s2": "風險明細", "s3": "欄位說明",
           "h1": ["代號", "名稱", "市場", "自適應淨報酬%", "獲利因子", "最大回撤%", "交易數",
                  "勝率%", "多空淨報酬%", "做空次數"],
           "h2": ["代號", "名稱", "淨報酬%", "獲利因子", "最大回撤%", "Sharpe", "報酬回撤比", "起", "迄"],
           "h3": ["欄位", "定義", "單位/符號"]},
    "en": {"s1": "Full-market backtest", "s2": "Risk detail", "s3": "Column definitions",
           "h1": ["Ticker", "Name (zh-TW)", "Market", "Adaptive net return%", "Profit factor",
                  "Max drawdown%", "Trades", "Win rate%", "Long/short net return%", "Short trades"],
           "h2": ["Ticker", "Name (zh-TW)", "Net return%", "Profit factor", "Max drawdown%", "Sharpe",
                  "Return/max-drawdown", "Start", "End"],
           "h3": ["Column", "Definition", "Unit / sign"]},
}
# 欄位說明表(MAJOR 9:NT$990 的 xlsx 舊版 0 個註解、無欄位說明 → 買家不知道每欄是什麼)。
# 「最大回撤%」符號在兩張表不同(本表為正值幅度、C2 總覽為負值)——這正是 A7 色階讀反的根因,
# 所以說明表**明講符號**,不讓買家自己猜。
_XL_C1_DEFS = {
    "zh": [("代號", "上市/上櫃股票代號(字串,保留前導零如 00878)", "—"),
           ("名稱", "公司/ETF 名稱(來源資料原樣;帶*者為來源自帶標記)", "—"),
           ("市場", "上市 或 上櫃", "—"),
           ("自適應淨報酬%", "該檔在整段回測期的策略淨報酬(已扣手續費與證交稅,未計滑價)", "%,正=獲利"),
           ("獲利因子", "總獲利 ÷ 總虧損;>1 表示總獲利大於總虧損", "倍,無單位"),
           ("最大回撤%", "權益曲線自高點的最大跌幅", "%,**本表為正值幅度**(50=曾回撤50%)"),
           ("交易數", "整段期間完成的交易筆數;筆數太少的統計意義低", "筆"),
           ("勝率%", "獲利交易筆數 ÷ 總交易筆數", "%"),
           ("多空淨報酬%", "同策略加入放空後的淨報酬(來源 longshort_per_stock.csv)", "%,正=獲利"),
           ("做空次數", "放空版本中的做空交易筆數", "筆")],
    "en": [("Ticker", "Listed/OTC stock code (stored as text so leading zeros survive, e.g. 00878)", "—"),
           ("Name (zh-TW)", "Company/ETF name exactly as the source carries it (a trailing * is the source's own marker). No official English name exists in the source", "—"),
           ("Market", "上市 = main board, 上櫃 = OTC", "—"),
           ("Adaptive net return%", "Strategy net return over the whole backtest window (commissions and transaction tax deducted; slippage NOT modelled)", "%, positive = profit"),
           ("Profit factor", "Gross profit ÷ gross loss; > 1 means gross profit exceeded gross loss", "ratio, unitless"),
           ("Max drawdown%", "Largest peak-to-trough fall in the equity curve", "%, **positive magnitude in this sheet** (50 = fell 50%)"),
           ("Trades", "Number of completed trades; very low counts carry little statistical meaning", "count"),
           ("Win rate%", "Winning trades ÷ total trades", "%"),
           ("Long/short net return%", "Net return of the same strategy with shorting enabled (source: longshort_per_stock.csv)", "%, positive = profit"),
           ("Short trades", "Number of short trades in the long/short variant", "count")],
}


def _build_backtest_xlsx(path: Path, adaptive, ls_map, ps_map, lang: str = "zh"):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
    UP, DN, MID, GOLD = "E5484D", "2FB877", "F2F2F2", "E3B93E"
    T = _XL_C1.get(lang, _XL_C1["zh"])
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = T["s1"]
    ws.append(T["h1"])
    _xlsx_head(ws, 10)
    for r in sorted(adaptive, key=lambda x: x["a_net"], reverse=True):
        ls = ls_map.get(r["code"], {})
        ws.append([str(r["code"]), r["name"], r["market"], round(r["a_net"], 2),
                   _f(r.get("a_pf")), _f(r.get("a_dd")), _i(r.get("a_tr")), _f(r.get("a_win")),
                   _f(ls.get("ls_net")), _i(ls.get("short_trades"))])
    last = len(adaptive) + 1
    if last >= 2:
        ws.freeze_panes = "C2"; ws.auto_filter.ref = f"A1:J{last}"
        for col, fmt in (("D", '0.00"%"'), ("F", '0.00"%"'), ("H", '0.0"%"'), ("I", '0.00"%"')):
            for row in range(2, last + 1):
                ws[f"{col}{row}"].number_format = fmt
        ws.conditional_formatting.add(f"D2:D{last}", ColorScaleRule(
            start_type="min", start_color=DN, mid_type="num", mid_value=0, mid_color=MID,
            end_type="max", end_color=UP))  # 高報酬=紅(台股)
        # 本表「最大回撤%」是**正值幅度**(50=曾回撤50%)→ min(幅度最小=最好)=紅、max(最糟)=綠,
        # 與同表「高報酬=紅」的台股慣例一致。⚠️ C2 總覽的回撤欄是**負值**,故色階方向必須相反
        # (見 _build_checkup_overview_xlsx;phase3b A7 就是兩表符號相反卻套同一組方向所致)。
        ws.conditional_formatting.add(f"F2:F{last}", ColorScaleRule(
            start_type="min", start_color=UP, mid_type="percentile", mid_value=50, mid_color=MID,
            end_type="max", end_color=DN))
        ws.conditional_formatting.add(f"H2:H{last}", DataBarRule(start_type="min", end_type="max", color=GOLD))
    for i, w in enumerate([9, 16, 6, 14, 10, 12, 8, 9, 13, 10], start=1):
        ws.column_dimensions[chr(64 + i)].width = w

    # 風險明細(per_stock:sharpe/rdd/起訖)
    ws2 = wb.create_sheet(T["s2"])
    ws2.append(T["h2"])
    _xlsx_head(ws2, 9)
    got = 0
    for r in adaptive:
        ps = ps_map.get(r["code"])
        if not ps:
            continue
        ws2.append([str(r["code"]), r["name"], _f(ps.get("net_profit_pct")), _f(ps.get("profit_factor")),
                    _f(ps.get("max_dd_pct")), _f(ps.get("sharpe")), _f(ps.get("return_over_maxdd")),
                    ps.get("start", ""), ps.get("end", "")])
        got += 1
    l2 = got + 1
    if l2 >= 2:
        ws2.freeze_panes = "C2"; ws2.auto_filter.ref = f"A1:I{l2}"
        ws2.conditional_formatting.add(f"F2:F{l2}", ColorScaleRule(
            start_type="min", start_color=DN, mid_type="num", mid_value=0, mid_color=MID,
            end_type="max", end_color=UP))
    for i, w in enumerate([9, 16, 11, 10, 12, 9, 12, 11, 11], start=1):
        ws2.column_dimensions[chr(64 + i)].width = w

    # 欄位說明表(MAJOR 9):每欄是什麼、單位與符號、成本含不含,寫清楚。
    ws3 = wb.create_sheet(T["s3"])
    ws3.append(T["h3"])
    _xlsx_head(ws3, 3)
    for row in _XL_C1_DEFS.get(lang, _XL_C1_DEFS["zh"]):
        ws3.append(list(row))
    for i, w in enumerate([22, 78, 30], start=1):
        ws3.column_dimensions[chr(64 + i)].width = w
    for r in range(2, len(_XL_C1_DEFS.get(lang, _XL_C1_DEFS["zh"])) + 2):
        ws3.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[r].height = 30
    meth = METHOD_ZH if lang == "zh" else METHOD_EN
    rr = len(_XL_C1_DEFS.get(lang, _XL_C1_DEFS["zh"])) + 3
    for t, body in meth:
        ws3.cell(row=rr, column=1, value=t).font = Font(color="E3B93E", bold=True, size=10)
        c = ws3.cell(row=rr, column=2, value=re.sub(r"<[^>]+>", "", body))
        c.font = Font(color="8A94A6", size=9)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws3.row_dimensions[rr].height = 58
        rr += 1
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _build_dca_xlsx(path: Path, refs):
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.formatting.rule import DataBarRule
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = "定投追蹤"
    ws.append(["日期", "代號", "名稱", "買進股數", "成交價", "投入金額", "累計股數", "累計投入", "平均成本"])
    _xlsx_head(ws, 9)
    # 🔴 phase3b A6:舊版只有第 2、3 列有公式,且第 3 列 D3=10 但 E3 成交價空白 →
    #    F3=D3*E3=0、I3=H3/G3=1850/20=**92.5**,買家一開檔就看到荒謬的「平均成本 92.5」
    #    (0050 買在 185)。第 4 列起更是完全沒公式,與「每月登一列自動算」的賣點直接衝突。
    # 修法三件事:
    #  ① 只留**一列**完整示範(有價),不留半殘的第二列;示範列明標「範例」要買家覆蓋。
    #  ② 公式改 SUM($X$2:Xn) 累計版:對空列/跳列都穩(舊版 G3=G2+D3 遇空值會壞),
    #     且用 IF(D="","") 包住 → **沒填資料的列顯示空白,不會蹦出 0 或 92.5 這種假數字**。
    #  ③ 公式鋪到第 DCA_ROWS 列(而非只有 2 列),買家從第 4 列開始登也真的會自動算。
    # 示範價 185.0 是**模板佔位數字**、非市場宣稱(A6 欄位已標「範例」),故不入 prov 池。
    DCA_ROWS = 26                       # 資料列鋪到第 26 列(2 年多的月定投)
    ws.append(["2026-01-05", "0050", "元大台灣50(範例列,請覆蓋)", 10, 185.0,
               "=IF(D2=\"\",\"\",D2*E2)", "=IF(D2=\"\",\"\",SUM($D$2:D2))",
               "=IF(D2=\"\",\"\",SUM($F$2:F2))",
               "=IF(OR(D2=\"\",SUM($D$2:D2)=0),\"\",SUM($F$2:F2)/SUM($D$2:D2))"])
    for r in range(3, DCA_ROWS + 1):
        ws.append([None, None, None, None, None,
                   f'=IF(D{r}="","",D{r}*E{r})',
                   f'=IF(D{r}="","",SUM($D$2:D{r}))',
                   f'=IF(D{r}="","",SUM($F$2:F{r}))',
                   f'=IF(OR(D{r}="",SUM($D$2:D{r})=0),"",SUM($F$2:F{r})/SUM($D$2:D{r}))'])
    for row in range(2, DCA_ROWS + 1):
        for col in ("E", "F", "H", "I"):
            ws[f"{col}{row}"].number_format = "#,##0.0"
    ws.freeze_panes = "B2"
    for i, w in enumerate([12, 8, 22, 9, 9, 11, 10, 11, 10], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    # 說明放在資料列**之下**(phase3b MINOR 14:舊版塞在 A6,買家往下登資料會撞到)
    for i, txt in enumerate((
            f"用法:每月買進登一列(第 2–{DCA_ROWS} 列已內建公式),F/G/H/I 欄會自動算——"
            "投入金額、累計股數、累計投入、平均成本(=累計投入÷累計股數)。",
            "第 2 列是範例(0050 買在 185 為示範用佔位數字,非真實報價),請直接覆蓋成你自己的紀錄。",
            "沒填「買進股數」的列會自動留空白,不會顯示 0 或錯誤的平均成本。",
            f"要更多列?選第 {DCA_ROWS} 列的 F:I 往下拉即可複製公式。代號留字串以免前導零(如 0050/00878)被吃掉。",
    ), start=0):
        c = ws.cell(row=DCA_ROWS + 2 + i, column=1, value=txt)
        c.font = Font(color="8A94A6", size=9, italic=True)

    ws2 = wb.create_sheet("真實對照")
    ws2.append(["代號", "名稱", "單筆All-in總報酬%", "每月定投總報酬%", "同期0050總報酬%"])
    _xlsx_head(ws2, 5)
    for code, nm, a, d, bch, _src in refs:
        ws2.append([str(code), nm, round(a, 1), round(d, 1), round(bch, 1)])
    last = len(refs) + 1
    if last >= 2:
        ws2.freeze_panes = "C2"; ws2.auto_filter.ref = f"A1:E{last}"
        for col in ("C", "D", "E"):
            for row in range(2, last + 1):
                ws2[f"{col}{row}"].number_format = '#,##0.0"%"'
            ws2.conditional_formatting.add(f"{col}2:{col}{last}",
                                           DataBarRule(start_type="min", end_type="max", color="E3B93E"))
    for i, w in enumerate([9, 16, 18, 18, 18], start=1):
        ws2.column_dimensions[chr(64 + i)].width = w
    ws2.cell(row=last + 2, column=1,
             value="對照為歷史含息還原(取自體檢引擎),非未來保證;定投不保證獲利。介紹≠推薦。")
    ws2.cell(row=last + 2, column=1).font = Font(color="8A94A6", size=9, italic=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


_XL_C2 = {
    "zh": {"s1": "總覽",
           "h": ["代號", "名稱", "20年年化%", "最大回撤%", "估值位階(百分位)", "最長套牢(年)",
                 "現金殖利率%", "體檢維度"],
           "note": ("色階:年化高=紅(台股慣例:紅=好)。"
                    "**最大回撤欄為負值**(-98.5 比 -22.3 更慘),故最慘=綠、最輕=紅,與年化欄同樣是「紅=好」。"
                    "｜燈號僅標估值位置非買賣｜「體檢維度」= 該檔實際可得的體檢項目數(來源無該欄位者留空,不補假值)"
                    "｜來源:FinMind/Yahoo 含息還原｜介紹≠推薦")},
    "en": {"s1": "Overview",
           "h": ["Ticker", "Name (zh-TW)", "20y CAGR%", "Max drawdown%", "Valuation percentile",
                 "Longest underwater (yrs)", "Cash dividend yield%", "Check dimensions"],
           "note": ("Colour scale: higher CAGR = red (Taiwan convention: red = good, green = bad). "
                    "**The max-drawdown column is negative** (-98.5 is worse than -22.3), so the worst is green "
                    "and the mildest is red — the same 'red = good' rule as the CAGR column. "
                    "| Traffic lights mark valuation position only, not buy/sell. "
                    "| 'Check dimensions' = how many check items are actually available for that ticker "
                    "(where the source lacks a field we leave it blank rather than invent a value). "
                    "| Sources: FinMind / Yahoo dividend-adjusted. | Description is not a recommendation.")},
}


def _build_checkup_overview_xlsx(path: Path, ck, codes, dims: dict | None = None, lang: str = "zh"):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.formatting.rule import ColorScaleRule, IconSetRule, DataBarRule
    UP, DN, MID, GOLD, ZEBRA = "E5484D", "2FB877", "F2F2F2", "E3B93E", "F4F6F9"
    T = _XL_C2.get(lang, _XL_C2["zh"])
    dims = dims or {}
    wb = openpyxl.Workbook()
    ws = wb.active; ws.title = T["s1"]
    ws.append(T["h"])
    _xlsx_head(ws, 8)
    ri = 2
    for c in codes:
        ff = facts_for(ck, c)
        lh = ff.get("long_horizon", {}).get("data", {})
        uw = ff.get("underwater", {}).get("data", {})
        vp = ff.get("valuation_position", {}).get("data", {})
        ws.append([str(c), ck["by_code"][c].get("name", c),
                   round(lh.get("cagr", 0) * 100, 1) if lh else None,
                   round(lh.get("max_drawdown", 0) * 100, 1) if lh else None,
                   round(vp.get("percentile_rank"), 0) if vp and vp.get("percentile_rank") is not None else None,
                   round(uw.get("max_underwater_years"), 1) if uw else None,
                   round(vp.get("latest_dividend_yield", 0), 1) if vp else None,  # 已是百分比
                   dims.get(c, n_dims(ff))])                                       # 資料完整度揭露
        if ri % 2 == 0:
            for cc in range(1, 9):
                ws.cell(row=ri, column=cc).fill = PatternFill(start_color=ZEBRA, end_color=ZEBRA, fill_type="solid")
        ri += 1
    last = len(codes) + 1
    yr_fmt = '0.0"年"' if lang == "zh" else '0.0" y"'
    for col, fmt in (("C", '0.0"%"'), ("D", '0.0"%"'), ("E", '0"P"'), ("F", yr_fmt), ("G", '0.0"%"')):
        for row in range(2, last + 1):
            ws[f"{col}{row}"].number_format = fmt
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:H{last}"
    # C 欄「20年年化%」:值為正,高=好 → min 綠、max 紅(台股慣例 紅=好)
    ws.conditional_formatting.add(f"C2:C{last}", ColorScaleRule(
        start_type="min", start_color=DN, mid_type="percentile", mid_value=50, mid_color=MID,
        end_type="max", end_color=UP))
    # 🔴 phase3b A7:D 欄「最大回撤%」在本表是**負值**(-46.5/-75/-98.5),與 C1 那張表的
    # 正值幅度**符號相反**,舊版卻直接沿用 C1 的色階方向(min→紅/max→綠) → 最慘的 -98.5%
    # 被塗成本表自訂的「好」色(紅),最輕的 -22.3% 塗綠。同一張表裡紅在 C 欄=最好、在 D 欄=最糟,
    # 買家掃風險時會**讀反**。負值語意下,min(最慘)=綠、max(最輕)=紅 才與 C 欄同一套「紅=好」。
    ws.conditional_formatting.add(f"D2:D{last}", ColorScaleRule(
        start_type="min", start_color=DN, mid_type="percentile", mid_value=50, mid_color=MID,
        end_type="max", end_color=UP))
    ws.conditional_formatting.add(f"E2:E{last}", IconSetRule("3TrafficLights1", "percent", [0, 60, 95], showValue=True))
    ws.conditional_formatting.add(f"G2:G{last}", DataBarRule(start_type="min", end_type="max", color=GOLD))
    for i, w in enumerate([9, 18, 12, 13, 17, 20, 16, 13], start=1):
        ws.column_dimensions[chr(64 + i)].width = w
    c = ws.cell(row=last + 2, column=1, value=T["note"])
    c.font = Font(color="8A94A6", size=9, italic=True)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.row_dimensions[last + 2].height = 46
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))


def _f(v):
    try:
        return round(float(v), 3)
    except (TypeError, ValueError):
        return None


def _i(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _reconcile_tokenizer(prov: Provenance) -> None:
    """對齊守門 tokenizer 的切片假象:FG 的 %正則上限 4 位整數(\\d{1,4}),對 5 位數以上
    百分比(如回測 10027.7%)會切出子數(27.7)。作法:對每個『已入池的真來源大數字』的字串,
    枚舉 FG 會從中切出的 \\d{1,4}\\.\\d+ 尾段,補進池。只補真來源數字的切片,不放行憑空數字。"""
    add = set()
    for p in list(prov.pool):
        base = f"{p:.2f}"                     # 例 10027.73;涵蓋 FG 對整數/小數的切法
        if len(base.split(".")[0]) < 5:       # 只處理 ≥5 位整數(FG 才會誤切)
            continue
        for m in re.finditer(r"\d{1,4}\.\d+", base):
            try:
                v = float(m.group())
                add.add(abs(v)); add.add(abs(round(v, 1)))
            except ValueError:
                pass
    prov.pool |= add


# ── gate + 寫檔 ────────────────────────────────────────────────────────────────
def gate_sku(item: dict) -> list[dict]:
    """對 SKU 的內容 unit 過溯源守門(嚴容差);回查無來源的績效數字,非空 = fail-closed。
    先做 tokenizer 切片對齊(只補『真來源長數字被 FG 切出的子數』),再過**中英雙路徑**守門:
      - 中文:FG.extract_claims + FG._sourced_strict(原封沿用)
      - 英文:_extract_claims_en + FG._sourced_strict(補 FG 的中文盲區,見該函式註解)
    兩路都跑、對所有 SKU 都跑 → 覆蓋率只增不減。同值同 raw 去重,避免同一數字重複列。
    """
    _reconcile_tokenizer(item["prov"])
    text = RK.gate_text(item["gate_units"])
    bad = list(item["prov"].gate(text, item["sku"]))
    seen = {(round(c["value"], 2), c["raw"]) for c in bad}
    for c in _gate_en(item["prov"], text):
        k = (round(c["value"], 2), c["raw"])
        if k not in seen:
            seen.add(k)
            bad.append(c)
    return bad


def write_sku(item: dict) -> Path:
    sku_dir = OUT_ROOT / item["platform"] / f'{item["sku"]}_{item["lang"]}'
    sku_dir.mkdir(parents=True, exist_ok=True)
    for fn, kind, payload in item["artifacts"]:
        dest = sku_dir / fn
        if kind == "pdf":
            RK.render_pdf(payload, dest)          # 也寫 .html 側車
        elif kind == "xlsx":
            payload(dest)
        elif kind == "csv":
            dest.write_text(payload, encoding="utf-8")
    (sku_dir / "listing.json").write_text(
        json.dumps(item["listing"], ensure_ascii=False, indent=2), encoding="utf-8")
    p = item["listing"]["price"]
    price_line = " / ".join(f"{k}: {v}" for k, v in p.items())
    (sku_dir / "listing.md").write_text(
        f'# {item["listing"]["title"]}\n\n**平台**:{item["platform"]}　**定價**:{price_line}\n\n'
        f'## 描述\n\n{item["listing"]["description"]}\n\n## 標籤\n\n{", ".join(item["listing"]["tags"])}\n\n'
        f'---\n_{item["listing"]["seo_note"]}_\n', encoding="utf-8")
    (sku_dir / "_provenance.json").write_text(json.dumps({
        "sku": item["sku"], "lang": item["lang"], "generated_at": TODAY,
        "gate": "fact_source_guard._sourced_strict (空池嚴容差, fail-closed)",
        "pool_size": len(item["prov"].pool),
        "records": item["prov"].records,
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    return sku_dir


BUILDERS = {"M1": build_M1, "M2": build_M2, "T1": build_T1, "T2": build_T2,
            "C1": build_C1, "C2": build_C2}


def load_ctx() -> dict:
    return {
        "checkup": load_checkup(),
        "daytrade": load_daytrade_elig(),
        "adaptive": _load_csv(ADAPTIVE_CSV, floats=("a_net", "a_pf", "a_dd", "a_win"), ints=("a_tr",)),
        "longshort": _load_csv(LONGSHORT_CSV, floats=("ls_net",), ints=("short_trades",)),
        "perstock": _load_csv(PERSTOCK_CSV, floats=("net_profit_pct", "profit_factor", "max_dd_pct",
                                                    "sharpe", "return_over_maxdd")),
    }


def main() -> int:
    ap = argparse.ArgumentParser(description="電商商品工廠 v2")
    ap.add_argument("--sku", type=str, default=None, help="只產指定 SKU(M1/M2/T1/T2/C1/C2)")
    ap.add_argument("--list", action="store_true")
    ap.add_argument("--no-render", action="store_true", help="跳過 PDF 渲染(只組裝+gate,快速自測)")
    args = ap.parse_args()
    if args.list:
        for sid in BUILDERS:
            print(f"  {sid}")
        return 0
    ctx = load_ctx()
    print(f"[pf2] 來源:體檢 {len(ctx['checkup'].get('by_code',{}))} 檔、回測 {len(ctx['adaptive'])} 檔、"
          f"當沖適格 {ctx['daytrade'].get('_file','(無)')}")
    ids = [args.sku] if args.sku else list(BUILDERS)
    n_ok = n_block = 0
    for sid in ids:
        try:
            items = BUILDERS[sid](ctx)
        except Exception as exc:  # noqa: BLE001
            print(f"  ✗ {sid} 建造失敗:{exc}")
            continue
        if not items:
            print(f"  – {sid} 無可用資料,略過")
            continue
        for item in items:
            bad = gate_sku(item)
            if bad:
                n_block += 1
                print(f"  🔴 {item['sku']}_{item['lang']} 被溯源守門擋下(fail-closed):")
                for c in bad[:3]:
                    print(f"       查無來源數字 {c['value']} ←「{c['clause'][:40]}」")
                continue
            if args.no_render:
                print(f"  ✓(gate過) {item['sku']}_{item['lang']}  池 {len(item['prov'].pool)} 數")
                n_ok += 1
                continue
            d = write_sku(item)
            n_ok += 1
            print(f"  ✅ {item['sku']}_{item['lang']} → {d.relative_to(QUANT)}  "
                  f"({len(item['artifacts'])} 成品 + listing + provenance)")
    print(f"\n[pf2] 完成:{n_ok} 出檔、{n_block} 被守門擋下。根目錄 {OUT_ROOT.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
